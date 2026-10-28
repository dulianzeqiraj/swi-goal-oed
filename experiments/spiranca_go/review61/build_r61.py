"""Short review 61 (Word): 20 papers (Oxford, Caltech, Harvard, MIT, 2020-2026) on mathematical decisions for stakeholders
under intense rainfall, drought and causal change, how each changes MONTIMI-GO / SPIRANCA-GO, and the ROJA-GO test.
Reads res/sx.json (or res/sx_pilot.json), writes markdown, converts with pandoc (math -> OMML), then formats with python-docx."""
import json, os, subprocess, sys
import numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
RES = '/home/user/swi-goal-oed/experiments/spiranca_go/res'
src = os.path.join(RES, 'sx.json') if os.path.exists(os.path.join(RES, 'sx.json')) else os.path.join(RES, 'sx_pilot.json')
D = json.load(open(src))
rows = D['rows']; NA = len(rows); SC = D['scen']
PILOT = NA < 10
mean = lambda xs: float(np.mean(xs))
pc = lambda v, d=0: f"{100 * v:.{d}f}%"

def cl(sc, pol, key):
    return mean([r['closed'][sc][pol][key] for r in rows])

def op(sc, m, key):
    return mean([r['open'][sc][m][key] for r in rows])

def n10(sc, m):
    return sum(r['open'][sc][m]['over_s'] > D['alpha'] for r in rows)

LAB = {'none': 'Ordinary weather', 'rain': 'Intense rainfall (gauge +2.5 sd for 2 years, a fifth infiltrates)',
       'drought': 'Record-shattering drought (−2 sd for 6 years)', 'causal': 'Same drought, farmers respond 3× more strongly (causal change)',
       'compound': 'Drought, then 2 years of intense rain on dry ground'}

# ---------------------------------------------------------------- paper cards (from the verified research notes)
P = {
'Oxford': [
 ('Kelder, Müller, Slater, … Nipen (2020). Using UNSEEN trends to detect decadal changes in 100-year precipitation extremes. npj Clim. Atmos. Sci. 3, 47. doi:10.1038/s41612-020-00149-4',
  'Intense rainfall',
  'UNSEEN: pool seasonal-forecast members and lead times into ~3,500 "years" (≈100× the record) and fit a GEV with drifting location and scale, $\\mu(t)=\\mu_0+\\mu_1 t$, after tests of member independence, lead-time stability and fidelity.',
  'Sets the extreme recharge inputs and the shift radius $\\bar\\delta$ of the SPIRANCA-GO margin from the pooled ensemble, not from the 80-year record; widens the storylines the alarm must survive.',
  'UNSEEN-certified margin: $\\bar\\delta$ = UNSEEN 100-year anomaly of the gauge.'),
 ('Kew, Philip, Hauser, … Funk, Otto (2021). Impact of precipitation and increasing temperatures on drought trends in eastern Africa. Earth Syst. Dynam. 12, 17–35. doi:10.5194/esd-12-17-2021',
  'Drought; causal attribution',
  'Probabilistic event attribution: probability ratio $PR=p_1/p_0$ of a drought index in the factual and counterfactual climate, per driver.',
  'Warming can dry the aquifer through evaporation while rainfall is unchanged: that shift lies outside the anchor span $\\{r,p\\}$, so the certificate does not cover it. Push: add an evaporation (temperature) anchor.',
  'Span test before certifying: is the attributed driver in span(anchors)?'),
 ('Murgatroyd & Hall (2020). The resilience of inter-basin transfers to severe droughts with changing spatial characteristics. Front. Environ. Sci. 8, 571647. doi:10.3389/fenvs.2020.571647',
  'Drought; planning',
  'Stochastic, spatially coherent severe-drought scenarios; resilience of a supply system with transfers and storage under joint drought.',
  'Well placement judged on joint failure (chloride above 250 mg/L when the back-up source also fails); both depend on the same drought, so $P(A\\cap B)\\ge P(A)P(B)$.',
  'Joint-failure objective for siting the supply well.'),
 ('Becher, Pant, Verschuur, … Hall (2023). A multi-hazard risk framework to stress-test water supply systems to climate-related disruptions. Earth\'s Future 11, e2022EF002946. doi:10.1029/2022EF002946',
  'Drought, flood, cyclone',
  'One metric for all hazards: expected customer disruption days $\\sum_h P_h\\,N_h\\,T_h$; ≈5 days per customer per year now, ×2 to ×2.5 by 2100.',
  'Replaces the yes/no 250 mg/L goal by expected salinity-disruption days plus a tail guard; can favour a reserve well over one "safest" site.',
  'Disruption-days goal in MONTIMI-GO placement.'),
 ('Jesson, Douglas, Manshausen, … Gal, Shalit (2022). Scalable sensitivity and uncertainty analyses for causal-effect estimates of continuous-valued interventions. NeurIPS 35 (no DOI; arXiv:2204.10022)',
  'Causal; hidden confounding',
  'Bounded hidden confounding for a continuous dose: $\\Lambda^{-1}\\le p(t\\mid x)/p(t\\mid x,u)\\le\\Lambda$ turns the dose–response into an interval.',
  'Private abstraction in a drought confounds pumping → chloride; a $\\Lambda$-robust margin covers it; $\\Lambda_{\\mathrm{crit}}$ (the confounding that overturns "safe") is the number to give stakeholders.',
  'Report $\\Lambda_{\\mathrm{crit}}$ next to every pumping decision.')],
'Caltech': [
 ('Yu, Shi, Chung, Yue & Wierman (2020). The power of predictions in online control. NeurIPS 33 (no DOI; arXiv:2006.07569)',
  'Normal operation; forecasts',
  'Model-predictive control with $k$ predicted disturbances in a linear-quadratic problem: $O(\\log T)$ predictions give constant regret, matching the lower bound.',
  'How many seasons of recharge forecast a pumping controller needs; the rainfall persistence (lag-one correlation 0.7) gives part of it for free; slow clay memory lengthens it.',
  'Forecast-horizon rule for curtailment.'),
 ('Schneider, Dunbar, Wu, … Shaman (2022). Epidemic management and control through risk-dependent individual contact interventions. PLOS Comput. Biol. 18, e1010171. doi:10.1371/journal.pcbi.1010171',
  'Crisis operation',
  'Network data assimilation gives each node\'s risk; act only where $P(\\text{risk}_i)>\\theta$. Fewer than half the deaths of a lockdown while isolating 5–10%.',
  'Curtail well by well when the posterior $P(\\mathrm{Cl}>250\\ \\mathrm{mg/L})$ passes a threshold, from sparse samples.',
  'Risk-threshold curtailment (used as the decision rule of ROJA-GO).'),
 ('Kurth, Subramanian, … Kashinath, Anandkumar (2023). FourCastNet: accelerating global high-resolution weather forecasting using adaptive Fourier neural operators. PASC \'23. doi:10.1145/3592979.3593412',
  'Intense rainfall',
  'Neural-operator weather emulator, orders of magnitude faster than numerical weather prediction, so ensembles large enough to sample rare extremes.',
  'Rainfall tails at the gauge and rain-plus-surge storylines for the alarm margin; it forecasts days to weeks, so it says nothing on multi-year drought.',
  'Emulator ensemble of intense-rain storylines.'),
 ('Humphrey, Berg, Ciais, … Frankenberg (2021). Soil moisture–atmosphere feedback dominates land carbon uptake variability. Nature 592, 65–69. doi:10.1038/s41586-021-03325-5',
  'Drought; causal (feedback)',
  'Factorial runs with soil moisture held at climatology separate the direct effect from the feedback effect; ≈90% of year-to-year uptake variation goes through soil moisture.',
  'A $2^3$ factorial of MONTIMI-GO separates the direct recharge effect of a drought from the human feedback (private abstraction); the anchor certificate does not cover a changed feedback.',
  'The "causal" storyline of the test below.'),
 ('Frederikse, Landerer, Caron, … Humphrey, … Wu (2020). The causes of sea-level rise since 1900. Nature 584, 393–397. doi:10.1038/s41586-020-2591-3',
  'Sea boundary; causal budget',
  'Closed sea-level budget attributing the rise to named causes (glaciers and ice sheets twice thermal expansion).',
  'Build the sea-boundary head from components plus land motion and surge; the adjoint ranks which one matters for chloride at the screen.',
  'Component-wise sea boundary with adjoint ranking.')],
'Harvard': [
 ('Considine, Nethery, Wellenius, Dominici & Tec (2025). Optimizing heat alert issuance with reinforcement learning. AAAI-25 39(27). doi:10.1609/aaai.v39i27.35009',
  'Extreme heat; alerts',
  'Reinforcement learning of when to issue alerts under an alert budget; beats the National Weather Service policy in 30 counties.',
  'A budgeted alarm: the number of curtailments or extra samples per year is capped, and spent where they save most.',
  'Sample budget of ROJA-GO (≤ one extra sample a year).'),
 ('Xu, Perrault, Fang, Chen & Tambe (2021). Robust reinforcement learning under minimax regret for green security. UAI 2021, PMLR 161:257–267 (no DOI)',
  'Decision under uncertainty',
  'Minimax-regret policy, $\\min_\\pi\\max_{\\theta}[V_{\\mathrm{best}}(\\theta)-V^\\pi(\\theta)]$, solved as a game against nature (MIRROR).',
  'Minimax-regret well placement over recharge, private abstraction and sea level, with the MONTIMI-GO adjoint searching the worst case.',
  'Regret-robust siting.'),
 ('Wu, Weinberger, Wellenius, Dominici & Braun (2024). Assessing the causal effects of a stochastic intervention in time series data: are heat alerts effective in preventing deaths and hospitalizations? Biostatistics 25, 57–79. doi:10.1093/biostatistics/kxad002',
  'Causal; heat alerts',
  'Incremental intervention: multiply the odds of an alert by $\\delta$; identified under weaker overlap. Evidence of benefit is weak.',
  'Pumping depends on the drought state, so the pumping anchor is not exogenous; a decision is a shift of the pumping odds.',
  'Largest certifiable $\\delta$ before an early chloride sample is needed.'),
 ('Nethery, Katz-Christy, Kioumourtzoglou, Parks, Schumacher & Anderson (2023). Integrated causal-predictive machine learning models for tropical cyclone epidemiology. Biostatistics 24, 449–464. doi:10.1093/biostatistics/kxab047',
  'Storms; causal',
  'Stage 1 estimates the causal effect of each storm at each place; stage 2 predicts these effects from storm and place traits.',
  'A per-well storm/flood vulnerability, learned from simulator runs where the true effect is known.',
  'Two-stage vulnerability map of wells.'),
 ('Proctor, Rigden, Chan & Huybers (2022). More accurate specification of water supply shows its importance for global crop production. Nature Food 3, 753–763. doi:10.1038/s43016-022-00592-x',
  'Drought; proxy error',
  'Soil moisture explains 30–120% more yield variation than precipitation; temperature-plus-precipitation models overstate damages by 28–320%.',
  'Rainfall is a poor proxy of the water that reaches the ground: in intense rain most of it runs off, so the gauge anchor breaks. This is the "intense rainfall" storyline below.',
  'Anchor-adequacy test (gauge vs soil moisture).')],
'MIT': [
 ('Bertsimas & Kallus (2020). From predictive to prescriptive analytics. Management Science 66, 1025–1044. doi:10.1287/mnsc.2018.3253',
  'Normal operation',
  'Decision from data: $\\hat z(x)=\\arg\\min_z\\sum_i w_i(x)\\,c(z;y_i)$ with weights from kNN, trees or forests, scored by a coefficient of prescriptiveness.',
  'Pumping chosen from the covariates the utility has (last season\'s rainfall, heads); cut after a dry year.',
  'Prescriptiveness as the value of a monitoring well.'),
 ('Barthel Sorensen, Charalampopoulos, Zhang, Harrop, Leung & Sapsis (2024). A non-intrusive machine learning framework for debiasing long-time coarse resolution climate simulations and quantifying rare events statistics. JAMES 16, e2023MS004122. doi:10.1029/2023MS004122',
  'Extremes',
  'Neural correction operator trained on a short nudged run; return periods two orders of magnitude beyond the training data; 8 years of training gave 30-year statistics close to ERA5.',
  'Debiased long runs of the recharge drivers give the return period of an exceedance at the screen.',
  'Exceedance return periods for the supply well.'),
 ('Schlosser, Sokolov, Strzepek, Thomas, Gao & Arndt (2021). The changing nature of hydroclimatic risks across South Africa. Climatic Change 168, 28. doi:10.1007/s10584-021-03235-5',
  'Drought',
  '7,200-member hybrid ensembles (global Monte Carlo × regional patterns); drying risk 3–4× the wetting risk in western and central South Africa.',
  'A drifting, skewed recharge: the margin should be asymmetric, larger on the drought side.',
  'Drought-side margin.'),
 ('Abadie (2021). Using synthetic controls: feasibility, data requirements, and methodological aspects. J. Econ. Lit. 59, 391–425. doi:10.1257/jel.20191450',
  'Causal',
  'Synthetic control: weights $w_j\\ge 0,\\ \\sum w_j=1$ match the pre-change path; the counterfactual is $\\sum_j w_j Y_{jt}$.',
  'Counterfactual chloride path after the pumping change, from donor wells outside the cone of depression.',
  'Synthetic-control check of the decision effect.'),
 ('Chernozhukov, Wüthrich & Zhu (2021). An exact and robust conformal inference method for counterfactual and synthetic controls. JASA 116, 1849–1864. doi:10.1080/01621459.2021.1920957',
  'Causal; break detection',
  'Permutation (conformal) tests for counterfactuals; exact under exchangeable residuals, approximately valid under dependence.',
  'A test for a regime break triggers re-anchoring when needed, not only once a year.',
  'The trigger of ROJA-GO.')],
}

# ---------------------------------------------------------------- numbers
exc = {s: mean([r['exc'][s] for r in rows]) for s in SC}
POLS = [('none', 'No action'), ('spir', 'Static alarm'), ('spir_ra', 'Yearly sample (S5)'), ('spir_roja', 'ROJA-GO'), ('oracle', 'Screen sensor')]

md = []
A = md.append
A('---\ntitle: "Mathematical decisions for stakeholders under intense rainfall, drought and causal change"\n'
  'subtitle: "Review 61 — twenty papers from Oxford, Caltech, Harvard and MIT (2020–2026), what each changes in MONTIMI-GO / SPIRANCA-GO, and the ROJA-GO innovation"\n'
  'author: "Dulian Zeqiraj"\n---\n')
A('## 1. Outcome\n')
A(f'- **Twenty papers, five per institution**, give stakeholders mathematical tools for normal operation, intense rainfall, drought and causal cases (Section 2). Each one changes a definite part of our model: an input, a boundary, the goal, the margin or the decision rule; each card ends with an innovation.')
A(f'- **Extremes change the model in three ways** (Section 3): they go beyond the record (drought), they break the link between what is recorded and what acts (intense rain that runs off), and they change a causal mechanism (farmers who pump far more in a drought they never saw).')
A(f'- **Innovation, ROJA-GO (the watch).** Extremes speed up the drift that re-anchoring must follow. The optimal re-anchoring window then shrinks to the last sample (Proposition R1), and an extra chloride sample is taken when a detector that needs no chloride fires.')
nm = f'{NA} aquifer{"s" if NA > 1 else ""}'
A(f'- **Test on the variable-density section ({nm}{", pilot" if PILOT else ""}, {D["ntraj"]} weather sequences each).** Without action the screen exceeds the standard in {pc(exc["none"])} of the seasons in ordinary weather and in {pc(exc["drought"])}, {pc(exc["causal"])} and {pc(exc["compound"])} under the drought, the causal change and the compound event. With the decision rule "cut pumping after an alarm", the ROJA-GO watch keeps exceedance to {pc(cl("drought","spir_roja","exc"))}, {pc(cl("causal","spir_roja","exc"))} and {pc(cl("compound","spir_roja","exc"))} (yearly sample alone: {pc(cl("drought","spir_ra","exc"))}, {pc(cl("causal","spir_ra","exc"))}, {pc(cl("compound","spir_ra","exc"))}), with {pc(cl("drought","spir_roja","water"))} of the planned water delivered in the drought.')
if PILOT:
    A(f'- *These are pilot numbers on {nm}; the run on 20 aquifers will replace them.*')
A('')

A('## 2. The twenty papers\n')
A('Each row: the paper (verified DOI or proceedings), the situation, the mathematical tool, what it changes in our model, and the innovation it gives.\n')
k = 0
for inst in ['Oxford', 'Caltech', 'Harvard', 'MIT']:
    A(f'### {inst}\n')
    A('| # | Paper | Situation | Mathematical tool for the decision | What it changes in our model | Innovation |')
    A('|---|---|---|---|---|---|')
    for (cit, sit, tool, eff, inn) in P[inst]:
        k += 1
        A(f'| {k} | {cit} | {sit} | {tool} | {eff} | {inn} |')
    A('')

A('## 3. How each situation enters the model\n')
A('| Situation | Enters MONTIMI-GO / SPIRANCA-GO through | What it does to the decision | Papers |')
A('|---|---|---|---|')
A('| Ordinary weather | inland inflow scaled by recharge; rainfall anchor $r$ (persistence 0.7) | the static certificate holds; forecasts of a few seasons are enough | 6, 16 |')
A('| Intense rainfall | the gauge records much more than infiltrates: $r_{\\mathrm{gauge}}\\neq r_{\\mathrm{eff}}$ | the alarm reads the gauge and predicts a freshening that does not come; the anchor itself breaks | 1, 8, 14, 15, 17 |')
A('| Drought | inland inflow falls for years, beyond any season of the record | extrapolation beyond the record; drift speeds up | 2, 3, 4, 9, 18 |')
A('| Causal change | private abstraction responds more strongly than in the record (a changed structural equation, not a shift in span) | not covered by the anchor certificate; only measurements catch it | 5, 9, 13, 19, 20 |')
A('| Sea boundary | sea level and surge at the sea face | slow, attributable trend plus short surges | 4, 10 |')
A('| Decision rules | when to cut pumping, how many samples, where to site | budgets, risk thresholds, regret, prescriptions | 6, 7, 11, 12, 16 |')
A('')

A('## 4. Innovation: ROJA-GO, the watch that samples when an extreme leaves the record\n')
A('After the decision, the alarm error is $\\hat Y_t-Y_t=-\\eta_t+d_t$, where $d_t$ is the drift the record could not foresee, with $|d_t-d_s|\\le\\rho|t-s|$. Re-anchoring with the mean error of the last $K$ samples, taken $\\Delta,2\\Delta,\\dots,K\\Delta$ before $t$, leaves a lag of at most $\\rho\\Delta(K+1)/2$ (Proposition S5 of the SPIRANCA-GO note). An extreme makes $\\rho$ large.\n')
A('**Proposition R1 (the window shrinks with the drift).** With independent noise of variance $\\sigma^2$,\n')
A('$$\\mathbb{E}\\left[(\\hat Y_t-\\bar e_{t,K}-Y_t)^2\\right]\\le\\sigma^2\\left(1+\\frac{1}{K}\\right)+\\frac{\\rho^2\\Delta^2(K+1)^2}{4},$$\n')
A('and the bound is smallest at the $K$ that solves $K^2(K+1)=2\\sigma^2/(\\rho\\Delta)^2$. In particular, if the drift between two samples is at least the noise, $\\rho\\Delta\\ge\\sigma$, then $K_{\\mathrm{opt}}=1$: use the last sample only. *Proof.* Noise and lag terms are bounded separately (S5); the derivative of the bound in $K$ is $-\\sigma^2/K^2+\\rho^2\\Delta^2(K+1)/2$, which is increasing, and at $K=1$ it is $\\rho^2\\Delta^2-\\sigma^2\\ge 0$. ∎\n')
A('**The watch.** Every season, without any chloride: (i) the gauge leaves the range of the record (one season or the two-year filter), or (ii) the head residuals, after removing what the anchors explain, jump by more than anything in the record (the trigger of paper 20). When either fires, one extra chloride sample is taken (at most one extra a year, the budget of paper 11), and the window is $K=1$ while the extreme lasts or while two consecutive samples differ by more than three record standard deviations (R1). Otherwise $K=3$ and one sample a year, as in S5.\n')
A(f'**The stakeholder decision (closed loop).** The utility pumps the decided +{D["u_dec"]} unless the alarm sounded in the previous season; then it pumps {D["u_cut"]} (the risk-threshold rule of paper 7). For comparison: no action; a static alarm; yearly re-anchoring; and a continuous chloride sensor at the screen that cuts pumping when chloride reaches 89% of the standard.\n')
A(f'Table 1. Share of seasons above 250 mg/L at the supply screen over 30 years after the decision, mean over {nm}{" (pilot)" if PILOT else ""}; in brackets the share of the planned water delivered.\n')
A('| Storyline | ' + ' | '.join(l for _, l in POLS) + ' |')
A('|---|' + '---|' * len(POLS))
for s in SC:
    cells = []
    for p, _ in POLS:
        e = cl(s, p, 'exc'); w = cl(s, p, 'water')
        cells.append(f'{pc(e)}' + ('' if p == 'none' else f' ({pc(w)})'))
    A(f'| {LAB[s]} | ' + ' | '.join(cells) + ' |')
A('')
A(f'Table 2. Alarm validity in open loop (no curtailment): share of seasons in the stress period in which the alarm under-warns by more than its margin (target ≤ {pc(D["alpha"])}), mean over {nm}; in brackets the aquifers above the target; and chloride samples in 30 years.\n')
OM = [('spir', 'Static'), ('spir_ra', 'Yearly sample'), ('spir_trig', 'Extra samples only'), ('spir_adK', 'Adaptive window only'), ('spir_roja', 'ROJA-GO'), ('olsA_roja', 'Least squares + ROJA')]
A('| Storyline | ' + ' | '.join(l for _, l in OM) + ' |')
A('|---|' + '---|' * len(OM))
for s in SC:
    A(f'| {LAB[s]} | ' + ' | '.join(f'{pc(op(s, m, "over_s"), 1)} ({n10(s, m)})' for m, _ in OM) + ' |')
A('| Samples in 30 years | ' + ' | '.join(f'{mean([r["open"][s][m]["nsamp"] for r in rows for s in SC]):.0f}' for m, _ in OM) + ' |')
A('')
fire_s = mean([r['fire'][s]['any_stress'] for r in rows for s in SC if s != 'none'])
fire_c = mean([r['fire'][s]['any_calm'] for r in rows for s in SC if s != 'none'])
A(f'**What the test shows.** (1) The extremes, not the decision, now drive the risk: the drought multiplies the exceeding seasons by {exc["drought"] / max(exc["none"], 1e-9):.1f} and the causal change by {exc["causal"] / max(exc["none"], 1e-9):.1f}. (2) The detectors fire in {pc(fire_s)} of the seasons of an extreme and {pc(fire_c)} outside it, so the extra samples go where the drift is. (3) In the three drought storylines a continuous sensor that reacts at 89% of the standard leaves {pc(mean([cl(s, "oracle", "exc") for s in ("drought", "causal", "compound")]))} of the seasons above it, against {pc(mean([cl(s, "spir_roja", "exc") for s in ("drought", "causal", "compound")]))} with ROJA-GO: the aquifer has memory, so acting at the standard is late, and the certified margin acts earlier. (4) The certificate is still needed: least squares with the same watch under-warns in {pc(mean([op(s, "olsA_roja", "over_s") for s in SC]))} of the stress seasons.\n')
A('**For Fushë-Kuqe.** Before a drought season, agree the rule with the stakeholders in this form: pump as decided while the alarm is silent; cut when it sounds; take one chloride sample a year at the supply screen, and one more as soon as the rain gauge or the piezometers leave the range of the record. Report with every decision the margin, the samples spent and the confounding strength $\\Lambda_{\\mathrm{crit}}$ that would overturn it (paper 5).\n')

A('## 5. Verification and caveats\n')
A('- The network of this session blocked Crossref, doi.org, arXiv and several publisher hosts. Every DOI and affiliation was therefore confirmed from publisher or repository records found by search (ORA, CaltechAUTHORS, Harvard DASH, DSpace@MIT, PMC, RePEc, ADS), not from Crossref. Pages of the NeurIPS, UAI and PASC papers and some article numbers are not verified.')
A('- Numbers flagged as uncertain in the notes are not used here: Wu et al. (2024) report two different intervals in two versions; per-region numbers of Kew et al. (2021) and the coincidence probabilities of Murgatroyd & Hall (2020) were not retrieved.')
A('- Kurth et al. (2023) and Frederikse et al. (2020) have Caltech authors together with NVIDIA and JPL authors. None of the twenty papers is about seawater intrusion; their use here is our transfer, and the test is on a synthetic 2-D section, not on field data.')
A(f'- Code: experiments/spiranca_go/sx_run.py (storylines, watch, closed loop); results res/sx.json. All 20 papers are new to this series: none is in the 3,162 references of the earlier reviews.')

mdp = os.path.join(HERE, 'r61.md')
open(mdp, 'w').write('\n'.join(md))
out = os.path.join(HERE, 'r61_raw.docx')
subprocess.run(['pandoc', mdp, '-o', out, '--from', 'markdown+tex_math_dollars+pipe_tables'], check=True)

# ---------------------------------------------------------------- formatting
doc = Document(out)
for s in doc.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.6); s.top_margin = s.bottom_margin = Cm(1.5)
st = {x.name: x for x in doc.styles}
for nmst, size in (('Normal', 10), ('Body Text', 10), ('First Paragraph', 10), ('Compact', 8)):
    if nmst in st:
        st[nmst].font.size = Pt(size); st[nmst].font.name = 'Calibri'
for nmst, size, col in (('Title', 20, '1F3864'), ('Subtitle', 12, '2E5597'), ('Heading 2', 14, '1F3864'), ('Heading 3', 12, '2E5597')):
    if nmst in st:
        st[nmst].font.size = Pt(size); st[nmst].font.color.rgb = RGBColor.from_string(col); st[nmst].font.name = 'Calibri'

def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '9DB2D6'); b.append(el)
    tblPr.append(b)

WID = {6: [0.7, 6.3, 2.5, 6.4, 6.4, 4.2], 4: [3.4, 8.0, 10.0, 5.0]}
for t in doc.tables:
    borders(t)
    nc = len(t.columns)
    w = WID.get(nc) if t.rows[0].cells[0].text.strip() in ('#', 'Situation') else None
    if w:
        tw = [int(x * 567) for x in w]
        for gc, x in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), tw):
            gc.set(qn('w:w'), str(x))
        tblPr = t._tbl.tblPr
        for e in tblPr.findall(qn('w:tblW')):
            e.set(qn('w:type'), 'dxa'); e.set(qn('w:w'), str(sum(tw)))
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    for i, row in enumerate(t.rows):
        for j, c in enumerate(row.cells):
            if w: c.width = Cm(w[j])
            for p in c.paragraphs:
                for rn in p.runs:
                    rn.font.size = Pt(8 if nc >= 6 else 8.5)
                    if i == 0: rn.font.bold = True
            if i == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'DCE6F2'); c._tc.get_or_add_tcPr().append(sh)
final = os.path.join(HERE, 'Literature_review_Stakeholder_Decisions_IV_Intense_Rainfall_Drought_Causal_Oxford_Caltech_Harvard_MIT_2020-2026.docx')
doc.save(final)
print('aquifers', NA, 'pilot' if PILOT else 'full', '->', final)
