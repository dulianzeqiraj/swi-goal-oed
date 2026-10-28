"""Builds the reference list of the HALOMAR Technical Annex (Word, Arial 11): every work the annex cites, in the
author-year style of the in-text citations. Entries are read from refs_data.py."""
import os
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from refs_data import SCI, LAW, COST, NOTES  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else 'HALOMAR_References_Technical_Annex.docx'


def para(doc, text, bold=False, italic=False, size=11, space_after=4, hanging=True, align=None):
    p = doc.add_paragraph()
    if hanging:
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = 'Arial'; r.font.size = Pt(size)
    return p


def main():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(2.0)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    para(doc, 'References', bold=True, size=13, space_after=8, hanging=False)
    para(doc, 'A. Scientific literature', bold=True, space_after=6, hanging=False)
    for e in sorted(SCI, key=lambda t: t[0]):
        para(doc, e[1])
    para(doc, 'B. European law, policy and standards', bold=True, space_after=6, hanging=False).paragraph_format.space_before = Pt(8)
    for e in LAW:
        para(doc, e)
    para(doc, 'C. Related COST Actions', bold=True, space_after=6, hanging=False).paragraph_format.space_before = Pt(8)
    for e in COST:
        para(doc, e)
    if NOTES:
        para(doc, 'Notes for the applicant (delete before submission)', bold=True, space_after=6, hanging=False).paragraph_format.space_before = Pt(10)
        for n in NOTES:
            para(doc, '• ' + n, italic=True, size=10)
    doc.save(OUT)
    n = len(SCI) + len(LAW) + len(COST)
    print('saved', OUT, n, 'entries')


if __name__ == '__main__':
    main()
