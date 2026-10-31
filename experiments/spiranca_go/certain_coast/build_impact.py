"""CERTAIN-COAST (NATO SPS proposal): Impact. What is at stake (verified figures), the impact pathway from activities to
security impacts (figure), who benefits, impact indicators with baselines and targets, the certified value of a warning
for each user (cost-loss model with the certified miss rate: equations in OMML and a figure), contribution to NATO and
partner security, scale-out, conditions for impact and what is new.
Results, deliverables, milestones, risks and running cost are read from the other builders for consistency.

usage: python build_impact.py OUT.docx"""
import json
import os
import sys
import tempfile

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, '..', '..', 'ml_paradigms', 'word'))
from build_sota import run, para, heading, rich, table, shade, link  # noqa: E402
from build_letters import run as lrun  # noqa: E402
import re  # noqa: E402


def ph(p, text, size=11, bold=False):
    """Text with [placeholders] highlighted (numeric citations such as [1, 2] are left plain) and **bold** segments."""
    for part in re.split(r'(\[[^\]]+\])', text):
        if not part:
            continue
        if part.startswith('[') and not re.fullmatch(r'\[[\d,\s–-]+\]', part):
            lrun(p, part, size=size, bold=bold, hl=True)
        else:
            for j, seg in enumerate(part.split('**')):
                if seg:
                    lrun(p, seg, size=size, bold=bold or j % 2 == 1)
import build_workplan as W  # noqa: E402
import omml  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')

# illustrative values for the worked example (the real ones come from back-testing, D3.2, and from the users, T3.1)
S_EX, F_EX, ALPHAS = 0.05, 0.05, (0.05, 0.10)

STAKE = [
    '**The 2019 earthquake.** The Mw 6.4 earthquake of 26 November 2019 near Durrës killed 51 people and affected 202,291; '
    'the Post-Disaster Needs Assessment put its total effects at €985 million (€844 million of damage and €141 million of '
    'losses) [1, 2]. Field surveys found liquefaction at coastal, riverine and lagoonal sites [3].',
    '**The port.** The Port of Durrës handled 94.2% of the freight carried by sea in Albania in the second quarter of 2026 [4]; '
    'a disruption of its quays is a disruption of the country\'s maritime trade.',
    '**Water and energy.** The wellfields supply drinking water to the Durrës area [population served, to confirm with the '
    'utility], and a well that turns saline takes years to recover; the energy landfall connects supply arriving from the '
    'sea to the national system.',
    '**The gap.** These assets are watched separately, and no warning at the site states how often it misses. A decision-maker '
    'therefore cannot tell what a warning is worth, or what silence means.',
]

BENEFIT = [
    ['Water utility (wellfields)', 'Certified intrusion warning before wells turn saline; a pumping rule safe for all three hazards; fewer wells lost to salt', 'Warnings from M23; rule by M33'],
    ['Port authority', 'Liquefaction and extreme-water-level warnings for the quays; assessment within days after an earthquake or surge', 'From M23; post-event protocol from M12'],
    ['Energy landfall operator', 'One stability check combining shaking, liquefaction, water level and subsidence; continuous monitoring', 'From M23'],
    ['National Civil Protection Agency', 'One consistent picture of water, port and energy after an event, with stated reliability', 'From M23'],
    ['National authorities', 'Relative sea level measured at the assets for planning; a policy brief on coastal subsidence and critical infrastructure', 'M12 (D1.1); M36 (D6.2)'],
    ['People of the Durrës area', 'Safer drinking water and continuity of port and energy services (indirect)', 'During and after the project'],
    ['Partner-country institution', 'Researchers trained in every work package; methods and software to apply on its own coasts', 'M3–M36 and after'],
    ['Scientific community', 'Coupled model, certified alarms, open software, data and a synthetic test coast', 'M24–M36'],
    ['NATO and partners', 'A tested way to protect critical coastal infrastructure with warnings of known reliability', 'M36 and after'],
]

KPI = [
    ['Hazards covered by one warning with a stated miss rate', 'None at the site', '3 (intrusion, liquefaction, extreme sea level)', '3, run by the operators', 'D3.1; D5.1'],
    ['Miss rate of each alarm', 'Not stated', '≤ αₖ (indicatively 5–10%) in back-testing and live', 'Still ≤ αₖ in the continued bulletin log', 'D3.2; D5.1; operators\' log'],
    ['Months of live operation with logged outcomes', '0', '12', '[≥ 36 in total]', 'D5.1; bulletin log'],
    ['Value of the warning to each user (Section 5)', 'Unknown', 'Certified lower bound reported for each user', 'Recomputed yearly from the log', 'D5.1; operators\' log'],
    ['Assets with an operating procedure adopted', '0', '3 (at least one per asset)', '3, reviewed after each event', 'D6.2'],
    ['Trained operator and civil-protection staff', '0', '≥ 20 in total; ≥ 2 per operator', '≥ 2 per operator kept', 'Training records (T6.2)'],
    ['Early-career researchers trained', '—', '≥ 6 (NATO and partner country)', '[Careers in the field]', 'T6.1 records'],
    ['Open-access publications', '—', '≥ 6 (indicatively)', '[Further papers from the data]', 'Publication list'],
    ['Open software and data with a DOI', '0', '1 release, installed by an outsider (D6.1)', '[Used by ≥ 1 group outside the consortium]', 'Repository statistics'],
    ['Post-event protocol executed', '—', 'For every event above its trigger', 'Continued by the operators', 'D5.1; event reports'],
    ['Running cost carried by the operators', '—', 'Maintenance agreement signed (M33)', 'About €{cost:,.0f} per year paid', 'Agreement; accounts'],
    ['Coasts applying the method', '0', '1 (Durrës)', '[≥ 2]', 'Transfer reports'],
]

USERS = [
    ['Water utility', 'Reduce or move pumping away from the coastal wells', 'Water bought or pumped elsewhere for the warning period', 'A well turned saline for years', '[r₁]'],
    ['Port authority', 'Restrict heavy loads and berthing on the affected quays', 'Hours of reduced quay operation', 'Damage to quays, cranes and cargo', '[r₂]'],
    ['Energy landfall operator', 'Inspect and take protective measures at the landfall', 'Inspection and reduced operation', 'Interruption of supply; repair', '[r₃]'],
    ['Civil protection', 'Pre-position teams and advise the public', 'Staff time and public advice', 'Slower response; harm to people', '[r₄]'],
]

SECURITY = [
    '**Resilience of an Ally.** The project supports three of NATO\'s seven baseline requirements for civil preparedness: resilient food and water resources, resilient energy supplies and resilient civil transport [5]; see the Security Relevance section.',
    '**SPS Key Priorities.** Environment, Climate Change and Security; Energy Security; and Innovation and Emerging and Disruptive Technologies [6].',
    '**Responsible use of AI.** Machine learning enters a security decision chain only with a certified error bound and a verified miss rate, so its use is traceable and accountable.',
    '**Partner capacity.** Researchers of the partner country are trained in every work package and take the methods and open software home.',
]

SCALE = [
    '**What a new coast needs.** Free Sentinel-1 data, GNSS and a few corner reflectors, piezometers with conductivity loggers, existing seismic stations or fibre where available, and the open software of D6.1. The synthetic test coast lets a new team check its installation before using its own data.',
    '**How it is passed on.** The transfer workshop in Month 35, the training material and the operating procedures are written for teams outside the consortium; D6.1 is accepted only when someone outside the team installs and runs it.',
    '**Where next.** Other subsiding coasts with wellfields, ports or energy landfalls, in the Adriatic, the Mediterranean and the partner country. [Candidate sites and follow-up funding to be identified with the end-users and the partner-country team.]',
]

CONDITIONS = [
    ['Operators use the warnings', 'Warnings are ignored or not trusted', 'Letters of commitment; thresholds and cost-loss ratios agreed with the operators (T3.1); operators co-sign MS1, MS4, MS6, MS7; every bulletin logged with its outcome', 'R15, R16'],
    ['The system keeps running', 'No budget or staff after the grant', 'Maintenance agreement by M33; running cost about €{cost:,.0f} per year shared by the operators; at least two trained people per operator', 'R18'],
    ['Guarantees stay valid', 'Conditions move beyond the record', 'Bulletins flag "outside record"; margins widen; guarantees rechecked yearly from the log', 'R8'],
    ['Openness does not expose assets', 'Vulnerability details released', 'Two-tier dissemination: methods and aggregated results open, asset-level results restricted; operators review releases', 'R20'],
    ['Others can repeat it', 'Method tied to one site or team', 'Open software and data with a DOI; synthetic test coast; outsider installation test (D6.1); transfer workshop', '—'],
]

NEW = [
    '**A certified value for each user.** The certified miss rate is turned into a guaranteed lower bound on the economic value of the warning for each user\'s own decision (Section 5), so the operators know in advance what the warning is at least worth to them.',
    '**Impact measured by the same log that checks the guarantees.** Every bulletin is logged with the action taken and the outcome; the same record verifies the miss rates, measures the value delivered and shows what was avoided.',
    '**Impact defined in the users\' terms.** The cost of acting and the loss avoided are set with each operator before live operation (T3.1), not chosen afterwards by the project.',
    '**One warning, three assets, one picture.** The impact on water, port and energy is delivered by one system, so after an event civil protection receives one consistent assessment instead of three.',
    '**Transfer tested within the project.** An outsider must install and run the system on the synthetic coast before the release is accepted, so transfer to other coasts is shown, not only promised.',
]

SOURCES = [
    ('EU Delegation to Albania: Release of the Post-Disaster Needs Assessment for Albania\'s recovery (February 2020)', 'https://www.eeas.europa.eu/delegations/albania/release-post-disaster-needs-assessment-albania-recovery_en'),
    ('Government of Albania, EU, UN and World Bank: Albania Post-Disaster Needs Assessment, Volume A (2020)', 'https://eeas.europa.eu/sites/eeas/files/albania_post-disaster_recovery_a.pdf.pdf'),
    ('Mavroulis, S., Lekkas, E. and Carydis, P. (2021). Liquefaction phenomena induced by the 26 November 2019, Mw = 6.4 Durrës (Albania) earthquake and liquefaction susceptibility assessment in the affected area. Geosciences 11(5), 215', 'https://doi.org/10.3390/geosciences11050215'),
    ('INSTAT: Maritime Transport Statistics, Q2 2026', 'https://www.instat.gov.al/en/themes/industry-trade-and-services/transport-accidents-and-characteristics-of-road-vehicles/publikimet/2026/maritime-transport-statistics-q2-2026/'),
    ('NATO: Resilience, civil preparedness and Article 3 (baseline requirements)', 'https://www.nato.int/cps/en/natohq/topics_132722.htm'),
    ('NATO Science for Peace and Security Programme (overview and Key Priorities)', 'https://www.nato.int/cps/en/natohq/85291.htm'),
    ('Richardson, D. S. (2000). Skill and relative economic value of the ECMWF Ensemble Prediction System. Quarterly Journal of the Royal Meteorological Society, 126, 649–667', 'https://www.ecmwf.int/en/elibrary/76153-skill-and-relative-economic-value-ecmwf-ensemble-prediction-system'),
    ('Murphy, A. H. (1977). The value of climatological, categorical and probabilistic forecasts in the cost-loss ratio situation. Monthly Weather Review, 105(7), 803', None),
]


# ------------------------------------------------------------------------------------------------- value model
def v_min(r, s, F, a):
    """Lower bound on the relative economic value when the hit rate is at least 1 - a (cost-loss model)."""
    r = np.asarray(r, dtype=float)
    den = np.minimum(r, s) - r * s
    return 1.0 - (r * F * (1 - s) + a * s * (1 - r)) / den


def positive_range(s, F, a):
    r = np.logspace(-4, 0, 200001)[:-1]
    v = v_min(r, s, F, a)
    ok = r[v > 0]
    return (float(ok.min()), float(ok.max())) if ok.size else (None, None)


def fig_value(path):
    fig, ax = plt.subplots(figsize=(6.7, 3.7), dpi=300)
    r = np.logspace(np.log10(0.004), np.log10(0.9), 600)
    cols = {0.05: '#1F3864', 0.10: '#2E6DB4'}
    for a in ALPHAS:
        v = v_min(r, S_EX, F_EX, a)
        ax.plot(r, np.where(v > -0.5, v, np.nan), color=cols[a], lw=2.2 if a == 0.05 else 1.8,
                ls='-' if a == 0.05 else (0, (5, 2)), label=f'miss rate certified ≤ {a:.0%}   (peak {1 - a - F_EX:.2f})')
        lo, hi = positive_range(S_EX, F_EX, a)
        if a == 0.10:
            ax.axvspan(lo, hi, color='#E6F4EC', zorder=0)
            ax.text(np.sqrt(lo * hi) * 1.9, 0.30, f'warning guaranteed to be worth using\n{lo:.3f} < r < {hi:.2f} (miss ≤ 10%)',
                    ha='center', va='bottom', fontsize=8.8, color='#1E6B47')
        ax.plot(S_EX, 1 - a - F_EX, 'o', ms=6, color=cols[a], zorder=5)
    ax.text(S_EX, 0.95, r'peak at $r = s$:  $V_{\min} = 1 - \alpha_k - F$', fontsize=8.8, color='#1F3864', ha='center',
            va='bottom')
    ax.axhline(0, color='#52606D', lw=0.9)
    ax.axvline(S_EX, color='#9aa3ad', lw=0.9, ls=':')
    ax.set_xscale('log')
    ax.set_xlim(0.004, 0.9); ax.set_ylim(-0.5, 1.1)
    ax.set_xticks([0.005, 0.01, 0.02, 0.05, 0.1, 0.2, 0.5])
    ax.set_xticklabels(['0.005', '0.01', '0.02', '0.05', '0.1', '0.2', '0.5'], fontsize=9)
    ax.tick_params(axis='y', labelsize=9)
    ax.set_xlabel('Cost-loss ratio of the user\'s decision, r = C / L', fontsize=9.5)
    ax.set_ylabel(r'Guaranteed relative value, $V_{\min}$', fontsize=9.5)
    ax.set_yticks([-0.4, -0.2, 0, 0.2, 0.4, 0.6, 0.8, 1.0])
    ax.grid(color='#e3e6ea', lw=0.8); ax.set_axisbelow(True)
    for sp in ('top', 'right'):
        ax.spines[sp].set_visible(False)
    lg = ax.legend(loc='lower center', frameon=True, fontsize=8.8, bbox_to_anchor=(0.55, 0.03), framealpha=1.0)
    lg.get_frame().set_edgecolor('none'); lg.get_frame().set_facecolor('white')
    fig.tight_layout(); fig.savefig(path, dpi=300); plt.close(fig)


# --------------------------------------------------------------------------------------------- impact pathway
COLS = [
    ('Activities', 'what we do · M1–36', '#52606D',
     ['Measure the coast\n(WP1, WP4)', 'Model and certify\n(WP2, WP3)', 'Operate live with\nthe users (WP5)', 'Train and hand\nover (WP6)']),
    ('Outputs', 'what exists · M6–36', '#2E6DB4',
     ['Relative sea level\nat each asset (D1.1)', 'Coupled model and\ncertified surrogate', 'Alarms with certified\nmiss rates (D3.1–D3.2)', 'Open software and\nprocedures (D6.1, D6.2)']),
    ('Outcomes', 'users act · M23 →', '#1E8A5A',
     ['Operators act on\nwarnings of known\nreliability', 'Pumping rule safe\nfor all three hazards', 'One picture for civil\nprotection after events', 'Operators run the\nsystem unaided']),
    ('Impacts', 'security · years 1–5', '#B45309',
     ['Drinking water\nkept fresh', 'Port and energy\nlandfall kept working', 'Faster, consistent\nemergency response', 'Method used on\nother coasts']),
]
ASSUME = ['if data access is\ngranted (D7.1, M3)', 'if gates MS3–MS5\nare passed', 'if operators adopt\nthe system (MS7)']


def fig_pathway(path):
    Wd, Ht = 10.0, 6.3
    fig = plt.figure(figsize=(6.7, 6.7 * Ht / Wd), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, Wd); ax.set_ylim(0, Ht); ax.axis('off')
    bw, gap = 2.14, 0.40
    x0s = [0.08 + i * (bw + gap) for i in range(4)]
    for (title, sub, col, boxes), x0 in zip(COLS, x0s):
        ax.add_patch(FancyBboxPatch((x0, Ht - 0.95), bw, 0.8, boxstyle='round,pad=0,rounding_size=0.08', fc=col, ec='none'))
        ax.text(x0 + bw / 2, Ht - 0.43, title, ha='center', va='center', fontsize=10.5, color='white', fontweight='bold')
        ax.text(x0 + bw / 2, Ht - 0.77, sub, ha='center', va='center', fontsize=8.6, color='white')
        for j, b in enumerate(boxes):
            y = Ht - 1.25 - (j + 1) * 1.02
            ax.add_patch(FancyBboxPatch((x0, y), bw, 0.9, boxstyle='round,pad=0,rounding_size=0.06', fc='white', ec=col, lw=1.3))
            ax.text(x0 + bw / 2, y + 0.45, b, ha='center', va='center', fontsize=8.6, color='#1f2328', linespacing=1.15)
    ymid = Ht - 1.25 - 2.5 * 1.02 + 0.45
    for i in range(3):
        xa = x0s[i] + bw + 0.03; xb = x0s[i + 1] - 0.03
        for yy in (ymid + 1.02, ymid - 1.02):
            ax.add_patch(FancyArrowPatch((xa, yy), (xb, yy), arrowstyle='-|>', mutation_scale=12, color='#9aa3ad', lw=1.4))
        ax.text((xa + xb) / 2, 0.42, ASSUME[i], ha='center', va='center', fontsize=8.6, color='#52606D', style='italic',
                linespacing=1.1, bbox=dict(boxstyle='round,pad=0.25', fc='#F4F6F8', ec='none'))
    fig.savefig(path, dpi=300); plt.close(fig)


# ------------------------------------------------------------------------------------------------- Word helpers
def no_borders(t):
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'nil'); b.append(e)
    t._tbl.tblPr.append(b)


def eq_row(doc, eqs, tex, num):
    t = doc.add_table(rows=1, cols=2); no_borders(t); t.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for j, w in enumerate((15.3, 1.7)):
        t._tbl.tblGrid.findall(qn('w:gridCol'))[j].set(qn('w:w'), str(int(w * 567)))
        t.rows[0].cells[j].width = Cm(w)
    i = len(eqs); eqs.append({'id': i, 'tex': tex, 'display': True})
    c0 = t.rows[0].cells[0].paragraphs[0]; c0.add_run(f'§§EQ{i}§§')
    c1 = t.rows[0].cells[1].paragraphs[0]; c1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(c1, f'({num})', size=10.5)
    from docx.enum.table import WD_ALIGN_VERTICAL
    for c in t.rows[0].cells:
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return t


def head_repeat(t):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t


def hl_cells(t):
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                if '[' not in p.text or not p.runs:
                    continue
                size = p.runs[0].font.size.pt if p.runs[0].font.size else 9.5
                bold = bool(p.runs[0].bold); txt = p.text
                for rn in list(p.runs):
                    rn._element.getparent().remove(rn._element)
                ph(p, txt, size=size, bold=bold)
    return t


def bullets(doc, items, size=10.5, fmt=None):
    for x in items:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); ph(q, x.format(**(fmt or {})), size=size)


def main(out):
    import build_sustainability as BS
    cost = sum(float(c[1].replace(' ', '')) for c in BS.COSTS)
    work = tempfile.mkdtemp()
    f_path = os.path.join(work, 'fig_pathway.png'); fig_pathway(f_path)
    f_val = os.path.join(work, 'fig_value.png'); fig_value(f_val)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.7)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    eqs = []

    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Impact', bold=True, size=18, color='1F3864')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'CERTAIN-COAST will change how three critical assets of an Ally are protected: the wellfields that supply drinking '
            'water, the country\'s main port and an energy landfall, on one subsiding coast. The impact is that the people who '
            'run them act on **warnings whose reliability is known in advance and checked in operation**, and keep doing so '
            'after the project. This section sets out what is at stake, how the results lead to that change, who benefits, how '
            'the impact is measured, what a certified warning is worth to each user, and how the approach reaches other coasts.',
         size=10.5)

    heading(doc, '1. What is at stake', 13.5, before=4)
    bullets(doc, STAKE)

    heading(doc, '2. Impact pathway', 13.5, before=8)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'The project turns activities into outputs that the users can act on, and the users\' actions into security '
            'impacts. Each step depends on a condition that the project controls or monitors (shown at the bottom of Figure 1).',
         size=10.5)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(f_path, width=Cm(16.0))
    p = para(doc, 4)
    run(p, 'Figure 1. Impact pathway: from the work packages to outputs (deliverables), outcomes (what the users do '
           'differently) and impacts. Italic notes: the condition for each step.', italic=True, size=9.5, color='52606D')

    h = heading(doc, '3. Who benefits', 13.5, before=0); h.paragraph_format.page_break_before = True
    head_repeat(table(doc, ['Who', 'What changes for them', 'When'], BENEFIT, [4.0, 9.6, 3.4], size=9.5))

    heading(doc, '4. Impact indicators and targets', 13.5, before=10)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'Each indicator has a baseline, a target at the end of the project and a target three years after it; values in '
            'brackets are proposals to be agreed with the end-users. Indicators after the project are measured from the '
            'bulletin log that the operators keep.', size=10.5)
    rows = [[a, b, c, d.format(cost=cost), e] for a, b, c, d, e in KPI]
    t = hl_cells(head_repeat(table(doc, ['Indicator', 'Baseline', 'End of project (M36)', 'Three years after', 'Verified by'],
                                   rows, [4.4, 2.2, 4.0, 3.6, 2.8], size=9)))
    for r_ in t.rows[1:]:
        shade(r_.cells[0], 'EEF2F8')

    h = heading(doc, '5. What a certified warning is worth to each user', 13.5, before=0)
    h.paragraph_format.page_break_before = True
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'A warning has value only through a decision. In the cost-loss model of forecast value [7, 8], a user can take a '
            'protective action at cost C or risk a loss L if the event occurs unprotected. For an event that occurs in a share '
            's of the warning windows, a warning with hit rate H (warnings among events) and false-alarm rate F (warnings among '
            'non-events) gives an expected expense per unit loss, with r = C/L,', size=10.5)
    eq_row(doc, eqs, r'\frac{E}{L} \;=\; r\,\big[\,H s + F\,(1-s)\,\big] \;+\; (1-H)\,s', 1)
    p = para(doc, 4, before=2); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'and its relative value V compares this expense with acting on the base rate alone (V = 0) and with a perfect '
            'warning (V = 1):', size=10.5)
    eq_row(doc, eqs, r'V \;=\; \frac{\text{min}(r,\,s) \;-\; E/L}{\text{min}(r,\,s) \;-\; r\,s}', 2)
    p = para(doc, 4, before=2); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, '**The new step.** E/L falls as H rises, and CERTAIN-COAST certifies that the miss rate 1 − H of alarm k is at most '
            'αₖ. The value of the warning therefore has a **guaranteed lower bound for every user**, whatever their cost-loss ratio:',
         size=10.5)
    eq_row(doc, eqs, r'V \;\ge\; V_{\text{min}}(r) \;=\; 1 \;-\; \frac{r\,F\,(1-s) \;+\; \alpha_k\, s\,(1-r)}{\text{min}(r,\,s) \;-\; r\,s}', 3)
    lo, hi = positive_range(S_EX, F_EX, 0.10)
    lo5, hi5 = positive_range(S_EX, F_EX, 0.05)
    p = para(doc, 4, before=2); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, f'At r = s the bound is 1 − αₖ − F. For a user who acts on any of the three alarms, αₖ is replaced by the joint '
            f'bound Σαₖ. F and s are measured in back-testing (D3.2) and in live operation (D5.1); r is set with each user '
            f'before live operation (T3.1). **Illustration** (values not yet measured): with s = {S_EX:.2f} and F = {F_EX:.2f}, '
            f'a miss rate certified at 10% guarantees V ≥ {1 - 0.10 - F_EX:.2f} at r = s, and guarantees that the warning is '
            f'worth using (V > 0) for every user with {lo:.3f} < r < {hi:.2f}; at 5% the range widens to {lo5:.3f} < r < {hi5:.2f} '
            f'(Figure 2).', size=10.5)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(f_val, width=Cm(15.5))
    p = para(doc, 4)
    run(p, f'Figure 2. Guaranteed relative value of a warning as a function of the user\'s cost-loss ratio, for miss rates '
           f'certified at 5% and 10% (illustrative s = {S_EX:.2f}, F = {F_EX:.2f}). Shaded: users for whom the warning is '
           f'guaranteed to be worth using at 10%.', italic=True, size=9.5, color='52606D')
    p = para(doc, 4, before=4); p.paragraph_format.keep_with_next = True
    rich(p, '**The decisions behind each warning.** The cost-loss ratio of each user\'s decision is agreed in T3.1 together with '
            'the thresholds; the guaranteed value is then reported in D5.1 for each user.', size=10.5)
    tu = hl_cells(head_repeat(table(doc, ['User', 'Action on a warning', 'Cost C of acting', 'Loss L avoided', 'r = C/L'], USERS,
                                    [3.0, 4.4, 3.6, 4.2, 1.8], size=9)))
    for r_ in tu.rows[:-1]:
        for c in r_.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True

    heading(doc, '6. Contribution to NATO and partner security', 13.5, before=10)
    bullets(doc, SECURITY)

    heading(doc, '7. Beyond Durrës: scale-out and transfer', 13.5, before=8)
    bullets(doc, SCALE)

    heading(doc, '8. Conditions for impact and how they are secured', 13.5, before=8)
    rows = [[a, b, c.format(cost=cost), d] for a, b, c, d in CONDITIONS]
    head_repeat(table(doc, ['Condition', 'What could block it', 'How it is secured', 'Risk'], rows, [3.4, 3.4, 8.6, 1.6], size=9))

    heading(doc, '9. What is new in this impact plan', 13.5, before=10)
    bullets(doc, NEW)

    heading(doc, 'Sources', 11.5, before=8)
    for i, (txt, url) in enumerate(SOURCES, start=1):
        q = para(doc, 2, 0.5); run(q, f'[{i}] {txt}.' + (' ' if url else ''), size=9.5)
        if url:
            link(q, url, url, size=8.5)

    draft = os.path.join(work, 'draft.docx'); doc.save(draft)
    eq_json = os.path.join(work, 'eqs.json')
    with open(eq_json, 'w', encoding='utf-8') as fh:
        json.dump(eqs, fh, ensure_ascii=False)
    omml.main(draft, eq_json, out)
    print('saved', out, '| equations', len(eqs), '| range 10%:', (round(lo, 4), round(hi, 3)), '| 5%:', (round(lo5, 4), round(hi5, 3)),
          '| cost', cost)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Impact.docx')
