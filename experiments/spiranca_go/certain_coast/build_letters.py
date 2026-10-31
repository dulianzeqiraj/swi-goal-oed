"""CERTAIN-COAST (NATO SPS proposal): Letters of commitment (templates). A cover page (how to use, list of letters,
what makes them strong) and seven one-page templates: four end-users (water utility, port authority, energy landfall
operator, civil protection), the partner-country Project Director's institution, and two optional letters (additional
NATO co-director, Albanian research partner). Placeholders are in square brackets and highlighted in yellow; each
organisation adapts the text to its real commitments, prints it on its letterhead and signs it.
Consistent with the Work Plan months, the Partnership (contents of a commitment letter), the Ethics and Data Management
Plan (restricted data, review of publications) and the Sustainability plan (maintenance agreement by Month 33).

usage: python build_letters.py OUT.docx"""
import re
import sys
import os

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import para, heading, table  # noqa: E402
from docx.shared import RGBColor
from docx.oxml.ns import qn

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')
ADDRESS = ['NATO Science for Peace and Security (SPS) Programme', 'NATO Headquarters', 'Boulevard Léopold III',
           '1110 Brussels, Belgium']


def run(p, text, bold=False, italic=False, size=11, color=None, hl=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = 'Arial'; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if hl:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def ph(p, text, size=11, bold=False):
    """Text with [placeholders] highlighted and **bold** segments."""
    for part in re.split(r'(\[[^\]]+\])', text):
        if not part:
            continue
        if part.startswith('['):
            run(p, part, size=size, bold=bold, hl=True)
        else:
            for j, seg in enumerate(part.split('**')):
                if seg:
                    run(p, seg, size=size, bold=bold or j % 2 == 1)


COMMON_END = [
    'We understand that the CERTAIN-COAST bulletins are advisory: operational decisions remain ours, and any warning to the public is issued only by the competent authorities.',
    'We understand that our participation creates no financial obligation; our contribution is in kind. Data we provide will be treated as restricted, used only under a data access agreement for the purposes of the project, and any publication that uses them will be sent to us for review of sensitive content [30] days before submission.',
]

LETTERS = [
    dict(key='utility', title='End-user: water utility of the Durrës wellfields',
         org='[Name of the water utility]', role='end-user (wellfields)',
         intro='[Name of the water utility] supplies drinking water from the wellfields of the Durrës coast. Seawater intrusion, land subsidence and the effects of earthquakes on groundwater are real risks to this supply. We therefore support the application of CERTAIN-COAST to the NATO SPS Programme and commit to take part as an end-user.',
         provide=['access to our production and monitoring wells for the installation and maintenance of loggers, from Month [4] (installation by Month 12);',
                  'our records of groundwater heads, chloride and pumping rates under a data access agreement to be signed by Month 3;',
                  'staff time to agree the chloride threshold and the acceptable false-alarm rate (Months 13–15) and to develop a joint pumping rule (Months 24–33);',
                  'participation of [number] staff in the workshops and in the operator training (Months 24–36).'],
         use='We intend to receive and use the certified intrusion bulletins during the twelve months of live operation (Months 23–34), to issue them ourselves in the final months, and to consider adopting an operating procedure for the system. We are willing to discuss an arrangement for its maintenance after the project.'),
    dict(key='port', title='End-user: Durrës Port Authority',
         org='[Durrës Port Authority]', role='end-user (port)',
         intro='[Durrës Port Authority] operates the main port of Albania on a subsiding coast exposed to earthquakes, liquefaction and extreme sea levels. We support the application of CERTAIN-COAST to the NATO SPS Programme and commit to take part as an end-user.',
         provide=['access to agreed locations in the port for sensors (GNSS, corner reflectors, seismic stations, piezometers), under our port-facility security rules, from Month [4];',
                  'available geotechnical records (CPT and borehole logs) and design levels of the quays under a data access agreement to be signed by Month 3;',
                  'staff time to agree the liquefaction and water-level thresholds (Months 13–15) and to take part in the post-event protocol;',
                  'participation of [number] staff in the workshops and in the operator training (Months 24–36).'],
         use='We intend to receive and use the liquefaction and extreme-water-level bulletins for the port during the live operation (Months 23–34) and to consider adopting an operating procedure for the system. We are willing to discuss an arrangement for its maintenance after the project.'),
    dict(key='energy', title='End-user: operator of the energy landfall',
         org='[Name of the energy landfall operator]', role='end-user (energy landfall)',
         intro='[Name of the energy landfall operator] operates the energy landfall at [location], where supply from the sea meets the national system. Its stability depends on the ground and on the sea level at the coast. We support the application of CERTAIN-COAST to the NATO SPS Programme and commit to take part as an end-user.',
         provide=['access to agreed locations near the landfall for sensors, under our security rules and with our escort, from Month [4];',
                  'the design levels and available geotechnical information needed for the stability checks, under a data access agreement to be signed by Month 3;',
                  'staff time to agree the thresholds for the landfall (Months 13–15);',
                  'participation of [number] staff in the workshops and in the operator training (Months 24–36).'],
         use='We intend to receive and use the stability bulletins for the landfall, combining shaking, liquefaction, water level and subsidence, during the live operation (Months 23–34), and to consider adopting an operating procedure. We are willing to discuss an arrangement for its maintenance after the project.'),
    dict(key='civil', title='End-user: National Civil Protection Agency',
         org='[National Civil Protection Agency]', role='end-user (civil protection)',
         intro='[National Civil Protection Agency] is responsible for preparedness and response to disasters in Albania. The Durrës coast, where drinking water, the main port and an energy landfall are close together on soft and subsiding ground, is a priority for us since the earthquake of 26 November 2019. We support the application of CERTAIN-COAST to the NATO SPS Programme and commit to take part as an end-user.',
         provide=['a representative on the project steering committee, meeting every six months;',
                  'the triggers for the post-event protocol (earthquake, shaking or water level) and participation in its design;',
                  'participation of [number] staff in the workshops, the operator training and the final demonstration.'],
         use='We intend to receive the bulletins and post-event reports during the live operation (Months 23–34) and to consider including the post-event protocol in our procedures. Any warning to the public will remain our responsibility and will be issued through our own channels.'),
    dict(key='ppd', title='Partner-country Project Director\'s institution',
         org='[Partner-country university or research institute]', role='partner-country co-director',
         intro='[Partner-country university or research institute], in [eligible partner country], commits to take part in CERTAIN-COAST as the institution of the partner-country Project Director, [title and name], who is employed by us in [eligible partner country].',
         provide=['leadership of the coupled model and certified surrogate (WP2), the certified multi-hazard alarm (WP3) and training, transfer and dissemination (WP6);',
                  'a transfer test of the method on a subsiding coast of [eligible partner country], using open satellite data and available records;',
                  'the salaries of our staff and the use of our facilities and computing, as an in-kind contribution;',
                  'participation of [number] early-career researchers in the project, including exchanges to the Durrës coast.'],
         use='We accept the terms of the NATO SPS Multi-Year Project Handbook and of the award, will sign the consortium agreement, and will follow the project\'s data management plan, ethics measures and two-tier dissemination rules.',
         custom_end=True),
    dict(key='ncd', title='Optional: additional NATO co-director\'s institution',
         org='[Institution in a NATO member country]', role='additional NATO co-director',
         intro='[Institution in a NATO member country], in [NATO member country], commits to take part in CERTAIN-COAST as the institution of an additional co-director, [title and name].',
         provide=['expertise and supervision for InSAR time-series processing (WP1) and for seismic velocity monitoring with ambient noise (T4.4);',
                  'the salaries of our staff and the use of our facilities, as an in-kind contribution;',
                  'hosting of early-career researchers for training visits.'],
         use='We accept the terms of the NATO SPS Multi-Year Project Handbook and of the award, will sign the consortium agreement, and will follow the project\'s data management plan, ethics measures and dissemination rules.',
         custom_end=True),
    dict(key='arp', title='Optional: Albanian research partner (data and stations)',
         org='[Albanian geoscience or seismological institute]', role='research partner',
         intro='[Albanian geoscience or seismological institute] supports the application of CERTAIN-COAST to the NATO SPS Programme and commits to take part as a research partner.',
         provide=['access to data from our [seismic stations / tide gauge / hydrometeorological stations] on the Durrës coast under a data agreement to be signed by Month 3;',
                  'technical support for the installation and upkeep of the project\'s seismic and GNSS stations;',
                  'participation in the post-event protocol after strong earthquakes or storms.'],
         use='We intend to continue supporting the geodetic and seismic stations after the project, as foreseen in the sustainability plan, and will follow the project\'s data management plan and dissemination rules.',
         custom_end=True),
]

HOWTO = [
    '**Send each template** to the organisation named in it, with the project summary (Abstract) and the relevant section of the Partnership document.',
    '**The organisation adapts the text** to what it can really commit to, replaces every yellow placeholder, prints the letter on its own letterhead, and has it signed (and stamped, where usual) by a person authorised to commit it.',
    '**Do not sign on behalf of an organisation** and do not send unsigned letters as if they were signed; reviewers check that commitments are genuine.',
    '**Collect the signed letters as PDF** and upload them with the application; keep the originals.',
    '**Check the call** for any required format, language (English) or deadline for the letters before sending them out.',
]

STRONG = [
    '**Specific and checkable.** Each letter states what the organisation will provide and when, using the months of the Work Plan (access by Month 3, thresholds in Months 13–15, live operation in Months 23–34), so reviewers can see a real plan rather than general support.',
    '**Security conditions stated up front.** Each end-user letter sets the rules for access and data (restricted tier, review of publications), consistent with the Data Management Plan.',
    '**Responsibilities clear.** The letters say that the bulletins are advisory, that decisions remain with the operators, and that public warnings remain with the authorities.',
    '**Looking beyond the project.** End-users state their intention to use the system, adopt procedures and discuss its maintenance after the project, which supports the Sustainability plan.',
]


def page_break(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)


def letter(doc, L):
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, '[LETTERHEAD OF ' + L['org'].strip('[]').upper() + ']', bold=True, size=10, hl=True)
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; ph(p, '[Place], [date]', size=10.5)
    for i, line in enumerate(ADDRESS):
        p = para(doc, 0 if i < len(ADDRESS) - 1 else 10); run(p, line, size=10.5)
    p = para(doc, 8)
    run(p, 'Subject: ', bold=True, size=10.5)
    run(p, f'Letter of commitment, {L["role"]}, for the NATO SPS Multi-Year Project application CERTAIN-COAST: {FULL}', bold=True, size=10.5)
    p = para(doc, 6); run(p, 'Dear Sir or Madam,', size=10.5)
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; ph(p, L['intro'], size=10.5)
    p = para(doc, 3); run(p, 'In particular, we will provide:', size=10.5)
    for t in L['provide']:
        q = para(doc, 2, 0.6); run(q, '• ', size=10.5); ph(q, t, size=10.5)
    p = para(doc, 6, before=4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; ph(p, L['use'], size=10.5)
    if not L.get('custom_end'):
        for t in COMMON_END:
            p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; ph(p, t, size=10.5)
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ph(p, 'Our contact person for the project is [name, position, e-mail, telephone].', size=10.5)
    p = para(doc, 14); run(p, 'Yours faithfully,', size=10.5)
    for t in ('[Signature and stamp]', '[Name]', '[Position]', L['org']):
        p = para(doc, 0); ph(p, t, size=10.5)


def main(out):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)

    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Letters of Commitment', bold=True, size=18, color='1F3864')
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, 'Templates for the organisations to adapt, print on their letterhead and sign. Placeholders are ', italic=True, size=10, color='52606D')
    run(p, '[highlighted]', italic=True, size=10, hl=True); run(p, '.', italic=True, size=10, color='52606D')
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ph(p, 'The NATO SPS call requires a commitment letter from each end-user of a Multi-Year Project. CERTAIN-COAST has four '
          'end-users (the water utility, the port authority, the energy landfall operator and civil protection) and a '
          'partner-country co-director; two further letters are optional.')

    heading(doc, 'The letters', 13, before=4)
    rows = [[str(i), L['title'], L['org'].strip('[]'), 'Required' if i <= 5 else 'Optional', '[to send]']
            for i, L in enumerate(LETTERS, start=1)]
    t = table(doc, ['', 'Letter', 'Signed by', 'Status', 'Sent / received'], rows, [0.8, 5.6, 5.0, 2.0, 3.2], size=9.5)
    for r in t.rows[1:]:
        for par in r.cells[4].paragraphs:
            for rn in par.runs:
                rn.font.highlight_color = WD_COLOR_INDEX.YELLOW
    para(doc, 4)
    heading(doc, 'How to use these templates', 13, before=6)
    for i, x in enumerate(HOWTO, start=1):
        q = para(doc, 3, 0.5); run(q, f'{i}. ', bold=True); ph(q, x)
    heading(doc, 'What makes these letters strong', 13, before=6)
    for x in STRONG:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); ph(q, x)

    for i, L in enumerate(LETTERS, start=1):
        page_break(doc)
        p = para(doc, 6); run(p, f'[Template {i} of {len(LETTERS)}: {L["title"]}. Delete this line before printing.]', italic=True, size=9, hl=True)
        letter(doc, L)
    doc.save(out)
    print('saved', out, '|', len(LETTERS), 'letters')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Letters_of_Commitment.docx')
