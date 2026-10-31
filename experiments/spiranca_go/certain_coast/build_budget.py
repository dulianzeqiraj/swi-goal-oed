"""CERTAIN-COAST (NATO SPS proposal): Budget. Indicative 36-month budget in the SPS cost categories (equipment, consumables,
training and travel, stipends for young researchers, dissemination, other direct costs, NATO-country Project Director
management allowance), computed from item lines (quantity x unit cost) with year, work-package and recipient splits;
summary tables, a category chart, detailed justification, in-kind contributions, budget control and what is new.
SPS rules used are listed with their sources; points to confirm with the SPS Programme are flagged.

All unit costs are planning estimates to be replaced by quotes. No salaries or overheads are included.

usage: python build_budget.py OUT.docx"""
import os
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade, link  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')

T3 = (1 / 3, 1 / 3, 1 / 3)
# category, item, qty, unit (EUR), work packages, year split, recipient (NPD, PPD or Both = 50/50), justification
LINES = [
    ('Equipment', 'Monitoring boreholes: drilling, casing, screens, sealing', 8, 4000, ['WP4'], (1, 0, 0), 'NPD',
     'Piezometers where the network design (D4.1) ranks them highest; sealed to standard (R12)'),
    ('Equipment', 'Pressure, conductivity and temperature loggers', 12, 1200, ['WP4'], (1, 0, 0), 'NPD',
     'Heads and salinity in new and existing wells: the core data for intrusion and the water table'),
    ('Equipment', 'Barometric loggers', 3, 400, ['WP4'], (1, 0, 0), 'NPD', 'Barometric correction of the pressure loggers'),
    ('Equipment', 'Telemetry units (mobile data loggers)', 12, 600, ['WP4'], (1, 0, 0), 'NPD',
     'Near-real-time data for the live bulletins (WP5)'),
    ('Equipment', 'GNSS stations: receiver, antenna, monument', 3, 9000, ['WP1'], (1, 0, 0), 'NPD',
     'Vertical land motion at the three assets; ties InSAR to a global frame (T1.2)'),
    ('Equipment', 'InSAR corner reflectors', 6, 1500, ['WP1'], (1, 0, 0), 'NPD', 'Stable radar targets on soft ground at the assets (R4)'),
    ('Equipment', 'Seismic stations: three-component sensor and digitiser', 4, 7500, ['WP4'], (1, 0, 0), 'NPD',
     'Continuous ambient noise for seismic velocity at the port and landfall (T4.4)'),
    ('Equipment', 'Seismic node rental, two campaigns', 2, 4000, ['WP4'], (0, 1, 0), 'NPD',
     'Dense temporary coverage to calibrate velocity against the water table'),
    ('Equipment', 'Fibre-optic sensing campaign: rental and processing (conditional)', 1, 15000, ['WP4'], (0, 1, 0), 'NPD',
     'Released only if fibre access is agreed by Month 9 (R9); otherwise moved to seismic nodes'),
    ('Equipment', 'Power and enclosures: solar panels, batteries, locked boxes', 10, 700, ['WP4'], (1, 0, 0), 'NPD',
     'Autonomous and secure operation of the field stations (R11)'),
    ('Equipment', 'Field chemistry instruments: multiparameter meter, sampling pump', 1, 6000, ['WP4'], (1, 0, 0), 'NPD',
     'Conductivity and temperature in the field; sampling for chloride'),
    ('Equipment', 'Spare loggers and sensors', 1, 6000, ['WP4'], (0.5, 0.5, 0), 'NPD',
     'Replacement within two weeks after damage or theft (R11, R21)'),
    ('Equipment', 'Computing server with GPU', 1, 12000, ['WP2'], (1, 0, 0), 'PPD',
     'Thousands of coupled-model runs and training of the certified surrogate (T2.4)'),
    ('Equipment', 'Redundant data server and backup storage', 1, 5000, ['WP5'], (0, 1, 0), 'NPD',
     'Operational redundancy for the live bulletins (R17)'),
    ('Equipment', 'Software licences', 1, 3000, ['WP2'], (1, 0, 0), 'PPD',
     'Processing tools without an open-source equivalent; the project itself releases open software'),
    ('Consumables', 'Laboratory analyses: chloride and major ions', 300, 40, ['WP4'], (0.3, 0.4, 0.3), 'NPD',
     'Calibrates conductivity loggers against chloride and constrains the model (T2.3)'),
    ('Consumables', 'Sampling consumables: bottles, filters, standards', 1, 4000, ['WP4'], (0.4, 0.3, 0.3), 'NPD',
     'Quality-controlled sampling over three years'),
    ('Consumables', 'Field consumables: batteries, cables, fixings', 1, 4000, ['WP4'], (0.4, 0.3, 0.3), 'NPD',
     'Maintenance of the network to at least 90% data availability'),
    ('Consumables', 'Mobile data plans: 15 units over 36 months', 540, 5, ['WP4'], (0.3, 0.35, 0.35), 'NPD',
     'Data transmission for telemetry'),
    ('Training and travel', 'Consortium meetings (kick-off, mid-term, final): travel and subsistence', 18, 1200, ['WP7'], T3, 'Both',
     'Six people from the two countries at three meetings'),
    ('Training and travel', 'Young-researcher exchanges, two months each', 6, 3500, ['WP6'], (0.2, 0.4, 0.4), 'PPD',
     'Partner-country early-career researchers work on site with the data and the operators (T6.1)'),
    ('Training and travel', 'Training school: venue, materials, trainee travel', 1, 15000, ['WP6'], (0, 1, 0), 'NPD',
     'Hands-on school with the coast kit in Month 14 (T6.1)'),
    ('Training and travel', 'Operator and civil-protection training sessions', 1, 5000, ['WP6'], (0, 0.4, 0.6), 'NPD',
     'Staff able to run the system unaided (T6.2)'),
    ('Training and travel', 'Stakeholder workshops: kick-off, mid-term, final, transfer', 4, 3000, ['WP6'], (0.25, 0.25, 0.5), 'NPD',
     'Thresholds, dry run, hand-over and transfer to other coasts (Dissemination Plan)'),
    ('Training and travel', 'Conference presentations', 8, 1800, ['WP6'], (0.25, 0.375, 0.375), 'Both',
     'At least eight presentations of the results (Dissemination Plan)'),
    ('Stipends', 'Stipends for young researchers without a PhD: 4 researchers × 30 months × €400', 120, 400,
     ['WP1', 'WP2', 'WP3', 'WP4', 'WP5'], (0.3, 0.4, 0.3), 'Both',
     'Early-career researchers trained by doing the work of WP1–WP5; rate as stated by the SPS Programme'),
    ('Dissemination', 'Open-access publication fees', 6, 2000, ['WP6'], (0, 1 / 3, 2 / 3), 'Both',
     'At least six open-access papers (Expected Results, Dissemination Plan)'),
    ('Dissemination', 'Website, translation and printing in Albanian and English', 1, 5000, ['WP6'], (0.4, 0.2, 0.4), 'NPD',
     'Public summaries, policy brief and procedures'),
    ('Dissemination', 'Public and press events', 1, 1000, ['WP6'], (0.25, 0.25, 0.5), 'NPD', 'Kick-off, live start and final event'),
    ('Other direct costs', 'Shipping and customs for equipment', 1, 4000, ['WP4'], (1, 0, 0), 'NPD', 'Delivery of sensors to the site'),
    ('Other direct costs', 'Insurance of field equipment', 3, 1000, ['WP4'], T3, 'NPD', 'One year of cover per year of operation'),
    ('Management', 'Management allowance of the NATO-country Project Director', 1, 15000, ['WP7'], T3, 'NPD',
     'Coordination, reporting and financial records; about 4% of the total'),
]
CATS = ['Equipment', 'Consumables', 'Training and travel', 'Stipends', 'Dissemination', 'Other direct costs', 'Management']
WPS = ['WP1', 'WP2', 'WP3', 'WP4', 'WP5', 'WP6', 'WP7']
WPNAMES = {'WP1': 'Coastal motion and relative sea level', 'WP2': 'Coupled model and certified surrogate',
           'WP3': 'Certified multi-hazard alarm', 'WP4': 'Monitoring network', 'WP5': 'Live demonstration',
           'WP6': 'Training, transfer and dissemination', 'WP7': 'Management, data and risk'}

RULES = [
    ('There is no hard budget limit; typical three-year Multi-Year Projects are in the range of €250,000–400,000.', 0),
    ('SPS funds project-specific costs (equipment, consumables, project-related training and travel, including conferences and seminars, and limited stipends for young researchers), not salaries or overheads.', 1),
    ('The NATO-country Project Director may receive a management allowance (up to 10% of the budget in the 2017 handbook).', 2),
    ('Stipends for young researchers without a PhD are given as €400 per month in a 2023 SPS presentation; an older SPS page limited stipends to 15% of the grant.', 3),
    ('Each project has a NATO-country Project Director and a partner-country Project Director; the NATO-country Project Director is accountable for the funds.', 4),
]
SOURCES = [
    ('SPS Multi-Year Project Handbook (2023)', 'https://www.nato.int/content/dam/nato/webready/documents/sps/grants/2023/SPS_MYP_Handbook.pdf'),
    ('SPS Multi-Year Project Handbook (2020)', 'https://nato.int/content/dam/nato/webready/documents/sps/grants/2020/2020-sps-myp-handbook.pdf'),
    ('SPS Multi-Year Project Handbook (2017)', 'https://nato.int/content/dam/nato/webready/documents/sps/grants/2017/1705_SPS_MYP_Handbook.pdf'),
    ('SPS project kick-off presentation (2023)', 'https://www.ftmc.lt/uploads/images/pages/projektai/SPS%20Presentation%20-%20KoM%20-%20Richard%20Brewin.pdf'),
    ('NATO SPS grant mechanisms', 'https://www.nato.int/en/about-us/organization/nato-structure/science-for-peace-and-security-hub/grant-mechanisms'),
    ('NATO SPS page on funding (2012, older rules)', 'https://www.nato.int/cps/en/natohq/88198.htm'),
]
CONFIRM = [
    '**Equipment on the NATO-country side.** All field equipment is installed in Albania, a NATO member, because the critical infrastructure is there. An older SPS page says funds to NATO countries are mainly for travel and consumables and that equipment needs justification. Confirm with the SPS Programme (sps.info@hq.nato.int) before submission; the justification is that the network is site-specific, protects Allied critical infrastructure and is handed over to Albanian operators.',
    '**Partner-country share.** The partner country receives about {ppd_pct}% of the budget in this version. If the Programme expects more, the options are to place the surrogate and alarm work (WP2, WP3) and more stipends at the partner institution, or to add a transfer study on a partner-country coast.',
    '**Eligibility of drilling.** Borehole drilling is a service needed to install the piezometers; confirm that it can be charged under equipment.',
    '**Current rates and ceilings.** Check the stipend rate, the management allowance ceiling and any travel rules in the current handbook.',
    '**Unit costs.** All unit costs are planning estimates; replace them with quotes before submission (Risk Assessment, R19).',
]
INKIND = [
    '**Salaries** of all researchers and project directors, paid by their institutions (not eligible for SPS funding).',
    '**Operators:** access to wells, port and landfall, existing wells and records, staff time for workshops and training.',
    '**Free data:** Copernicus Sentinel-1 radar images, tide-gauge records and existing seismic stations.',
    '**Facilities:** offices, laboratories and existing field equipment of the partner institutions.',
]
CONTROL = [
    '**Spending follows the network ranking.** Sensors are bought in the order of their value for the warnings per euro (D4.1); if prices rise, the lowest-value sensors are dropped first (R19).',
    '**One conditional line.** The fibre-optic campaign (€15,000) is released only if fibre access is agreed by Month 9; otherwise it moves to seismic nodes (R9).',
    '**Front-loaded equipment.** Most equipment is bought in Year 1 so the network streams by Month 12 (MS2).',
    '**Reporting.** Spending against this plan is reported with each progress report; changes between categories follow the SPS rules.',
]
NEW = [
    '**The budget is a design variable.** Each sensor line is funded in the order of its value for the warnings per euro, computed by the goal-oriented network design (Methodology, Section 6), so the budget buys the largest reduction of alarm margins it can.',
    '**A budget cut has a known effect.** Because sensors are ranked, removing the last lines removes the least valuable sensors, and the effect on the miss and false-alarm rates can be stated before the cut.',
    '**Money is released on evidence.** The costly fibre campaign is conditional on access by Month 9, and live operation starts only after the go/no-go milestone (MS4), so funds are not committed to work that cannot happen.',
    '**What is bought stays.** The network, servers and software are handed over to the operators, so the equipment budget becomes a lasting national capability rather than project hardware.',
]


def recipient_split(rec, total):
    return {'NPD': (total, 0), 'PPD': (0, total), 'Both': (total / 2, total / 2)}[rec]


def compute():
    rows = []
    for cat, item, q, u, wps, ys, rec, just in LINES:
        tot = q * u
        rows.append(dict(cat=cat, item=item, q=q, u=u, tot=tot, wps=wps, y=[tot * f for f in ys], rec=rec,
                         npd=recipient_split(rec, tot)[0], ppd=recipient_split(rec, tot)[1], just=just))
    return rows


def eur(x):
    return f'{x:,.0f}'.replace(',', ' ')


def fig_categories(path, rows, total):
    vals = [(c, sum(r['tot'] for r in rows if r['cat'] == c)) for c in CATS]
    vals.sort(key=lambda t: t[1])
    fig, ax = plt.subplots(figsize=(6.7, 3.3), dpi=300)
    y = range(len(vals))
    ax.barh(list(y), [v for _, v in vals], color='#2e6db4', height=0.62)
    for i, (c, v) in enumerate(vals):
        ax.text(v + total * 0.01, i, f'€{eur(v)}  ({100 * v / total:.1f}%)', va='center', ha='left', fontsize=10, color='#1f2328')
    ax.set_yticks(list(y)); ax.set_yticklabels([c for c, _ in vals], fontsize=10.5)
    ax.set_xlim(0, max(v for _, v in vals) * 1.42)
    ax.set_xticks([])
    for s in ('top', 'right', 'bottom'):
        ax.spines[s].set_visible(False)
    ax.spines['left'].set_color('#9aa3ad')
    ax.tick_params(axis='y', length=0)
    fig.tight_layout()
    fig.savefig(path, dpi=300)
    plt.close(fig)


def tidy(t, total_row=False):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    keep = t.rows[:-1] if len(t.rows) <= 8 else (t.rows[0], t.rows[-2])   # short tables stay on one page
    for row in keep:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    if total_row:
        for c in t.rows[-1].cells:
            shade(c, 'E8EDF5')
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True
    return t


def bullets(doc, items):
    for t in items:
        p = para(doc, 3, 0.5); run(p, '• ', bold=True); rich(p, t)


def main(out):
    rows = compute()
    total = sum(r['tot'] for r in rows)
    ytot = [sum(r['y'][k] for r in rows) for k in range(3)]
    npd = sum(r['npd'] for r in rows); ppd = sum(r['ppd'] for r in rows)
    assert abs(sum(ytot) - total) < 1 and abs(npd + ppd - total) < 1
    work = tempfile.mkdtemp(); fig = os.path.join(work, 'fig_budget_categories.png'); fig_categories(fig, rows, total)

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Budget', bold=True, size=18, color='1F3864')
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    eq_share = 100 * sum(r['tot'] for r in rows if r['cat'] == 'Equipment') / total
    lead = (f'CERTAIN-COAST requests **€{eur(total)}** from the NATO SPS Programme for 36 months, within the typical range of '
            f'€250,000–400,000 for a three-year Multi-Year Project. The budget contains only project-specific costs, with no '
            f'salaries or overheads. Equipment is {eq_share:.0f}% of the total, because the project installs a monitoring network '
            f'at critical infrastructure and hands it over to the operators; the rest pays for consumables, training and travel, '
            f'stipends for young researchers, dissemination and management. Every figure is computed from quantity × unit cost; '
            f'unit costs are planning estimates to be replaced by quotes.')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; rich(p, lead, size=11.5)

    heading(doc, 'SPS budget rules used', 13, before=4)
    for t, _ in RULES:
        p = para(doc, 2, 0.5); run(p, '• ', bold=True); rich(p, t, size=10.5)
    p = para(doc, 2, before=3); run(p, 'Sources: ', bold=True, size=9.5, color='52606D')
    for i, (t, u) in enumerate(SOURCES):
        link(p, u, t, 9.5)
        if i < len(SOURCES) - 1:
            run(p, '; ', size=9.5, color='52606D')

    heading(doc, '1. Budget by category and year', 13.5, before=10)
    trows = []
    for c in CATS:
        cr = [r for r in rows if r['cat'] == c]
        ct = sum(r['tot'] for r in cr)
        trows.append([c] + [eur(sum(r['y'][k] for r in cr)) for k in range(3)] + [eur(ct), f'{100 * ct / total:.1f}%'])
    trows.append(['Total'] + [eur(v) for v in ytot] + [eur(total), '100%'])
    tidy(table(doc, ['Category (€)', 'Year 1', 'Year 2', 'Year 3', 'Total', 'Share'], trows, [5.2, 2.4, 2.4, 2.4, 2.6, 2.0], size=10), True)
    para(doc, 4)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(fig, width=Cm(17.0))
    p = para(doc, 6); run(p, 'Figure 1. Budget by category.', italic=True, size=10, color='52606D')

    heading(doc, '2. Budget by work package', 13.5, before=8)
    wrows = []
    for w in WPS:
        v = sum(r['tot'] / len(r['wps']) for r in rows if w in r['wps'])
        wrows.append([w, WPNAMES[w], eur(v), f'{100 * v / total:.1f}%'])
    wrows.append(['', 'Total', eur(total), '100%'])
    tidy(table(doc, ['', 'Work package', 'Total (€)', 'Share'], wrows, [1.4, 10.0, 3.2, 2.4], size=10), True)
    p = para(doc, 6, before=3)
    run(p, 'Stipends are shared equally across WP1–WP5, where the young researchers work. WP2 and WP3 are mostly researcher '
           'time, which the institutions pay in kind, so their SPS share is small.', italic=True, size=9.5, color='52606D')

    heading(doc, '3. Budget by recipient', 13.5, before=8)
    rrows = []
    for c in CATS:
        cr = [r for r in rows if r['cat'] == c]
        rrows.append([c, eur(sum(r['npd'] for r in cr)), eur(sum(r['ppd'] for r in cr))])
    rrows.append(['Total', f'{eur(npd)} ({100 * npd / total:.0f}%)', f'{eur(ppd)} ({100 * ppd / total:.0f}%)'])
    tidy(table(doc, ['Category (€)', 'NATO-country Project Director (Albania)', 'Partner-country Project Director'], rrows,
               [5.2, 6.0, 5.8], size=10), True)
    p = para(doc, 6, before=3)
    run(p, 'The partner country is to be named; lines marked "Both" are split equally. See the points to confirm below.',
        italic=True, size=9.5, color='52606D')

    heading(doc, '4. Points to confirm with the SPS Programme before submission', 13.5, before=8)
    bullets(doc, [t.format(ppd_pct=f'{100 * ppd / total:.0f}') for t in CONFIRM])

    heading(doc, '5. Detailed budget and justification', 13.5, before=10)
    for c in CATS:
        cr = [r for r in rows if r['cat'] == c]
        heading(doc, f'{c}  (€{eur(sum(r["tot"] for r in cr))})', 11.5, color='2E6DB4', before=6)
        drows = [[r['item'], f'{r["q"]:g}', eur(r['u']), eur(r['tot']), ', '.join(r['wps']) if len(r['wps']) < 3 else 'WP1–WP5',
                  {'NPD': 'NATO', 'PPD': 'Partner', 'Both': 'Both'}[r['rec']], r['just']] for r in cr]
        drows.append(['Subtotal', '', '', eur(sum(r['tot'] for r in cr)), '', '', ''])
        tidy(table(doc, ['Item', 'Qty', 'Unit (€)', 'Total (€)', 'WP', 'To', 'Justification'], drows,
                   [4.0, 1.0, 1.5, 1.8, 1.4, 1.6, 5.7], size=9), True)

    heading(doc, '6. In-kind contributions (not requested from SPS)', 13.5, before=10)
    bullets(doc, INKIND)
    heading(doc, '7. Budget control', 13.5, before=8)
    bullets(doc, CONTROL)
    heading(doc, '8. What is new in this budget', 13.5, before=8)
    bullets(doc, NEW)
    doc.save(out)
    print('saved', out, '| total', total, '| years', [round(v) for v in ytot], '| NPD', npd, 'PPD', ppd,
          '| equipment %.1f%%' % eq_share)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Budget.docx')
