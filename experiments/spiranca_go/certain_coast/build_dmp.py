"""CERTAIN-COAST (NATO SPS proposal): Data Management Plan. Follows the structure of the Horizon Europe DMP template
(data summary; FAIR: findable, accessible, interoperable, re-usable; other research outputs; allocation of resources;
data security; ethics; other issues), with a data inventory of 17 datasets, a data-flow figure (two tiers), versions
of the plan, what is new and a short summary for the application form. Consistent with T7.2 (data management and
access agreements, Month 3), the two-tier Dissemination Plan, the Ethics section (Law No. 124/2024), the Budget
(servers) and D6.1 (open software and data, Month 33).

usage: python build_dmp.py OUT.docx"""
import os
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')

LEAD = ('CERTAIN-COAST produces and uses satellite, geodetic, hydrogeological, chemical, seismic and model data about '
        'critical infrastructure, together with operator records and a small amount of personal data. Its data policy is '
        '"as open as possible, as closed as necessary": methods, code, regional products, a synthetic test coast and the '
        'verification statistics of the warnings are open and FAIR; operator records and asset-level results are restricted '
        'but still findable through their metadata; personal data are protected and deleted when no longer needed. The '
        'plan is delivered with the access agreements in Month 3 (T7.2) and updated in Month 18 and Month 36.')

# id, dataset, origin and type, format, volume, tier, owner
INVENTORY = [
    ['DS1', 'Sentinel-1 radar images', 'External, Copernicus (free)', 'SAFE / SLC', '~5 TB working copy', 'Not redistributed', 'Copernicus'],
    ['DS2', 'InSAR products: line-of-sight series, vertical and east velocities', 'Processed (WP1)', 'GeoTIFF, NetCDF, CSV', '~50 GB', 'Open (regional); restricted (asset points)', 'Project'],
    ['DS3', 'GNSS observations and daily time series', 'New, project stations (WP1)', 'RINEX, CSV', '~10 GB', 'Restricted inside sites; open outside', 'Project, operators'],
    ['DS4', 'Tide-gauge sea level', 'External, existing gauge', 'CSV', '<1 GB', 'Provider\'s terms', 'Provider'],
    ['DS5', 'Heads, conductivity and temperature from loggers', 'New (WP4)', 'CSV', '~2 GB', 'Restricted; open after consent with generalised locations', 'Operators'],
    ['DS6', 'Chloride and major-ion analyses', 'New (WP4)', 'CSV', '<1 GB', 'Restricted; open after consent', 'Operators'],
    ['DS7', 'Historical heads, chloride and pumping records', 'Existing, operators', 'CSV, scanned records', '~5 GB', 'Restricted', 'Operators'],
    ['DS8', 'Geotechnical data (CPT, boreholes) at port and landfall', 'Existing, operators', 'PDF, AGS, CSV', '~2 GB', 'Restricted', 'Operators'],
    ['DS9', 'Continuous seismic waveforms and station metadata', 'New (WP4)', 'miniSEED, StationXML', '~0.2 TB', 'Restricted inside sites; open outside', 'Project'],
    ['DS10', 'Seismic node and fibre-optic campaigns', 'New (WP4)', 'miniSEED, HDF5', '1–10 TB', 'Restricted raw; derived open', 'Project'],
    ['DS11', 'Seismic velocity change (dv/v) series', 'Processed (T4.4)', 'CSV, NetCDF', '<1 GB', 'Open (aggregated)', 'Project'],
    ['DS12', 'Coupled-model set-up, calibrated parameters and runs', 'Generated (WP2)', 'Model files, NetCDF', '~3 TB', 'Restricted (site); open (synthetic coast)', 'Project'],
    ['DS13', 'Surrogate training sets and trained surrogates', 'Generated (WP2)', 'HDF5, ONNX', '~1 TB', 'Restricted (site); open (synthetic coast)', 'Project'],
    ['DS14', 'Bulletins and outcome log', 'Generated (WP5)', 'JSON, CSV', '<1 GB', 'Restricted; verification statistics open', 'Operators, project'],
    ['DS15', 'Back-testing protocol and results', 'Generated (WP3)', 'PDF, CSV', '<1 GB', 'Open (protocol frozen before results)', 'Project'],
    ['DS16', 'Participant data (workshops, training, interviews)', 'New, personal data', 'Encrypted spreadsheets, notes', '<1 GB', 'Personal: not shared', 'Partners'],
    ['DS17', 'Software: model, surrogate, alarms, network design', 'Generated (all WPs)', 'Source code, containers', '<10 GB', 'Open (D6.1)', 'Project'],
]

FINDABLE = [
    '**Persistent identifiers.** Every open dataset and every software release receives a DOI through Zenodo; restricted datasets receive a DOI for their metadata record, so they are findable even when the data are not open.',
    '**Rich metadata.** DataCite metadata for all records; ISO 19115-style metadata for spatial data; StationXML for seismic stations; each record names the project, the work package, the method and the software version.',
    '**Naming and versions.** Files follow the pattern CC_<WP>_<dataset>_<version>_<date>; versions are numbered and the change log is kept with the data.',
    '**Keywords** from the project vocabulary (seawater intrusion, liquefaction, relative sea level, InSAR, early warning, conformal risk control) to support search.',
]
ACCESSIBLE = [
    '**Open data** in Zenodo, free of charge and without registration; seismic data from stations outside the sites through an FDSN-compliant data centre, subject to agreement with that centre.',
    '**Restricted data** on the project servers; access on request through the steering committee with the owner\'s consent, under a data access agreement; metadata always public.',
    '**Embargo.** Open data are released at the latest when the related paper is published, or by Month 33 (D6.1), whichever comes first.',
    '**Release gate.** Every release is classified and checked before publication (Dissemination Plan): sensor locations inside sites are generalised and asset-level results are aggregated or anonymised.',
]
INTEROP = [
    '**Standard formats:** RINEX for GNSS, miniSEED and StationXML for seismic data, GeoTIFF and NetCDF following the CF conventions for gridded data, CSV with a data dictionary for tables, HDF5 for large arrays.',
    '**Reference systems stated for every dataset:** ETRS89 / UTM zone 34N (EPSG:25834) for coordinates; ellipsoidal heights from GNSS with the transformation to the local tide-gauge datum documented.',
    '**Units and variables** follow SI units and CF standard names where they exist; a project glossary defines the rest.',
]
REUSE = [
    '**Licences:** CC BY 4.0 for open data, Apache-2.0 for software, CC0 for metadata; third-party data keep their own terms.',
    '**Documentation:** a README and a data dictionary with every dataset; the processing chain and software version recorded as provenance.',
    '**Quality:** quality flags on every time series; calibration and maintenance logs for every sensor.',
    '**A synthetic test coast:** an open, synthetic coast with the same structure as the real one, so that every method can be reproduced without access to sensitive data.',
    '**Preservation:** open records in Zenodo for long-term preservation; restricted data kept by the owners and the partners for at least ten years after the project.',
]
OUTPUTS = [
    '**Software** in a public Git repository with tests, documentation and containers; each release archived with a DOI through the Zenodo integration (D6.1).',
    '**Models and workflows:** model configurations, calibration scripts and notebooks published with the synthetic coast; site-specific versions restricted.',
    '**Trained surrogates** exported in an open format (ONNX) for the synthetic coast.',
    '**Bulletin templates and procedures:** generic templates open; operator-specific procedures restricted (D6.2).',
]
RESOURCES = [
    ['Data storage and backup', 'Computing server and redundant data server (Budget: €12,000 and €5,000)', 'WP2, WP5'],
    ['Open repository', 'Zenodo, free of charge', '—'],
    ['Data manager', 'Named at the NATO-country institution (in kind)', 'WP7, T7.2'],
    ['Dataset owners', 'Each work-package leader for the data of the work package', 'WP1–WP6'],
    ['Long-term preservation', 'Zenodo for open data; partner institutions for restricted data', 'After Month 36'],
]
SECURITY = [
    '**No classified information.** The project uses and produces no NATO-classified or nationally classified information; restricted data are sensitive but unclassified.',
    '**Two servers in different locations** with daily backup, following the 3-2-1 rule (three copies, two media, one off site) (Risk Assessment, R17).',
    '**Encryption** at rest and in transit (VPN for telemetry; SFTP and HTTPS for transfers); restricted data never stored on personal cloud accounts.',
    '**Access control** by role, with personal accounts, two-factor authentication and access logs.',
    '**Integrity:** checksums (SHA-256) for every file; the bulletin log is append-only, so a warning cannot be changed after it is issued.',
    '**Incidents** reported to the data manager at once, and to the owners and, for personal data, to the competent authority as the law requires.',
]
ETHICS = [
    '**Personal data (DS16)** are handled under Albania\'s Law No. 124/2024, the GDPR for any EU partner and the partner country\'s law, with consent, minimisation, pseudonymisation and deletion at the end of the retention period (Ethics, Section 4).',
    '**Operator data** are used only under the access agreements of T7.2 and for the purposes stated in them.',
    '**Security-sensitive results** follow the two-tier dissemination and the release gate.',
]
OTHER = [
    '**Copernicus attribution** ("contains modified Copernicus Sentinel data") on every product derived from Sentinel-1.',
    '**Third-party data** (tide gauge, national networks) used under their providers\' terms.',
    '**National rules on geospatial data** checked with the competent Albanian authority before publishing maps.',
    '**Ownership and access** of all results set out in the consortium agreement.',
]
VERSIONS = [
    ['1.0', 'Month 3', 'With the access agreements (T7.2); data inventory, rules and responsibilities'],
    ['2.0', 'Month 18', 'Update after the network is installed and the model calibrated; repositories confirmed'],
    ['3.0', 'Month 36', 'Final version with what was released, where, and how it is preserved'],
]
NEW = [
    '**Replayable warnings.** Every bulletin is stored with its inputs, model and code version in an append-only log, so any warning can be reproduced exactly, years later.',
    '**A test registered before the results.** The back-testing protocol is published with a timestamp and a DOI before the results are computed, so the central claim of the project can be checked by anyone.',
    '**Findable even when closed.** Restricted datasets have public metadata and a DOI, so others know they exist and how to request access.',
    '**FAIR with a security gate.** The release gate generalises sensor locations and aggregates asset-level results, so the project can be open without exposing critical infrastructure.',
    '**A synthetic twin of the coast.** An open synthetic coast with the same structure lets anyone reproduce the methods without sensitive data.',
]
SHORT = ('CERTAIN-COAST follows the principle "as open as possible, as closed as necessary". Its 17 datasets (satellite, GNSS, '
         'groundwater, chemistry, seismic, model, bulletin and participant data, and software) are listed with format, '
         'volume, owner and access tier. Open data, a synthetic test coast, verification statistics and software are '
         'published in Zenodo with DOIs, standard formats (RINEX, miniSEED, NetCDF-CF, GeoTIFF), stated reference systems and '
         'open licences (CC BY 4.0, Apache-2.0). Operator records and asset-level results are restricted but findable through '
         'public metadata, and every release passes a security gate. No classified information is used. Data are kept on two '
         'servers with encryption, access control, checksums and an append-only bulletin log that makes every warning '
         'reproducible. Personal data follow Albania\'s Law No. 124/2024 and the GDPR. A data manager and the work-package '
         'leaders are responsible; the plan is delivered in Month 3 and updated in Months 18 and 36.')


def fig_flow(path):
    W, H = 6.7, 4.1
    fig = plt.figure(figsize=(W, H), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(0, H); ax.axis('off')

    def box(x, y, w, h, fc, ec, title, body, tc='#1f2328', ts=9.8, bs=8.8):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0,rounding_size=0.07', fc=fc, ec=ec, lw=1.3))
        ax.text(x + w / 2, y + h - 0.19, title, fontsize=ts, fontweight='bold', color=tc, ha='center', va='center')
        if body:
            ax.text(x + w / 2, y + (h - 0.30) / 2, body, fontsize=bs, color='#1f2328', ha='center', va='center', linespacing=1.25)

    def arrow(x0, y0, x1, y1, col='#1f2328'):
        ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle='-|>', mutation_scale=12, color=col, lw=1.3))

    srcs = [('Satellite', 'Sentinel-1 (DS1–2)'), ('Field sensors', 'GNSS, loggers,\nseismic (DS3–11)'),
            ('Operators', 'records, geotechnics\n(DS7–8)'), ('Models', 'runs, surrogates\n(DS12–13)')]
    for i, (t, b) in enumerate(srcs):
        y = H - 0.15 - (i + 1) * 0.88 + 0.06
        box(0.05, y, 1.62, 0.78, '#eef2f8', '#2e6db4', t, b)
        arrow(1.68, y + 0.39, 2.20, 2.05)
    box(2.22, 1.25, 1.95, 1.62, '#1F3864', '#1F3864', 'Project data platform', '', tc='white')
    ax.text(3.195, 1.85, 'two servers, backup,\nencryption, access\ncontrol, checksums,\nappend-only bulletin log',
            fontsize=8.6, color='white', ha='center', va='center', linespacing=1.25)
    ax.add_patch(FancyBboxPatch((2.22, 0.12), 1.95, 0.85, boxstyle='round,pad=0,rounding_size=0.07', fc='#fdecea', ec='#b03a3a', lw=1.3))
    ax.text(3.195, 0.73, 'Personal data (DS16)', fontsize=9.6, fontweight='bold', color='#b03a3a', ha='center', va='center')
    ax.text(3.195, 0.40, 'separate, encrypted;\ndeleted after retention', fontsize=8.6, color='#1f2328', ha='center', va='center', linespacing=1.2)

    ax.add_patch(FancyBboxPatch((4.30, 2.70), 0.56, 0.52, boxstyle='round,pad=0,rounding_size=0.06', fc='#fff1c2', ec='#c27c12', lw=1.2))
    ax.text(4.58, 2.96, 'release\ngate', fontsize=8.6, fontweight='bold', color='#1f2328', ha='center', va='center', linespacing=1.1)
    arrow(4.18, 2.55, 4.36, 2.70)
    box(5.12, 2.40, 1.53, 1.15, '#e6f4ec', '#1e8a5a', 'Open tier', 'Zenodo DOIs; FDSN\nnode; software;\nsynthetic coast;\nverification statistics')
    arrow(4.87, 2.96, 5.11, 2.96, '#1e8a5a')
    box(5.12, 0.55, 1.53, 1.15, '#f6efe6', '#9c6b30', 'Restricted tier', 'operators and\nauthorities; access\nby request; metadata\npublic with DOI')
    arrow(4.18, 1.45, 5.11, 1.15, '#9c6b30')
    fig.savefig(path, dpi=300)
    plt.close(fig)


def tidy(t):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    keep = t.rows[:-1] if len(t.rows) <= 6 else (t.rows[0], t.rows[1], t.rows[-2])
    for row in keep:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t


def bullets(doc, items, size=11):
    for t in items:
        p = para(doc, 3, 0.5); run(p, '• ', bold=True, size=size); rich(p, t, size=size)


def sub(doc, text):
    p = para(doc, 2, before=4); run(p, text, bold=True, size=11.5, color='2E6DB4'); p.paragraph_format.keep_with_next = True


def main(out):
    n_short = len(SHORT.split())
    work = tempfile.mkdtemp(); fig = os.path.join(work, 'fig_dmp_flow.png'); fig_flow(fig)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Data Management Plan', bold=True, size=18, color='1F3864')
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; rich(p, LEAD, size=11.5)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(fig, width=Cm(17.0))
    p = para(doc, 6); run(p, 'Figure 1. Data flow: sources, the project data platform, and the open and restricted tiers behind the release gate.',
                          italic=True, size=10, color='52606D')

    heading(doc, '1. Data summary', 13.5, before=6)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'The data serve one purpose: to measure, model and warn of seawater intrusion, liquefaction and extreme sea levels '
            'at the three assets, and to verify the warnings. The project re-uses free satellite data and existing operator '
            'records, and generates new field, model and bulletin data (Table 1, on the next page). The total volume is of the '
            'order of 20 TB, most of it satellite and fibre-optic data.')

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(1.5)
    p = para(doc, 4); run(p, 'Table 1. Data inventory.', bold=True, size=12, color='1F3864'); p.paragraph_format.keep_with_next = True
    t = tidy(table(doc, ['ID', 'Dataset', 'Origin and type', 'Format', 'Volume', 'Access tier', 'Owner'], INVENTORY,
                   [1.2, 6.4, 4.2, 3.4, 2.6, 5.6, 3.3], size=9))
    tier_col = {'Open': 'E6F4EC', 'Restricted': 'F6EFE6', 'Personal': 'FDECEA', 'Not': 'EEF2F8', 'Provider': 'EEF2F8'}
    for i, r in enumerate(INVENTORY, start=1):
        for k, v in tier_col.items():
            if r[5].startswith(k):
                shade(t.rows[i].cells[5], v)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(1.8)
    heading(doc, '2. FAIR data', 13.5, before=0)
    sub(doc, '2.1 Findable'); bullets(doc, FINDABLE)
    sub(doc, '2.2 Accessible'); bullets(doc, ACCESSIBLE)
    sub(doc, '2.3 Interoperable'); bullets(doc, INTEROP)
    sub(doc, '2.4 Re-usable'); bullets(doc, REUSE)
    heading(doc, '3. Other research outputs', 13.5, before=8); bullets(doc, OUTPUTS)
    heading(doc, '4. Allocation of resources', 13.5, before=8)
    tidy(table(doc, ['Need', 'Resource', 'Where'], RESOURCES, [4.2, 9.6, 3.2], size=10)); para(doc, 4)
    heading(doc, '5. Data security', 13.5, before=8); bullets(doc, SECURITY)
    heading(doc, '6. Ethics', 13.5, before=8); bullets(doc, ETHICS)
    heading(doc, '7. Other issues', 13.5, before=8); bullets(doc, OTHER)
    heading(doc, '8. Versions of this plan', 13.5, before=8)
    tidy(table(doc, ['Version', 'When', 'Content'], VERSIONS, [2.0, 2.6, 12.4], size=10)); para(doc, 4)
    heading(doc, '9. What is new in this data plan', 13.5, before=8); bullets(doc, NEW)
    heading(doc, f'10. Short summary for the application form ({n_short} words)', 13.5, before=8)
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; rich(p, SHORT)
    doc.save(out)
    print('saved', out, '| datasets', len(INVENTORY), '| short', n_short, 'words')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Data_Management_Plan.docx')
