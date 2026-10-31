"""CERTAIN-COAST (NATO SPS proposal): Risk assessment. Scoring method (likelihood x impact, 1-5), a risk matrix before and
after mitigation (figure), a register of 21 risks in seven categories with early-warning indicators and numeric
triggers, mitigation, contingency, owner and residual score, the risks the warning service itself could create,
monitoring and escalation (T7.3), and what is new. The seven risks of the Work Plan are included with the same ratings
(Low = 2, Medium = 3, High = 4).

usage: python build_risk.py OUT.docx"""
import os
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')

BANDS = [(4, 'Low', 'DCEFE3', '#DCEFE3'), (9, 'Medium', 'FFF1C2', '#FFF1C2'), (14, 'High', 'F9D3B4', '#F9D3B4'),
         (25, 'Very high', 'F2B8B5', '#F2B8B5')]


def band(score):
    for top, name, hx, rgb in BANDS:
        if score <= top:
            return name, hx, rgb


# id, risk and consequence, L, I, indicator (trigger), mitigation, contingency, owner, residual L, residual I
CATS = [
    ('A. Data and access', [
        ('R1', 'Operators delay access to well, port or landfall data; calibration and design start late',
         3, 4, 'Access agreements not signed by Month 3',
         'Letters of support at proposal stage; access agreements by Month 3 (T7.2, D7.1); data needs listed per operator',
         'Start WP1 and the design in WP4 on public and satellite data; escalate to the steering committee', 'WP7', 2, 3),
        ('R2', 'Security rules restrict access to the port or landfall (port-facility security, energy operator rules)',
         3, 3, 'Site access not granted by Month 6',
         'Early meetings with security officers; named staff, escorts and inductions; installation plans approved in advance',
         'Install on the nearest accessible ground; rely more on InSAR, corner reflectors and seismic noise', 'WP4', 2, 2),
        ('R3', 'Historical head and chloride records are incomplete or of uneven quality',
         4, 3, 'More than 30% of records fail quality control',
         'Data rescue from paper archives; quality flags; record errors carried as uncertainty in the assimilation',
         'Calibrate on the new network and InSAR; widen the stated uncertainty', 'WP2', 3, 2),
    ]),
    ('B. Science and methods', [
        ('R4', 'InSAR loses coherence on vegetated plains or on new port works',
         3, 3, 'Too few coherent points at an asset after T1.1 (Month 9)',
         'Corner reflectors and GNSS at the assets (T1.2); persistent scatterers on structures; both orbit directions',
         'Use GNSS and levelling at the asset; interpolate with stated uncertainty', 'WP1', 2, 2),
        ('R5', 'The coupled model fails its validation test on held-out data',
         3, 4, 'Coverage of the 90% intervals below 80% at Month 18 (MS3)',
         'Build in stages (flow, then salt, then compaction and liquefaction); test each against data before coupling',
         'Two-dimensional sections through each asset; larger error bound Δₖ; report the limitation', 'WP2', 2, 3),
        ('R6', 'The surrogate fails its error bound in part of the domain',
         2, 3, 'Δₖ larger than half the margin for more than 5% of forecasts',
         'Training runs concentrated where the bound is large; adjoint-weighted error indicator flags weak regions',
         'Full model on a reduced domain for flagged forecasts; margin widens automatically', 'WP2', 1, 2),
        ('R7', 'Certified alarms are too cautious: the false-alarm rate is unacceptable to the operators',
         3, 4, 'Back-tested false-alarm rate above the level agreed in T3.1',
         'Acceptable false-alarm rate agreed with each operator at T3.1; tolerances allocated by consequence; network designed to shrink margins',
         'Longer lead-time windows; staged alerts (watch, then warning); report the trade-off to the operators', 'WP3', 2, 3),
        ('R8', 'Conditions move beyond the observed record (drought, compound storm, post-seismic change) and weaken the guarantees',
         3, 4, 'Growing share of issue times outside the observed record',
         'Extreme-aware margin term; storylines beyond the record in back-testing (T3.4); nonstationary extremes with land motion',
         'Bulletins flag "outside record"; margins widen; result reported to the operators and civil protection', 'WP3', 3, 2),
        ('R9', 'Fibre-optic access is not granted, or seismic velocity responds too weakly to the water table',
         3, 2, 'No fibre agreement by Month 9; no significant relation with piezometers by Month 20',
         'Seismic stations and temporary nodes give the same velocity signal at lower resolution',
         'Water table from piezometers alone; seismic velocity kept for post-event damage', 'WP4', 2, 1),
    ]),
    ('C. Field work and equipment', [
        ('R10', 'Procurement or customs delays for sensors and loggers',
         3, 3, 'Delivery more than 8 weeks after the order',
         'Orders placed in Month 4; installation window to Month 12; sensors ranked highest in D4.1 installed first',
         'Install in order of value; borrow partner equipment for the first months', 'WP4', 2, 2),
        ('R11', 'Theft, vandalism or damage of field equipment',
         3, 2, 'Telemetry gap or tamper alert at a station',
         'Sites inside operator premises where possible; locked enclosures; telemetry alerts; spare loggers',
         'Replace from spares within two weeks; fill the gap from neighbouring sensors', 'WP4', 2, 1),
        ('R12', 'A new borehole connects fresh and saline layers and lets seawater move between them',
         2, 4, 'Change in conductivity or vertical flow in the borehole after installation',
         'Licensed driller; screens and seals designed by the project hydrogeologist; grouting checked; conductivity logged before and after',
         'Seal or decommission the borehole to standard; report to the utility', 'WP4', 1, 3),
        ('R13', 'Injury during field work (port traffic, storms, inspections after an earthquake)',
         2, 5, 'Near-miss report; active storm or post-earthquake warning',
         'Safety plan and operator inductions; two-person rule; no field work during active warnings; post-event visits only after clearance',
         'Stop field work; incident review; revise the safety plan', 'WP7', 1, 4),
    ]),
    ('D. The warning service in operation', [
        ('R14', 'No strong earthquake or surge occurs during the demonstration, so the guarantees are not tested on such events',
         4, 3, 'No triggering event by Month 30',
         'Back-testing over the historical record, including the 2019 earthquake, and on storylines beyond it (T3.4)',
         'Report event-based verification as back-testing only; continue the bulletin log after hand-over', 'WP3', 4, 2),
        ('R15', 'A missed alarm coincides with damage; trust and responsibilities are questioned',
         2, 5, 'Any missed exceedance during live operation',
         'Every bulletin states its certified miss rate; decisions stay with the operators; roles written into the bulletin protocol (T5.1)',
         'Review within one week; margins recalibrated; findings shared with the operators and civil protection', 'WP5', 2, 3),
        ('R16', 'Over-reliance: operators drop their own checks and rely on the bulletins alone',
         2, 4, 'Decisions recorded without the operator\'s own measurements',
         'System advisory alongside existing procedures; training on what the miss rate means (T6.2)',
         'Joint review with the operator; procedures restated in D6.2', 'WP6', 1, 3),
        ('R17', 'IT or cyber failure interrupts data streams or bulletins',
         3, 3, 'Data latency above the cadence of a hazard; failed backup',
         'Redundant server and storage; offline copy of the surrogate at each operator; access control and updates',
         'Manual bulletin from the last valid forecast with a wider margin; restore from backup', 'WP5', 2, 2),
    ]),
    ('E. People, partners and budget', [
        ('R18', 'Loss of key researchers or of trained operator staff',
         3, 3, 'Fewer than two trained people per operator; departure of a WP lead',
         'Several staff trained per operator; documented procedures and software; early-career researchers paired with seniors',
         'Reassign tasks within the consortium; extra training session', 'WP7', 2, 2),
        ('R19', 'Equipment prices rise above the budget lines',
         3, 3, 'Quotes more than 10% above the budget line',
         'Quotes at proposal stage; prices fixed at order; network design ranked by value per cost',
         'Drop the lowest-value sensors first (D4.1 ranking); report to the steering committee', 'WP7', 2, 2),
    ]),
    ('F. Security of information', [
        ('R20', 'Sensitive information on critical-infrastructure vulnerability is disclosed through open data or publications',
         3, 5, 'A planned release contains asset-level vulnerability results',
         'Data classification in the data plan (T7.2); asset-level results shared only with operators and authorities; open release of methods, code and non-sensitive aggregated data; operators review publications before submission',
         'Withdraw or redact the release; notify the operators; review the data plan', 'WP7', 1, 4),
    ]),
    ('G. External events', [
        ('R21', 'A strong earthquake or storm damages the network or disrupts the project',
         2, 4, 'Event above the post-event triggers',
         'Spare sensors; robust mounting; data stored off site; the post-event protocol (T5.4) turns the event into data',
         'Re-plan the affected tasks; repair from spares; report to the SPS Programme', 'WP5', 2, 2),
    ]),
]

LIKELIHOOD = [
    ['1', 'Rare', 'Less than 10% over the 36 months'],
    ['2', 'Unlikely', '10–30%'],
    ['3', 'Possible', '30–50%'],
    ['4', 'Likely', '50–80%'],
    ['5', 'Almost certain', 'More than 80%'],
]
IMPACT = [
    ['1', 'Minimal', 'No effect on deliverables'],
    ['2', 'Minor', 'Delay of less than one month in one task'],
    ['3', 'Moderate', 'Delay of one to three months, or one deliverable reduced in scope'],
    ['4', 'Major', 'A milestone missed or an objective partly not met'],
    ['5', 'Severe', 'An objective not met, harm to people, or a security breach'],
]

SELF = [
    '**False alarms that wear out trust.** Every guarantee on misses is paid for in false alarms. The acceptable false-alarm rate is agreed with each operator before the alarms are designed (T3.1), so the cost of caution is a decision, not a surprise (R7).',
    '**A miss that matters.** A certified miss rate of 5–10% means misses will happen. Each bulletin says so, responsibilities are written down, and every miss is reviewed and used to recalibrate (R15).',
    '**Over-reliance.** A trusted system can replace the checks it was meant to support. The bulletins are advisory, and training explains what a miss rate means in practice (R16).',
    '**Disclosure.** A map of where the port or the landfall would liquefy is useful to an operator and to an adversary. Methods and code are open; asset-level vulnerability results are not (R20).',
    '**New pathways for seawater.** A badly sealed monitoring borehole can let seawater move between layers. Drilling and sealing follow a written design and are checked by logging conductivity before and after (R12).',
]

MONITOR = [
    '**Owner and indicator.** Each risk has an owner (the leader of the work package named in the register) and an early-warning indicator with a numeric trigger.',
    '**Monthly review.** The coordinator and the work-package leaders review the indicators every month and update the scores (T7.3).',
    '**Escalation.** A risk is escalated to the steering committee when its indicator crosses its trigger or its score reaches 10 or more; the committee decides on the contingency within two weeks.',
    '**Milestone gates.** MS3 (Month 18, model calibrated) and MS4 (Month 22, go or no-go for live operation) are formal checkpoints: live operation starts only if the surrogate meets its error bound and the alarm prototype passes a dry run.',
    '**Reporting.** The register, with changes since the last report, is part of every progress report to the SPS Programme.',
]

NEW = [
    '**The main operational risk has a measured bound.** The risk that a hazard is missed is not only listed: it is bounded by the certified miss rates and checked against the bulletin log (Methodology, Section 5).',
    '**Risks the service itself creates are in the register.** False-alarm fatigue, over-reliance, disclosure and new seawater pathways are treated as project risks, with owners and triggers, not left to after the project.',
    '**Numeric triggers, watched like the hazards.** Each risk has an indicator and a threshold, reviewed monthly, so escalation is a rule rather than a judgement call.',
    '**The miss–false-alarm trade-off is agreed in advance.** The operators set the acceptable false-alarm rate before the alarms are built, so the largest scientific risk (R7) becomes a negotiated parameter.',
    '**Security by design.** The data plan separates open methods from sensitive asset-level results from the start, so openness and protection of critical infrastructure do not conflict.',
]


def all_risks():
    return [r for _, rs in CATS for r in rs]


def fig_matrix(path):
    risks = all_risks()
    W, H = 10.4, 5.85
    fig = plt.figure(figsize=(W, H), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(0, H); ax.axis('off')
    cs = 0.80                                            # cell size (inches)
    lnames = ['Rare', 'Unlikely', 'Possible', 'Likely', 'Almost\ncertain']
    inames = ['Minimal', 'Minor', 'Moderate', 'Major', 'Severe']
    for k, (title, li, ii) in enumerate([('Before mitigation', 2, 3), ('After mitigation', 8, 9)]):
        x0 = 1.25 + k * 4.95; y0 = 0.92
        ax.text(x0 + 2.5 * cs, y0 + 5 * cs + 0.25, title, fontsize=13, fontweight='bold', ha='center', va='center', color='#1F3864')
        cells = {}
        for r in risks:
            cells.setdefault((r[li], r[ii]), []).append(r[0])
        for L in range(1, 6):
            for I in range(1, 6):
                _, _, rgb = band(L * I)
                ax.add_patch(Rectangle((x0 + (I - 1) * cs, y0 + (L - 1) * cs), cs, cs, fc=rgb, ec='white', lw=2.0))
                ids = cells.get((L, I), [])
                if ids:
                    lines = [' '.join(ids[j:j + 2]) for j in range(0, len(ids), 2)]
                    ax.text(x0 + (I - 0.5) * cs, y0 + (L - 0.5) * cs, '\n'.join(lines), fontsize=9.6, fontweight='bold',
                            ha='center', va='center', color='#1f2328', linespacing=1.15)
        for I in range(1, 6):
            ax.text(x0 + (I - 0.5) * cs, y0 - 0.14, f'{I}', fontsize=10, ha='center', va='top', color='#1f2328', fontweight='bold')
            ax.text(x0 + (I - 0.5) * cs, y0 - 0.36, inames[I - 1], fontsize=8.6, ha='center', va='top', color='#52606D')
        ax.text(x0 + 2.5 * cs, y0 - 0.66, 'Impact', fontsize=10.5, ha='center', va='top', fontweight='bold', color='#1f2328')
        for L in range(1, 6):
            ax.text(x0 - 0.10, y0 + (L - 0.5) * cs, f'{L}', fontsize=10, ha='right', va='center', color='#1f2328', fontweight='bold')
            if k == 0:
                ax.text(x0 - 0.30, y0 + (L - 0.5) * cs, lnames[L - 1], fontsize=8.6, ha='right', va='center', color='#52606D')
        if k == 0:
            ax.text(0.12, y0 + 2.5 * cs, 'Likelihood', fontsize=10.5, rotation=90, ha='center', va='center', fontweight='bold', color='#1f2328')
    # legend
    lx = 2.55; ly = H - 0.22
    for j, (top, name, _, rgb) in enumerate(BANDS):
        lo = 1 if j == 0 else BANDS[j - 1][0] + 1
        ax.add_patch(Rectangle((lx + j * 1.75, ly - 0.09), 0.22, 0.18, fc=rgb, ec='#9aa3ad', lw=0.6))
        ax.text(lx + j * 1.75 + 0.30, ly, f'{name} ({lo}–{top})', fontsize=9.5, va='center', ha='left', color='#1f2328')
    ax.text(lx - 0.15, ly, 'Score = L × I:', fontsize=9.5, va='center', ha='right', color='#1f2328', fontweight='bold')
    fig.savefig(path, dpi=300)
    plt.close(fig)


def counts(li, ii):
    c = {name: 0 for _, name, _, _ in BANDS}
    for r in all_risks():
        c[band(r[li] * r[ii])[0]] += 1
    return c


def keep_together(t):
    """Header row repeats on a new page, and each category table is kept on one page with its header."""
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def landscape(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(1.5)


def portrait(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(1.8)


def bullets(doc, items):
    for t in items:
        p = para(doc, 3, 0.5); run(p, '• ', bold=True); rich(p, t)


def main(out):
    risks = all_risks()
    before, after = counts(2, 3), counts(8, 9)
    work = tempfile.mkdtemp(); fig = os.path.join(work, 'fig_risk_matrix.png'); fig_matrix(fig)

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Risk Assessment', bold=True, size=18, color='1F3864')
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    lead = (f'CERTAIN-COAST is exposed to risks in delivering the project and, because it issues warnings for critical '
            f'infrastructure, to risks that the warning service itself could create. The register below covers both: '
            f'{len(risks)} risks in seven categories, each scored before and after mitigation, with an owner, an early-warning '
            f'indicator with a numeric trigger, a mitigation that acts in advance and a contingency if the risk occurs. Before '
            f'mitigation, {before["Very high"]} risk is very high, {before["High"]} are high and {before["Medium"]} medium; after '
            f'mitigation none is high or very high, {after["Medium"]} remain medium and {after["Low"]} are low. The seven '
            f'risks of the Work Plan are included with the same ratings.')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; rich(p, lead, size=11.5)

    heading(doc, '1. How risks are assessed', 13.5, before=6)
    p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'Each risk is scored for likelihood L over the 36 months and for impact I on the objectives, both from 1 to 5. '
            'The score L × I sets its level: **low** (1–4), **medium** (5–9), **high** (10–14) or **very high** (15–25). '
            'The residual score is the score expected after the mitigation is in place.')
    table(doc, ['L', 'Likelihood', 'Probability over the project'], LIKELIHOOD, [1.0, 3.6, 12.4], size=10); para(doc, 4)
    table(doc, ['I', 'Impact', 'Effect on the project'], IMPACT, [1.0, 3.6, 12.4], size=10); para(doc, 4)

    heading(doc, '2. Highest risks and how they are reduced', 13.5, before=8)
    top = sorted(risks, key=lambda r: (-r[2] * r[3], r[0]))[:8]
    rows = [[r[0], r[1], f'{r[2]}×{r[3]} = {r[2] * r[3]}', f'{r[8]}×{r[9]} = {r[8] * r[9]}'] for r in top]
    t = table(doc, ['', 'Risk', 'Before', 'After'], rows, [1.1, 11.3, 2.3, 2.3], size=10)
    for i, r in enumerate(top, start=1):
        shade(t.rows[i].cells[2], band(r[2] * r[3])[1]); shade(t.rows[i].cells[3], band(r[8] * r[9])[1])

    landscape(doc)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(fig, width=Cm(26.4))
    p = para(doc, 6)
    run(p, 'Figure 1. Risk matrix before and after mitigation. Each cell lists the risks with that likelihood and impact; '
           'colours show the level of the score L × I. R14 (no strong event during the demonstration) stays at likelihood 4 '
           'because mitigation cannot create events; it reduces the impact through back-testing.', italic=True, size=10, color='52606D')

    heading(doc, '3. Risk register', 13.5, before=4)
    hdr = ['ID', 'Risk and consequence', 'L', 'I', 'Score', 'Early-warning indicator (trigger)', 'Mitigation (in advance)',
           'Contingency (if it occurs)', 'Owner', 'Residual']
    widths = [1.0, 4.8, 0.7, 0.7, 1.4, 3.9, 5.7, 4.6, 1.5, 2.4]
    for cname, rs in CATS:
        heading(doc, cname, 11.5, color='2E6DB4', before=6)
        rows = [[r[0], r[1], str(r[2]), str(r[3]), str(r[2] * r[3]), r[4], r[5], r[6], r[7], f'{r[8]}×{r[9]} = {r[8] * r[9]}']
                for r in rs]
        t = table(doc, hdr, rows, widths, size=8.5)
        for i, r in enumerate(rs, start=1):
            shade(t.rows[i].cells[4], band(r[2] * r[3])[1]); shade(t.rows[i].cells[9], band(r[8] * r[9])[1])
        keep_together(t)

    portrait(doc)
    heading(doc, '4. Risks the warning service itself could create', 13.5, before=0)
    p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'A warning system for critical infrastructure changes how people decide, and it produces information that has '
            'value beyond its users. These risks are part of the register, not an afterthought:')
    bullets(doc, SELF)
    heading(doc, '5. Monitoring, escalation and reporting (T7.3)', 13.5, before=8)
    bullets(doc, MONITOR)
    heading(doc, '6. What is new in this risk assessment', 13.5, before=8)
    bullets(doc, NEW)
    doc.save(out)
    print('saved', out, '| risks', len(risks), '| before', before, '| after', after)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Risk_Assessment.docx')
