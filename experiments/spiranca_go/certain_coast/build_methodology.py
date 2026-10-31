"""CERTAIN-COAST (NATO SPS proposal): Methodology. The six-step method chain (figure), the three hazards as one system
(indicator table), the working equations as native Word math (OMML, numbered): coastal motion and relative sea level,
the coupled model (variable-density flow, salt memory, compaction, moving sea boundary, coastal water level,
liquefaction with a moving water table), calibration, the certified surrogate, the certified alarms (online conformal
risk control, extreme-aware term, joint guarantee and allocation of tolerances, nonstationary extremes, back-testing),
goal-oriented network design, seismic velocity as a water-table gauge, live operation and the joint pumping rule,
pass tests and fallbacks, and what is new. Task numbers follow the Work Plan (build_workplan.py).

usage: python build_methodology.py OUT.docx"""
import json
import os
import re
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(1, os.path.join(HERE, '..', '..', 'ml_paradigms', 'word'))
from build_sota import run, para, heading, table  # noqa: E402
import omml  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')
COL = {'WP1': '#2e6db4', 'WP2': '#6a4fb3', 'WP3': '#b03a3a', 'WP4': '#1e8a5a', 'WP5': '#c27c12'}

# ---------------------------------------------------------------------------------------------------- content
# blocks: ('h', title) ('h2', title) ('p', text) ('b', text) ('eq', key, tex) ('table', header, rows, widths, size)
# ('fig',) ('cap', text). Text: $inline math$, **bold**, [[key]] -> equation number.
LEAD = ('CERTAIN-COAST turns three linked hazards into one warning chain whose reliability is stated in advance and checked in '
        'operation. The chain has six steps (Figure 1): observe what the warnings need (WP4), measure the moving coast (WP1), '
        'model the coast as one coupled system (WP2), copy the model fast with a certificate (WP2), certify each alarm (WP3), '
        'and operate, verify and hand over (WP5, WP6). Each step has a quantitative pass test and a fallback (Section 9). The '
        'equations are the working equations of the project; their parameters are estimated from the site data.')

INDICATORS = [
    ['Seawater intrusion (wellfields)', 'I₁: chloride cⱼ at each production well j', 'c* agreed with the utility (for example the drinking-water value of 250 mg/L)', 'Weeks to months; updated weekly', 'Chloride will exceed c* within the lead time unless pumping changes'],
    ['Liquefaction (port, energy landfall)', 'I₂: liquefaction potential index (LPI) under the design shaking at the current water table', 'Class boundary agreed with the operators', 'State updated daily; recomputed at once after an earthquake', 'If the design earthquake struck now, liquefaction would exceed the agreed level'],
    ['Extreme water level (all three)', 'I₃: total water level TWLₐ at asset a', 'Design or operational level of each asset', 'Hours to days for storms; return levels yearly', 'The water level will exceed the level the asset was built for'],
]

CHECKS = [
    ['Coastal motion (WP1)', 'Relative sea level and land motion at each asset; compaction parameters; preconsolidation head', 'Vertical-velocity uncertainty of the order of 1–2 mm/yr at asset points; InSAR and GNSS agree within uncertainty', 'Corner reflectors and GNSS at the assets; persistent scatterers on structures'],
    ['Coupled model (WP2)', 'I₁, I₂, I₃ from one forcing', 'Held-out heads, chloride, InSAR and seismic velocity fall inside the 90% predictive intervals at close to the nominal rate', 'Two-dimensional sections through each asset; wider error bound Δₖ'],
    ['Surrogate (WP2)', 'Îₖ with Δₖ', 'Coverage of Δₖ ≥ 1 − δ on test runs; thousands of storylines per day', 'Full model on a reduced domain for flagged forecasts'],
    ['Alarms (WP3)', 'Aₖ with αₖ', 'Back-tested miss rate ≤ αₖ (Clopper–Pearson interval); false-alarm rate reported', 'Margin widened: the miss target is kept and the extra false alarms are reported'],
    ['Network (WP4)', 'Installed sensor set', '≥ 90% data availability; forecast spread reduced as the design predicted', 'Design re-run with the remaining budget; seismic nodes if fibre is not available'],
    ['Seismic velocity (T4.4)', 'Calibrated δv/v–water-table relation', 'Relation significant against piezometers at the port or landfall', 'Piezometers alone for the water table; velocity kept for post-event damage'],
    ['Live operation (WP5)', 'Bulletin log; pumping rule', 'Observed miss rates within targets over 12 months', 'If no strong event occurs, guarantees rest on back-testing and storylines (T3.4)'],
]

NEW = [
    '**Relative sea level where it matters.** Sea level is referred to each asset, not to the tide gauge, by adding the differential land motion measured by InSAR and GNSS (Eq. [[rsl]]).',
    '**One model in which pumping moves the hazards in opposite directions.** The same water table drives intrusion, compaction and the liquefaction factor of safety, so every pumping decision is checked against all three at once, including a sudden pumping stop (Eq. [[rule]]).',
    '**A margin with three certified parts.** Surrogate error (Eq. [[cert]]), distribution-free control of the long-run miss rate (Eq. [[aci]]) and a physics-based widening beyond the observed record (Eq. [[ext]]), so the guarantee does not end where the data end.',
    '**Tolerances allocated by consequence.** The joint miss probability is bounded by Σαₖ, and the αₖ are set where tightening costs least in false alarms (Eq. [[alloc]]).',
    '**Sensors chosen for the warnings.** Every sensor type, from a piezometer to a fibre, is ranked by how much it shrinks the alarm margins per unit cost (Eq. [[design]]).',
    '**Seismic noise as a water-table gauge with a physical prior.** Grain-contact theory links seismic velocity to effective stress and so to the water table that controls liquefaction (Eq. [[dvv]]).',
]

REFS = [
    'Angelopoulos, A.N., Candès, E.J., Tibshirani, R.J. (2023). Conformal PID control for time series prediction. NeurIPS 36.',
    'Becker, R., Rannacher, R. (2001). An optimal control approach to a posteriori error estimation in finite element methods. Acta Numerica 10, 1–102.',
    'Boulanger, R.W., Idriss, I.M. (2014). CPT and SPT based liquefaction triggering procedures. Report UCD/CGM-14/01, University of California, Davis.',
    'Coles, S. (2001). An Introduction to Statistical Modeling of Extreme Values. Springer.',
    'Emerick, A.A., Reynolds, A.C. (2013). Ensemble smoother with multiple data assimilation. Computers & Geosciences 55, 3–15.',
    'Gibbs, I., Candès, E. (2021). Adaptive conformal inference under distribution shift. NeurIPS 34.',
    'Iwasaki, T., Tatsuoka, F., Tokida, K., Yasuda, S. (1978). A practical method for assessing soil liquefaction potential based on case studies at various sites in Japan. Proc. 2nd Int. Conf. on Microzonation, San Francisco.',
    'Riley, F.S. (1969). Analysis of borehole extensometer data from central California. IASH Publication 89, 423–431.',
    'Seed, H.B., Idriss, I.M. (1971). Simplified procedure for evaluating soil liquefaction potential. Journal of the Soil Mechanics and Foundations Division, ASCE 97(SM9), 1249–1273.',
    'Sens-Schönfelder, C., Wegler, U. (2006). Passive image interferometry and seasonal variations of seismic velocities at Merapi Volcano, Indonesia. Geophysical Research Letters 33, L21302.',
]

BLOCKS = [
    ('p', LEAD),
    ('fig',),
    ('cap', 'Figure 1. The CERTAIN-COAST method chain. Solid arrows: from data to bulletins. Dashed arrows: verified outcomes '
            'update the alarm margins and the model; the alarms tell the network which sensors are worth most.'),

    ('h', '1. The three hazards as one system'),
    ('p', 'The three assets share two physical variables: the relative sea level $R_a(t)$ at asset $a$ and the water table '
          '$z_w(\\mathbf{x},t)$. Sea level enters the aquifer at its seaward boundary and sets how far seawater advances; the water '
          'table sets the effective stress in the sands and so their resistance to liquefaction; land subsidence raises the '
          'relative sea level at every asset. Pumping lowers the water table: it increases intrusion and, below the '
          'preconsolidation head, permanent subsidence, but it raises the factor of safety against liquefaction. No hazard can '
          'be managed alone. CERTAIN-COAST defines one indicator per hazard, oriented so that larger means more dangerous (Table 1).'),
    ('cap', 'Table 1. Hazard indicators, thresholds and lead times.'),
    ('table', ['Hazard (assets)', 'Indicator', 'Threshold τₖ', 'Lead time, cadence', 'What an alarm means'], INDICATORS,
     [2.9, 3.6, 3.4, 3.0, 4.1], 9.5),
    ('p', 'An alarm for hazard $k$ is missed when $I_k$ exceeds $\\tau_k$ within the lead time and no alarm was issued. The **miss '
          'rate** is the long-run fraction of issue times with a missed exceedance; it is the quantity certified in Section 5, and '
          'the miss rate per event is reported alongside it. Earthquakes are not forecast: the liquefaction alarm describes the '
          'current state, how the ground would respond if the design earthquake struck now, and that state changes with the water table.'),

    ('h', '2. Measuring the moving coast (WP1; T1.1–T1.4)'),
    ('p', '**InSAR.** Sentinel-1 ascending and descending time series from 2014 onwards are processed with persistent-scatterer and '
          'small-baseline methods over the three asset areas. Combining the two viewing geometries gives the vertical and east '
          'components of motion at each point (the radar is nearly blind to north–south motion):'),
    ('eq', 'los', 'd^{\\,asc}_{LOS} = n^{asc}_{E}\\,u_E + n^{asc}_{U}\\,u_z, \\qquad d^{\\,dsc}_{LOS} = n^{dsc}_{E}\\,u_E + n^{dsc}_{U}\\,u_z'),
    ('p', '**Trajectory model.** At every point and GNSS station the vertical motion is fitted with a trend, a seasonal term, the '
          '2019 coseismic step and a postseismic transient, with $H$ the step function and $t_{eq}$ the date of the 26 November 2019 earthquake:'),
    ('eq', 'traj', 'u_z(t) = u_0 + v\\,t + A\\sin\\left(\\frac{2\\pi t}{T_{yr}} + \\phi\\right) + \\left[\\Delta_{eq} + P\\ln\\left(1 + \\frac{t - t_{eq}}{\\tau_p}\\right)\\right] H(t - t_{eq})'),
    ('p', 'Corner reflectors at the assets (T1.2) give stable radar targets on soft ground, and GNSS ties the InSAR solution to a global frame.'),
    ('p', '**Relative sea level at each asset.** The tide gauge records sea level $\\eta_{TG}(t)$ relative to its own benchmark, '
          'which also moves. The relative sea level at asset $a$ adds the differential land motion between the gauge and the asset:'),
    ('eq', 'rsl', 'R_a(t) = \\eta_{TG}(t) + u^{TG}_{z}(t) - u^{a}_{z}(t)'),
    ('p', 'It is propagated with its full covariance; the target is a vertical-velocity uncertainty of the order of 1–2 mm/yr at asset points.'),
    ('p', '**Elastic and inelastic compaction.** The aquifer system compacts reversibly while the head stays above the '
          'preconsolidation head $h_{pc}$, the lowest head it has experienced, and permanently below it, with $S_{kv} \\gg S_{ke}$ '
          'and $u_z = -\\Delta b$:'),
    ('eq', 'comp', '\\Delta b(t) = \\begin{cases} S_{ke}\\,[h_0 - h(t)], & h(t) \\ge h_{pc}(t) \\\\ S_{ke}\\,[h_0 - h_{pc}(t)] + S_{kv}\\,[h_{pc}(t) - h(t)], & h(t) < h_{pc}(t) \\end{cases}'),
    ('p', '$S_{ke}$ is estimated from the ratio of seasonal vertical motion to seasonal head change, and $S_{kv}$ and $h_{pc}$ from '
          'the long-term trend and the head history. Thick clay interbeds of thickness $b_0$ drain slowly, with time constant'),
    ('eq', 'tau', '\\tau = \\frac{S_{skv}\\,(b_0/2)^2}{K_v}'),
    ('p', 'so subsidence continues after heads recover; the lag is fitted from the InSAR record. The output is the head limit '
          '$h_{pc}(\\mathbf{x})$ above which pumping causes no new permanent subsidence, a hard constraint of the pumping rule (Section 7).'),

    ('h', '3. One coupled model of the coast (WP2; T2.1–T2.3)'),
    ('p', '**Variable-density flow.** Flow is written in equivalent freshwater head $h_f$, with density rising linearly with the '
          'seawater fraction $c$ and the production wells as sinks:'),
    ('eq', 'darcy', '\\mathbf{q} = -\\mathbf{K}_f\\left(\\nabla h_f + \\frac{\\rho - \\rho_f}{\\rho_f}\\,\\nabla z\\right), \\qquad \\rho = \\rho_f + (\\rho_s - \\rho_f)\\,c'),
    ('eq', 'mass', 'S_s\\,\\frac{\\partial h_f}{\\partial t} + \\nabla\\cdot\\mathbf{q} = -\\sum_j Q_j\\,\\delta(\\mathbf{x} - \\mathbf{x}_j)'),
    ('p', '$S_s$ contains the skeletal storage of Section 2, which couples flow and compaction.'),
    ('p', '**Salt with memory.** Clays and lenses store salt and release it long after the front has moved. Transport uses a '
          'mobile and an immobile domain exchanging salt at rate $\\omega$:'),
    ('eq', 'salt', '\\theta_m\\frac{\\partial c_m}{\\partial t} + \\theta_{im}\\frac{\\partial c_{im}}{\\partial t} = \\nabla\\cdot(\\theta_m\\mathbf{D}\\,\\nabla c_m) - \\mathbf{q}\\cdot\\nabla c_m, \\qquad \\theta_{im}\\frac{\\partial c_{im}}{\\partial t} = \\omega\\,(c_m - c_{im})'),
    ('p', 'The immobile domain explains why wells can stay saline for years after pumping is reduced.'),
    ('p', '**The sea as a moving boundary.** At the seaward boundary the equivalent freshwater head at elevation $z$ follows the '
          'relative sea level, tide and surge. This is where subsidence and sea-level rise enter the intrusion model:'),
    ('eq', 'sea', 'h_f(\\mathbf{x}_{sea}, t) = \\frac{\\rho_s}{\\rho_f}\\left[R(t) + \\eta_{tide}(t) + \\eta_{surge}(t)\\right] - \\frac{\\rho_s - \\rho_f}{\\rho_f}\\,z'),
    ('p', '**Extreme water level at each asset**, with the wave setup $\\eta_{wave}$ from offshore wave height, period and beach slope; '
          'its extremes are modelled in Section 5:'),
    ('eq', 'twl', 'TWL_a(t) = R_a(t) + \\eta_{tide}(t) + \\eta_{surge}(t) + \\eta_{wave}(t)'),
    ('p', '**Liquefaction with a moving water table.** The cyclic stress ratio and the factor of safety follow the simplified '
          'procedure (Seed and Idriss, 1971) in its CPT-based form (Boulanger and Idriss, 2014), with the water table taken from '
          'the flow model at every time step:'),
    ('eq', 'csr', 'CSR(z,t) = 0.65\\,\\frac{a_{max}}{g}\\,\\frac{\\sigma_v(z,t)}{\\sigma\'_v(z,t)}\\,r_d(z), \\qquad \\sigma\'_v = \\sigma_v - \\gamma_w\\,[z - z_w(\\mathbf{x},t)]^{+}'),
    ('eq', 'fs', 'FS(z,t) = \\frac{CRR_{7.5}(q_{c1Ncs})\\;MSF\\;K_\\sigma}{CSR(z,t)}, \\qquad LPI(\\mathbf{x},t) = \\int_0^{20}\\left[1 - FS(z,t)\\right]^{+}(10 - 0.5\\,z)\\,dz'),
    ('p', 'Here $\\sigma_v = \\gamma_m\\min(z, z_w) + \\gamma_{sat}(z - z_w)^{+}$, MSF is the magnitude scaling factor, $K_\\sigma$ the '
          'overburden correction, and LPI the liquefaction potential index (Iwasaki et al., 1978), used as the asset-level '
          'indicator $I_2$. Because $\\sigma\'_v$ grows as the water table falls, $\\partial FS/\\partial z_w > 0$: the pumping that '
          'drives intrusion protects against liquefaction. The design shaking $a_{max}$ comes from the national seismic hazard model '
          'with one-dimensional site response at the port and the landfall, and includes a 2019-type scenario.'),
    ('p', '**Calibration (T2.3).** The parameters $\\boldsymbol{\\theta}$ (conductivity field, $S_{ke}$, $S_{kv}$, immobile porosity '
          'and $\\omega$, dispersivities, boundary fluxes and the stress exponent of Section 6) are estimated jointly from heads, '
          'conductivity, chloride, InSAR and seismic velocity by an ensemble smoother with multiple data assimilation (Emerick and '
          'Reynolds, 2013), with inflation factors $\\beta_i$ and ensemble members $j$:'),
    ('eq', 'esmda', '\\boldsymbol{\\theta}^{a}_{j} = \\boldsymbol{\\theta}^{f}_{j} + \\mathbf{C}_{\\theta d}\\left(\\mathbf{C}_{dd} + \\beta_i\\mathbf{C}_D\\right)^{-1}\\left(\\mathbf{d}_{obs} + \\sqrt{\\beta_i}\\,\\mathbf{C}_D^{1/2}\\mathbf{z}_j - \\mathbf{g}(\\boldsymbol{\\theta}^{f}_{j})\\right), \\qquad \\sum_i \\beta_i^{-1} = 1'),
    ('p', 'The model passes when observations held out from calibration fall inside its 90% predictive intervals at close to the nominal rate.'),

    ('h', '4. A fast copy of the model, with a certificate (T2.4)'),
    ('p', 'Warnings need thousands of model runs per day for storylines, pumping options and ensemble members. A neural or kernel '
          'operator $\\hat{I}_k(\\xi)$ is trained on runs of the coupled model, where $\\xi$ collects the forcing history, the pumping '
          'schedule and the parameters. Each forecast carries an error bound built from two parts: a dual-weighted residual '
          'indicator $\\eta_k(\\xi)$, the residual of the full model equations at the surrogate solution weighted by the adjoint '
          'solution for $I_k$ (Becker and Rannacher, 2001), which is large where the surrogate is wrong for the quantity that '
          'matters; and the conformal quantile $\\hat{q}_{1-\\delta}$ of the normalised errors $|I_k - \\hat{I}_k|/\\eta_k$ on $n$ '
          'held-out runs of the full model:'),
    ('eq', 'cert', '\\Delta_k(\\xi) = \\hat{q}_{1-\\delta}\\,\\eta_k(\\xi), \\qquad \\Pr\\left[\\,|I_k(\\xi) - \\hat{I}_k(\\xi)| \\le \\Delta_k(\\xi)\\,\\right] \\ge 1 - \\delta'),
    ('p', 'The guarantee holds for any surrogate and requires only that new inputs are exchangeable with the held-out runs; '
          'inputs outside that range are covered by the extreme-aware term of Section 5. Forecasts whose $\\Delta_k$ is too large '
          'are recomputed with the full model on a reduced domain.'),

    ('h', '5. Certified alarms (WP3; T3.1–T3.5)'),
    ('h2', '5.1 The alarm rule'),
    ('p', 'Hazard $k$ is forecast at lead time $L_k$. The alarm fires when the forecast comes within a margin of the threshold, '
          'and the margin has three parts, each covering a different source of error:'),
    ('eq', 'alarm', 'A_k(t) = \\mathbf{1}\\left\\{\\hat{I}_k(t + L_k) \\ge \\tau_k - m_k(t)\\right\\}, \\qquad m_k(t) = \\Delta_k(t) + \\gamma_k(t) + e_k(t)'),
    ('h2', '5.2 Online conformal risk control'),
    ('p', 'After each outcome is observed, the margin rises if the alarm missed and falls slightly otherwise. The long-run miss '
          'rate then converges to its target for any sequence of data, without assuming stationarity (Gibbs and Candès, 2021; '
          'Angelopoulos et al., 2023), which is why the guarantee survives trends in sea level, pumping and climate:'),
    ('eq', 'aci', '\\gamma_k(t+1) = \\gamma_k(t) + \\lambda\\,[\\,miss_k(t) - \\alpha_k\\,], \\qquad \\left|\\frac{1}{T}\\sum_{t=1}^{T} miss_k(t) - \\alpha_k\\right| \\le \\frac{C_k}{\\lambda\\,T}'),
    ('p', 'Outcomes arrive after the lead time; the update is applied with that delay, which changes only the constant $C_k$.'),
    ('h2', '5.3 Extreme-aware term'),
    ('p', 'Conformal calibration learns from conditions already seen. When the forcing leaves the observed record, in a drought '
          'deeper than any recorded or a surge on a lower coast, the margin widens in proportion to how far it has left. Here '
          '$d$ is the distance of the current forcing from the observed record $\\Xi_{obs}$, and $\\hat{L}_k$ the sensitivity of '
          'the indicator to the forcing, estimated from the coupled model over a set of storylines $\\mathcal{S}$ with a safety '
          'factor; inside the record $e_k = 0$:'),
    ('eq', 'ext', 'e_k(t) = \\hat{L}_k\\,d\\left(\\xi(t), \\Xi_{obs}\\right), \\qquad \\hat{L}_k = \\max_{\\xi, \\xi\' \\in \\mathcal{S}} \\frac{|I_k(\\xi) - I_k(\\xi\')|}{\\|\\xi - \\xi\'\\|}'),
    ('h2', '5.4 Joint guarantee and allocation of tolerances'),
    ('p', 'The three alarms share inputs, so their misses are dependent. The probability that any of them is missed is bounded '
          'by the sum of the tolerances, and the tolerances are allocated where tightening them costs least in false alarms '
          '(FAR), weighted by consequence $c_k$, within caps $\\alpha_k^{max}$ agreed with the operators:'),
    ('eq', 'union', '\\Pr\\left(\\bigcup_k miss_k\\right) \\le \\sum_k \\alpha_k \\le \\bar{\\alpha}'),
    ('eq', 'alloc', '\\min_{\\alpha_1, \\dots, \\alpha_K} \\sum_k c_k\\,FAR_k(\\alpha_k) \\quad s.t. \\quad \\sum_k \\alpha_k \\le \\bar{\\alpha}, \\;\\; \\alpha_k \\le \\alpha_k^{max} \\quad \\Rightarrow \\quad c_k\\,FAR\'_k(\\alpha^{*}_k) = -\\mu'),
    ('p', 'At the optimum the consequence-weighted marginal false-alarm cost is equal across hazards ($\\mu$ is the multiplier of the '
          'budget). The joint miss rate is also tracked directly with Eq. [[aci]] applied to the event "any alarm missed", which '
          'tightens the union bound when the dependence is strong.'),
    ('h2', '5.5 Extremes on a sinking coast'),
    ('p', 'Exceedances of the still water level over a threshold follow a generalised Pareto distribution (Coles, 2001) whose '
          'threshold moves with the relative sea level of Section 2, so that land motion from InSAR enters the return levels '
          'with its uncertainty; $\\lambda_u$ is the yearly rate of exceedances and $z_N$ the $N$-year level:'),
    ('eq', 'gpd', '\\Pr\\left(Y > y \\mid Y > u(t)\\right) = \\left[1 + \\xi_g\\,\\frac{y - u(t)}{\\sigma_g(t)}\\right]^{-1/\\xi_g}, \\qquad u(t) = u_0 + R_a(t) - R_a(t_0)'),
    ('eq', 'retl', 'z_N(t) = u(t) + \\frac{\\sigma_g(t)}{\\xi_g}\\left[(N\\lambda_u)^{\\xi_g} - 1\\right]'),
    ('p', 'These distributions generate the storm storylines used for $\\hat{L}_k$ and for back-testing.'),
    ('h2', '5.6 Back-testing (T3.4)'),
    ('p', 'The full chain is replayed over the InSAR era with the system frozen at each issue time, so no later data are used, '
          'and on storylines beyond the record: droughts, compound storm and high tide, a 2019-type earthquake at a high water '
          'table, and increased pumping. Miss and false-alarm rates are reported with Clopper–Pearson intervals, using block '
          'resampling for serial dependence. In preparatory synthetic tests on a coupled variable-density model, an extreme-aware '
          'margin kept the share of seasons above the chloride limit at 4–9% across drought, human-feedback and compound '
          'storylines, where a static alarm allowed 13–21%.'),

    ('h', '6. A monitoring network designed for the warnings (WP4; T4.1, T4.4)'),
    ('p', '**Goal-oriented design.** Candidate sensors $\\mathcal{C}$ (piezometers with conductivity loggers, chloride sampling '
          'points, GNSS, corner reflectors, seismic stations and nodes, fibre-optic sensing) are ranked by how much they narrow '
          'the forecast of each indicator. A narrower forecast needs a smaller margin for the same miss rate, and so gives fewer '
          'false alarms. With a linearised model, the forecast spread of indicator $k$ after observing set $S$ uses the adjoint '
          'gradient $\\mathbf{g}_k = \\partial I_k/\\partial\\boldsymbol{\\theta}$, the prior covariance $\\boldsymbol{\\Gamma}_{pr}$, '
          'the sensor Jacobian $\\mathbf{J}_S$ and noise covariance $\\boldsymbol{\\Gamma}_S$; the network maximises the '
          'consequence-weighted reduction of margins within the budget $B$:'),
    ('eq', 'sigma', '\\sigma_k^2(S) = \\mathbf{g}_k^{\\top}\\left(\\boldsymbol{\\Gamma}_{pr}^{-1} + \\mathbf{J}_S^{\\top}\\boldsymbol{\\Gamma}_S^{-1}\\mathbf{J}_S\\right)^{-1}\\mathbf{g}_k'),
    ('eq', 'design', '\\max_{S \\subseteq \\mathcal{C}} \\sum_k w_k\\,z_{1-\\alpha_k}\\left[\\sigma_k(\\emptyset) - \\sigma_k(S)\\right] \\quad s.t. \\quad \\sum_{s \\in S} cost_s \\le B'),
    ('p', 'The selection is greedy by benefit per cost, checked against exhaustive search on small candidate sets, and repeated '
          'in Year 2 with the calibrated model and the value of each sensor reported by WP3.'),
    ('p', '**Seismic velocity as a water-table gauge (T4.4).** Ambient seismic noise recorded by stations, nodes or fibre is '
          'cross-correlated between sensors; a uniform relative change in seismic velocity stretches the correlation in time '
          '(Sens-Schönfelder and Wegler, 2006). Lowering the water table raises the effective stress, and grain-contact '
          '(Hertz–Mindlin) theory predicts velocity rising as the 1/6 power of effective stress, which gives a physical prior '
          '$n \\approx 1/6$ for the calibration against piezometers ($\\Delta T$ is temperature, $\\varepsilon$ noise):'),
    ('eq', 'dvv', '\\frac{\\delta v}{v} = -\\frac{\\delta t}{t}, \\qquad \\frac{\\delta v}{v}(t) \\approx n\\,\\frac{\\Delta\\sigma\'_v(t)}{\\bar{\\sigma}\'_v} + b\\,\\Delta T(t) + \\varepsilon(t), \\qquad \\Delta\\sigma\'_v \\approx (\\gamma_w - \\gamma_{sat} + \\gamma_m)\\,\\Delta z_w'),
    ('p', 'Once calibrated, velocity changes measure the water table between wells, at the port and landfall where liquefaction '
          'depends on it, and a drop in velocity after an earthquake measures the softening of the ground.'),

    ('h', '7. Live operation, joint pumping rule and post-event protocol (WP5)'),
    ('p', '**Operational cycle (T5.1, T5.2).** Data are assimilated, indicators forecast and margins applied at the cadence of '
          'each hazard (Table 1). Each bulletin states, per asset, the forecast, the margin, the alarm state and the certified '
          'miss rate, and is logged with its inputs and code version so that it can be reproduced and audited. Outcomes are '
          'recorded when known and feed the margin update of Eq. [[aci]].'),
    ('p', '**A pumping rule safe for all three hazards (T5.3).** With the utility, the wellfield schedule $\\mathbf{Q}(t)$ is '
          'chosen to supply as much water as possible while every certified alarm stays off and no new permanent subsidence occurs:'),
    ('eq', 'rule', '\\max_{\\mathbf{Q}(\\cdot)} \\sum_j \\int_0^{T} Q_j(t)\\,dt \\quad s.t. \\quad \\hat{I}_k(t;\\mathbf{Q}) + m_k(t) < \\tau_k \\;\\; \\forall k, t, \\qquad h(\\mathbf{x}, t; \\mathbf{Q}) \\ge h_{pc}(\\mathbf{x})'),
    ('p', 'The liquefaction constraint is also checked for a sudden pumping stop, after which the water table recovers and the '
          'factor of safety falls: the rule must stay safe in both directions. The problem is solved with the surrogate and the '
          'result is checked with the full model before it is proposed to the utility.'),
    ('p', '**Post-event protocol (T5.4).** When an earthquake, ground shaking or water level exceeds the triggers agreed with '
          'civil protection, heads, conductivity, chloride, GNSS and seismic velocity are measured within days, the next '
          'Sentinel-1 acquisitions are processed with priority, the model is updated, the alarms are recomputed, and the event '
          'joins the back-testing record.'),

    ('h', '8. Software, data and reproducibility (T6.3, T7.2)'),
    ('b', '**Open software.** The model, surrogate, alarms and network design are released under an open licence, in containers, with tests and documentation.'),
    ('b', '**Data.** Operator data are used under the access agreements of T7.2; derived products, the bulletin log and the project data are released openly following FAIR principles.'),
    ('b', '**Reproducibility.** Every bulletin and figure can be regenerated from the logged inputs and the code version.'),

    ('h', '9. Pass tests and fallbacks for each step'),
    ('cap', 'Table 2. What each step must achieve, and what happens if it does not.'),
    ('table', ['Step', 'Output', 'Pass test', 'Fallback'], CHECKS, [3.0, 3.4, 5.6, 5.0], 9.5),

    ('h', '10. What is new in this methodology'),
] + [('b', t) for t in NEW]


# ---------------------------------------------------------------------------------------------------- figure
def fig_chain(path):
    W, H = 6.7, 6.25
    fig = plt.figure(figsize=(W, H), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(0, H); ax.axis('off')
    stages = [
        ('Observe', 'WP4 · T4.1–T4.4', COL['WP4'],
         'InSAR · GNSS · tide gauge · piezometers\nwith conductivity · chloride · seismic\nnoise · fibre-optic sensing',
         'Data chosen by\nvalue for the\nwarnings'),
        ('Measure', 'WP1 · T1.1–T1.4', COL['WP1'],
         'LOS decomposition, trajectory model,\ndifferential land motion, elastic and\ninelastic compaction',
         '$R_a(t)$, $u_z(t)$,\n$h_{pc}(\\mathbf{x})$'),
        ('Model', 'WP2 · T2.1–T2.3', COL['WP2'],
         'Variable-density flow + salt memory +\ncompaction + FS($z_w$) + coastal water\nlevel; calibrated by ES-MDA',
         '$I_1$ chloride\n$I_2$ liquefaction\n$I_3$ water level'),
        ('Fast copy', 'WP2 · T2.4', COL['WP2'],
         'Neural or kernel operator; error bound =\nadjoint-weighted residual × conformal\nquantile',
         '$\\hat{I}_k \\pm \\Delta_k$\nin seconds'),
        ('Certify', 'WP3 · T3.1–T3.5', COL['WP3'],
         'Margin $m_k = \\Delta_k + \\gamma_k + e_k$;\n$\\alpha_k$ allocated by consequence;\njoint miss $\\leq \\Sigma\\,\\alpha_k$',
         'Alarm $A_k$,\nmiss rate $\\leq \\alpha_k$'),
        ('Operate', 'WP5 · T5.1–T5.4', COL['WP5'],
         'Bulletins to utility, port, landfall\noperator, civil protection; outcomes\nlogged; pumping rule; post-event',
         'Verified miss\nand false-alarm\nrates'),
    ]
    xL0, xT0, xT1, xM0, xM1, xO0, xO1 = 0.05, 0.42, 1.98, 2.06, 4.88, 4.96, 6.12
    bh, gap, top = 0.80, 0.19, H - 0.08
    ys = []
    for i, (name, wp, col, meth, out) in enumerate(stages):
        y1 = top - i * (bh + gap); y0 = y1 - bh; ym = (y0 + y1) / 2; ys.append((y0, y1, ym))
        ax.add_patch(FancyBboxPatch((xT0, y0), xT1 - xT0, bh, boxstyle='round,pad=0,rounding_size=0.08', fc=col, ec='none'))
        ax.text(xT0 + 0.13, ym + 0.13, f'{i + 1}', fontsize=15, fontweight='bold', color='white', va='center', ha='center')
        ax.text(xT0 + 0.32, ym + 0.13, name, fontsize=11.5, fontweight='bold', color='white', va='center', ha='left')
        ax.text(xT0 + 0.32, ym - 0.17, wp, fontsize=9, color='white', va='center', ha='left')
        ax.add_patch(FancyBboxPatch((xM0, y0), xM1 - xM0, bh, boxstyle='round,pad=0,rounding_size=0.08', fc='white', ec=col, lw=1.6))
        ax.text(xM0 + 0.10, ym, meth, fontsize=9.4, color='#1f2328', va='center', ha='left', linespacing=1.25)
        ax.add_patch(FancyBboxPatch((xO0, y0), xO1 - xO0, bh, boxstyle='round,pad=0,rounding_size=0.08', fc=col + '22', ec=col, lw=1.0))
        ax.text((xO0 + xO1) / 2, ym, out, fontsize=9.4, fontweight='bold', color='#1f2328', va='center', ha='center', linespacing=1.2)
        if i:
            ax.add_patch(FancyArrowPatch(((xT0 + xT1) / 2, ys[i - 1][0] - 0.01), ((xT0 + xT1) / 2, y1 + 0.01),
                                         arrowstyle='-|>', mutation_scale=13, color='#1f2328', lw=1.6))
            ax.add_patch(FancyArrowPatch(((xM0 + xM1) / 2, ys[i - 1][0] - 0.01), ((xM0 + xM1) / 2, y1 + 0.01),
                                         arrowstyle='-|>', mutation_scale=13, color='#9aa3ad', lw=1.3))

    def loop(x, i_from, i_to, label, side, dy=0.0):
        yf, yt = ys[i_from][2] + dy, ys[i_to][2]
        x_edge = xO1 if side == 'r' else xT0
        ax.plot([x_edge, x], [yf, yf], color='#1f2328', lw=1.3, ls=(0, (3, 2)))
        ax.plot([x, x], [yf, yt], color='#1f2328', lw=1.3, ls=(0, (3, 2)))
        ax.add_patch(FancyArrowPatch((x, yt), (x_edge + (0.01 if side == 'r' else -0.01), yt), arrowstyle='-|>',
                                     mutation_scale=12, color='#1f2328', lw=1.3))
        ax.text(x + (0.10 if side == 'r' else -0.11), (yf + yt) / 2, label, rotation=90, fontsize=8.8, va='center',
                ha='center', color='#1f2328', fontweight='bold')

    loop(6.22, 5, 4, 'margin update', 'r', dy=0.18)
    loop(6.48, 5, 2, 'new data assimilated', 'r', dy=-0.18)
    loop(0.24, 4, 0, 'value of each sensor', 'l')
    fig.savefig(path, dpi=300)
    plt.close(fig)


# ---------------------------------------------------------------------------------------------------- Word
SUBWORDS = ('asc', 'dsc', 'LOS', 'eq', 'yr', 'TG', 'pc', 'ke', 'kv', 'skv', 'sea', 'tide', 'surge', 'wave', 'max',
            'c1Ncs', 'sat', 'im', 'obs', 'pr')
WORDS = ('CSR', 'FS', 'CRR', 'MSF', 'LPI', 'TWL', 'FAR', 'miss', 'cost')


def upright(tex):
    """Descriptive names and subscripts upright; separators between joined equations written as words, so that both
    Word and LibreOffice show them clearly."""
    tex = re.sub(r'([_^])\{(?:\\,)?(' + '|'.join(SUBWORDS) + r')\}', r'\1{\\text{\2}}', tex)
    tex = re.sub(r'(?<![\\A-Za-z])(' + '|'.join(WORDS) + r')(?![A-Za-z])', r'\\text{\1}', tex)
    tex = tex.replace(', \\qquad ', ' \\quad\\text{ and }\\quad ').replace(' \\quad s.t. \\quad ', ' \\quad\\text{ subject to }\\quad ')
    return tex


class Eqs:
    def __init__(self):
        self.items = []

    def add(self, tex, display, size=None):
        tex = upright(tex)
        i = len(self.items)
        d = {'id': i, 'tex': tex, 'display': display}
        if size:
            d['size'] = size
        self.items.append(d)
        return f'§§EQ{i}§§' if display else f'§§IQ{i}§§'


def no_borders(t):
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'nil'); b.append(e)
    tblPr.append(b)


def fixed_widths(t, widths):
    t.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[j] * 567)))
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)


def mrich(p, text, E, size=11, numbers=None):
    """Text with $inline math$, **bold** and [[key]] equation numbers."""
    if numbers:
        text = re.sub(r'\[\[(\w+)\]\]', lambda m: f'({numbers[m.group(1)]})', text)
    for i, seg in enumerate(text.split('$')):
        if i % 2:
            run(p, E.add(seg, False, size=int(size * 2)), size=size)
            continue
        for j, s in enumerate(seg.split('**')):
            if s:
                run(p, s, bold=(j % 2 == 1), size=size)


def display_eq(doc, E, tex, number):
    t = doc.add_table(rows=1, cols=2)
    no_borders(t); fixed_widths(t, [15.6, 1.4])
    c0, c1 = t.rows[0].cells
    for c in (c0, c1):
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c0.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    run(p, E.add(tex, True))
    q = c1.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(q, f'({number})', size=11)
    trPr = t.rows[0]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))


def main(out):
    numbers, n = {}, 0
    for b in BLOCKS:
        if b[0] == 'eq':
            n += 1; numbers[b[1]] = n
    E = Eqs()
    work = tempfile.mkdtemp()
    fig = os.path.join(work, 'fig_method_chain.png')
    fig_chain(fig)

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Methodology', bold=True, size=18, color='1F3864')
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    for b in BLOCKS:
        kind = b[0]
        if kind == 'h':
            heading(doc, b[1], 13.5, before=10)
        elif kind == 'h2':
            heading(doc, b[1], 12, color='2E6DB4', before=6)
        elif kind == 'p':
            p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; mrich(p, b[1], E, numbers=numbers)
        elif kind == 'b':
            p = para(doc, 3, 0.5); run(p, '• ', bold=True); mrich(p, b[1], E, numbers=numbers)
        elif kind == 'eq':
            display_eq(doc, E, b[2], numbers[b[1]])
        elif kind == 'table':
            table(doc, b[1], b[2], b[3], size=b[4]); para(doc, 4)
        elif kind == 'fig':
            p = para(doc, 2, before=4); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(fig, width=Cm(17.0))
        elif kind == 'cap':
            p = para(doc, 6); p.paragraph_format.keep_with_next = (b[1].startswith('Table'))
            run(p, b[1], italic=True, size=10, color='52606D')
    heading(doc, 'Method references', 12, color='2E6DB4', before=10)
    for r_ in REFS:
        p = para(doc, 1, 0.5); run(p, r_, size=9.5, color='52606D')
    p = para(doc, 2, before=4)
    run(p, 'Standard method references only; the 30 key papers of 2020–2026 on which the innovations build are listed in the State of the Art.',
        italic=True, size=9.5, color='52606D')

    draft = os.path.join(work, 'draft.docx'); doc.save(draft)
    eq_json = os.path.join(work, 'eqs.json')
    with open(eq_json, 'w', encoding='utf-8') as fh:
        json.dump(E.items, fh, ensure_ascii=False)
    omml.main(draft, eq_json, out)
    shown = os.path.join(os.path.dirname(os.path.abspath(out)), 'fig_method_chain.png')
    os.replace(fig, shown)
    print('saved', out, '|', n, 'numbered equations,', len(E.items) - n, 'inline | figure', shown)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Methodology.docx')
