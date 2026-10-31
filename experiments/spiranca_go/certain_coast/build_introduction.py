"""CERTAIN-COAST (NATO SPS proposal): Introduction. The setting (verified figures), the problem as one coupled system
(figure with drivers, the two shared variables and the three hazards at the three assets), the gap, aim and objectives
(from build_objectives.py), the approach in brief, why NATO SPS, a reader's guide that maps each evaluation question to
the section that answers it, and what is new.

usage: python build_introduction.py OUT.docx"""
import os
import re
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade, link  # noqa: E402
import build_objectives as O  # noqa: E402
import build_workplan as W  # noqa: E402
import build_exec_summary as BE  # noqa: E402
import sota_data as D  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')

GUIDE = [
    ['Is it relevant to NATO security?', 'Security Relevance; Impact (Section 6)'],
    ['What is known, and what is new?', 'State of the Art; Innovation; References'],
    ['What exactly will be done?', 'Objectives; Methodology; Work Packages'],
    ['When, and how is progress judged?', 'Work Plan; Timeline; Milestones (evidence gates); Deliverables'],
    ['What can go wrong?', 'Risk Assessment (21 risks with triggers and fallbacks)'],
    ['Who does it, and are the users committed?', 'Partnership; CVs of the Project Directors; Letters of Commitment'],
    ['What will exist at the end, and for whom?', 'Expected Results; Impact; Dissemination Plan'],
    ['Is it value for money, and will it last?', 'Budget; Sustainability'],
    ['Is it done responsibly?', 'Ethics; Gender Balance; Data Management Plan'],
]

NEW = [
    '**The problem drawn as one system.** Figure 1 shows the drivers, the two shared variables and the three hazards with the direction of each effect, including the pumping trade-off that makes separate warnings unsafe.',
    '**Every objective traced forward.** Each objective names the deliverable that proves it and the innovations it carries, so the reader can follow a claim from this page to its test.',
    '**A reader\'s guide by question.** The guide maps the questions an evaluator asks to the sections that answer them, rather than listing chapters.',
]

SOURCES = [
    ('INSTAT: Maritime Transport Statistics, Q2 2026', 'https://www.instat.gov.al/en/themes/industry-trade-and-services/transport-accidents-and-characteristics-of-road-vehicles/publikimet/2026/maritime-transport-statistics-q2-2026/'),
    ('EU Delegation to Albania: Release of the Post-Disaster Needs Assessment for Albania\'s recovery (February 2020)', 'https://www.eeas.europa.eu/delegations/albania/release-post-disaster-needs-assessment-albania-recovery_en'),
    (D.LOCAL[0][1].rstrip('.'), D.LOCAL[0][2]),
    ('NATO: Resilience, civil preparedness and Article 3 (baseline requirements)', 'https://www.nato.int/cps/en/natohq/topics_132722.htm'),
    ('NATO Science for Peace and Security Programme (overview and Key Priorities)', 'https://www.nato.int/cps/en/natohq/85291.htm'),
]


# ------------------------------------------------------------------------------------------------- figure
def fig_system(path):
    Wd, Ht = 10.0, 6.75
    fig = plt.figure(figsize=(6.7, 6.7 * Ht / Wd), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, Wd); ax.set_ylim(0, Ht); ax.axis('off')
    cols = {'drv': ('#52606D', '#EEF0F2'), 'var': ('#1F3864', '#E8EDF5'), 'haz': ('#B03A3A', '#FBE9E7')}
    for x, lab in ((1.3, 'Drivers'), (5.0, 'Shared variables'), (8.7, 'Hazards at the assets')):
        ax.text(x, Ht - 0.28, lab, ha='center', va='center', fontsize=10, fontweight='bold', color='#1f2328')

    def box(cx, cy, w, h, title, sub, kind):
        ec, fc = cols[kind]
        ax.add_patch(FancyBboxPatch((cx - w / 2, cy - h / 2), w, h, boxstyle='round,pad=0,rounding_size=0.08', fc=fc, ec=ec, lw=1.4))
        ax.text(cx, cy + (0.15 if sub else 0), title, ha='center', va='center', fontsize=9.2, fontweight='bold', color=ec)
        if sub:
            ax.text(cx, cy - 0.2, sub, ha='center', va='center', fontsize=8.6, color='#1f2328')
        return (cx - w / 2, cx + w / 2, cy - h / 2, cy + h / 2)

    B = {}
    B['sea'] = box(1.3, 5.5, 2.4, 0.8, 'Sea level η', 'tides, storms, surge', 'drv')
    B['land'] = box(1.3, 4.15, 2.4, 0.8, 'Land motion u_z', 'compaction, co-seismic', 'drv')
    B['pump'] = box(1.3, 2.8, 2.4, 0.8, 'Pumping', 'the operators\' lever', 'drv')
    B['eq'] = box(1.3, 1.35, 2.4, 0.8, 'Earthquakes', 'shaking', 'drv')
    B['rsl'] = box(5.0, 4.75, 2.5, 0.9, 'Relative sea level', 'R(t) = η(t) − u_z(t)', 'var')
    B['wt'] = box(5.0, 2.5, 2.5, 0.9, 'Water table', 'pore pressure', 'var')
    B['ext'] = box(8.7, 5.5, 2.4, 0.8, 'Extreme water level', 'all three assets', 'haz')
    B['int'] = box(8.7, 3.6, 2.4, 0.8, 'Seawater intrusion', 'wellfields', 'haz')
    B['liq'] = box(8.7, 1.45, 2.4, 0.8, 'Liquefaction', 'port, energy landfall', 'haz')

    def arrow(p, q, lab, t=0.5, col='#52606D', dy=0.0):
        ax.add_patch(FancyArrowPatch(p, q, arrowstyle='-|>', mutation_scale=11, color=col, lw=1.3, shrinkA=2, shrinkB=2))
        if lab:
            x = p[0] + t * (q[0] - p[0]); y = p[1] + t * (q[1] - p[1]) + dy
            ax.text(x, y, lab, ha='center', va='center', fontsize=8.6, color=col, style='italic',
                    bbox=dict(boxstyle='round,pad=0.12', fc='white', ec='none'))

    R_, L_, RV, LH = 2.5, 3.75, 6.25, 7.5
    arrow((R_, 5.62), (LH, 5.62), 'surge, waves', t=0.5)
    arrow((R_, 5.38), (L_, 4.95), None)
    arrow((R_, 4.15), (L_, 4.6), None)
    arrow((1.3, 3.2), (1.3, 3.75), None)
    ax.text(1.42, 3.47, 'compaction', ha='left', va='center', fontsize=8.6, color='#52606D', style='italic')
    arrow((R_, 2.75), (L_, 2.55), 'lowers', t=0.5, dy=0.17)
    arrow((R_, 1.35), (LH, 1.4), 'shaking', t=0.5)
    red = '#B03A3A'
    arrow((RV, 4.95), (LH, 5.38), 'higher: more', t=0.5, col=red, dy=-0.02)
    arrow((RV, 4.55), (LH, 3.8), 'higher: more', t=0.5, col=red)
    arrow((RV, 2.7), (LH, 3.4), 'lower: more', t=0.5, col=red)
    arrow((RV, 2.3), (LH, 1.6), 'shallower: more', t=0.62, col=red, dy=-0.12)
    ax.text(5.0, 0.4, 'Pumping lowers the water table: more intrusion and subsidence, less liquefaction',
            ha='center', va='center', fontsize=8.8, color='#92400E', fontweight='bold',
            bbox=dict(boxstyle='round,pad=0.3', fc='#FDE7C2', ec='none'))
    fig.savefig(path, dpi=300); plt.close(fig)


# ------------------------------------------------------------------------------------------------- Word
def head_repeat(t):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t


def bullets(doc, items, size=10.5):
    for x in items:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); rich(q, x, size=size)


def main(out):
    F = BE.facts()
    work = tempfile.mkdtemp(); f1 = os.path.join(work, 'fig_system.png'); fig_system(f1)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.7)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)

    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Introduction', bold=True, size=18, color='1F3864')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')

    heading(doc, '1. Three critical assets on one subsiding coast', 13.5, before=2)
    p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'Within a few kilometres of the Durrës coast in Albania, a NATO Ally since 2009, lie the wellfields that supply '
            'drinking water to the area, the country\'s main port, which handled 94.2% of Albania\'s sea freight in the second '
            'quarter of 2026 [1], and an energy landfall where supply arriving from the sea meets the national system. They stand '
            'on soft, compacting sediments near active faults. The Mw 6.4 earthquake of 26 November 2019 near Durrës killed 51 '
            'people, affected 202,291 and caused total effects of about €985 million [2]; field surveys found liquefaction at '
            'coastal, riverine and lagoonal sites [3].', size=10.5)
    p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'The three assets face three hazards: **seawater intrusion** into the aquifer that feeds the wells, **liquefaction** '
            'of the ground beneath the port and the landfall during an earthquake, and **extreme water levels** from storms on a '
            'rising relative sea. Some threats act slowly, through subsidence and the landward advance of salt water, others '
            'suddenly, through earthquakes and surges, and the slow ones make the sudden ones worse.', size=10.5)

    heading(doc, '2. One system, not three problems', 13.5, before=8)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'The three hazards are moved by the same two variables, the **relative sea level** R(t), the sea level minus the '
            'vertical motion of the land, and the **water table** that sets pore pressure in the ground (Figure 1). Pumping, the '
            'one lever the operators control, pulls them in opposite directions: a lower water table draws in salt water and '
            'compacts the clays, raising R(t), but it also makes the ground less prone to liquefaction. A warning that looks at one '
            'hazard at a time can therefore recommend an action that shifts the risk from one asset to another.', size=10.5)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(f1, width=Cm(16.0))
    p = para(doc, 4)
    run(p, 'Figure 1. The coast as one system: four drivers act through two shared variables on three hazards at the three '
           'assets. Red arrows: the direction in which each shared variable changes a hazard.', italic=True, size=9.5, color='52606D')

    heading(doc, '3. The gap', 13.5, before=8)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_with_next = True
    rich(p, 'The State of the Art, built on 30 key papers of 2020–2026, finds three gaps:', size=10.5)
    bullets(doc, [
        '**No early warning treats the three hazards as one system,** although they share relative sea level and pore pressure and are moved by the same pumping lever.',
        '**No warning for any of them carries a guarantee.** Certified methods assume exchangeable data or hold only on time average; none covers a slow physics model of a coastal hazard or states beyond the observed record.',
        '**The new sensors are measured but not designed.** InSAR time series, seismic noise and fibre are not chosen jointly with wells and samples to reduce missed warnings.',
    ])
    p = para(doc, 4, before=2); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'As a result, a decision-maker at the wells, the quay or the landfall cannot tell how often a warning will miss, '
            'and therefore what silence means.', size=10.5)

    heading(doc, '4. Aim and objectives', 13.5, before=8)
    p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, O.AIM, size=10.5)
    rows = []
    for code, title, _what, _ind, dels, gap in O.OBJ:
        inn = re.findall(r'\b[MP]\d\b', gap)
        rows.append([code, title, dels, ', '.join(inn) if inn else '—'])
    t = head_repeat(table(doc, ['', 'Objective', 'Proved by', 'Innovations'], rows, [1.0, 6.4, 7.2, 2.4], size=9.5))
    for r_ in t.rows[1:]:
        shade(r_.cells[0], 'E8EDF5')

    heading(doc, '5. The approach in brief', 13.5, before=10)
    gates = {c: m for c, m, *_ in W.MILESTONES}
    p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, f'CERTAIN-COAST runs for 36 months in {len(W.WPS)} work packages. It **measures** the moving coast with InSAR, '
            'GNSS and a monitoring network designed for the warnings; **models** the three hazards in one coupled model with a '
            'fast, error-certified copy; **certifies** each alarm with a margin that bounds its miss rate; and **operates** the '
            'system live with the water utility, the port authority, the landfall operator and civil protection for twelve months '
            f'before handing it over. Three evidence gates, at Months {gates["MS3"]}, {gates["MS4"]} and {gates["MS5"]}, decide '
            f'whether each phase starts. The requested budget is {BE.eur(F["total"])}.', size=10.5)

    heading(doc, '6. Why NATO SPS', 13.5, before=8)
    bullets(doc, [
        '**Resilience of an Ally:** water, transport and energy are three of NATO\'s seven baseline requirements for civil preparedness [4].',
        '**SPS Key Priorities:** Environment, Climate Change and Security; Energy Security; Innovation and Emerging and Disruptive Technologies [5].',
        '**Cooperation:** NATO-country and partner-country scientists work on a shared security problem with the operators of the assets, and the partner country takes the methods home.',
    ])

    heading(doc, '7. Reader\'s guide', 13.5, before=8)
    p = para(doc, 4); p.paragraph_format.keep_with_next = True
    rich(p, 'The proposal answers the questions an evaluator asks; each is answered in the sections below.', size=10.5)
    t = head_repeat(table(doc, ['Question', 'Where it is answered'], GUIDE, [6.5, 10.5], size=9.5))
    for r_ in t.rows[1:]:
        shade(r_.cells[0], 'EAF0F8')
        for c in r_.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True

    heading(doc, '8. What is new in this introduction', 13.5, before=10)
    bullets(doc, NEW)

    heading(doc, 'Sources', 11.5, before=8)
    for i, (txt, url) in enumerate(SOURCES, start=1):
        q = para(doc, 2, 0.5); run(q, f'[{i}] {txt}. ', size=9.5); link(q, url, url, size=8.5)
    doc.save(out)
    print('saved', out)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Introduction.docx')
