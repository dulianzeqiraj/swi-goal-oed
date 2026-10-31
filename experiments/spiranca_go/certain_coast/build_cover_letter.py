"""CERTAIN-COAST (NATO SPS proposal): cover letter from the NATO-country Project Director to the NATO SPS Programme.
Page 1: the letter (one page), with the applicant's name and e-mail filled in, the budget computed from build_budget.py,
and highlighted placeholders for everything not known for certain. Page 2: notes for use (to delete before sending).

usage: python build_cover_letter.py OUT.docx"""
import os
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import para, heading  # noqa: E402
from build_letters import run, ph  # noqa: E402
import build_budget as BB  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')
ADDRESS = ['NATO Science for Peace and Security (SPS) Programme', 'NATO Headquarters', 'Boulevard Léopold III',
           '1110 Brussels, Belgium']
ENCLOSED = ('Application form; Executive Summary; State of the Art; Objectives; Methodology; Work Plan and Timeline; Expected '
            'Results; Security Relevance; Risk Assessment; Budget; Partnership; Dissemination Plan; Sustainability; Gender '
            'Balance; Ethics; Data Management Plan; References; CVs of the Project Directors; letters of commitment of the four '
            'end-users and of the partner-country institution.')

NOTES = [
    '**Where it goes.** The application is submitted through the SPS grant platform. Check whether the platform accepts a cover letter as an attachment; if it does not, use the text in your correspondence with the SPS Programme.',
    '**Complete the placeholders.** Replace every yellow placeholder (date, title, institution, partner country and Project Director, telephone), print on your institution\'s letterhead, sign, and save as PDF.',
    '**Co-signature.** The partner-country Project Director may co-sign; otherwise delete the second signature block.',
    '**Eligibility statement.** Keep the sentence on other SPS grants only after both co-directors have confirmed it; an individual co-director may not hold more than one SPS grant at a time.',
    '**Consistency.** The figures in the letter (36 months, budget) are taken from the Budget; if the budget changes, rebuild the letter or edit the figure.',
    '**Deadline.** Confirm the deadline and phase of the 2027-1 call on nato.int before sending.',
]
STRONG = [
    '**One page that leads with the security problem,** three critical assets of an Ally on one subsiding coast, before any method.',
    '**The innovation stated as a checkable promise:** every warning carries a certified miss rate that is verified against the log of live operation.',
    '**Commitments, not intentions:** the end-users are named and their letters are enclosed; eligibility is confirmed up front.',
]


def main(out):
    total = sum(r['tot'] for r in BB.compute())
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.3); s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.4)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10.5)
    sz = 10.3

    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ph(p, '[LETTERHEAD OF THE INSTITUTION: Faculty of Geology and Mining, Polytechnic University of Tirana, to confirm]', size=9.5, bold=True)
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; ph(p, 'Tirana, [date]', size=sz)
    for i, line in enumerate(ADDRESS):
        p = para(doc, 0 if i < len(ADDRESS) - 1 else 8); run(p, line, size=sz)
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p, 'Subject: ', bold=True, size=sz)
    run(p, f'Application for a Multi-Year Project, call 2027-1: CERTAIN-COAST, {FULL}', bold=True, size=sz)
    p = para(doc, 5); run(p, 'Dear Members of the SPS Programme and of the Independent Scientific Evaluation Group,', size=sz)

    body = [
        f'On behalf of the consortium, I am pleased to submit CERTAIN-COAST for consideration under the NATO Science for Peace '
        f'and Security call 2027-1, as a Multi-Year Project of 36 months with a requested budget of €{total:,.0f}.',
        'The Durrës coast of Albania brings together, within a few kilometres, the wellfields that supply drinking water, the '
        'country\'s main port and an energy landfall. They stand on soft, subsiding ground near active faults, where the '
        'earthquake of 26 November 2019 caused liquefaction. Seawater intrusion, liquefaction and extreme sea levels threaten '
        'them through the same variables, yet they are monitored separately, and no existing warning states how often it misses.',
        'CERTAIN-COAST will build, demonstrate and hand over one early-warning system for the three hazards, driven by one '
        'physical model of the coast and by satellite, geodetic and field data. **Its central innovation is that every warning '
        'carries a certified miss rate,** guaranteed by a margin that combines a surrogate error bound, online conformal risk '
        'control and an extreme-aware term, and verified against the log of twelve months of live operation with the operators.',
        'The project addresses the SPS Key Priorities on Environment, Climate Change and Security, on Energy Security, and on '
        'Innovation and Emerging and Disruptive Technologies. It protects the water, port and energy infrastructure of an Ally, '
        'three of NATO\'s baseline requirements for national resilience, and brings machine learning into a security decision '
        'chain in line with NATO\'s principles of responsible use of artificial intelligence.',
        'The project is directed by me, as NATO-country Project Director, and by [title and name] of [institution], '
        '[eligible partner country], as partner-country Project Director. The water utility of the Durrës wellfields, the Durrës '
        'Port Authority, the operator of the energy landfall and the National Civil Protection Agency take part as end-users and '
        'have confirmed their commitment in the enclosed letters.',
        'We confirm that both co-directors are employed by [academic / governmental / non-profit] institutions and that neither '
        'currently holds another SPS grant. We would be glad to provide any further information, and we thank you for '
        'considering our application.',
    ]
    for t in body:
        p = para(doc, 5); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; ph(p, t, size=sz)

    p = para(doc, 16); run(p, 'Yours faithfully,', size=sz)
    p = para(doc, 0); ph(p, '[Signature]', size=sz)
    p = para(doc, 0); ph(p, '[Title] ', size=sz); run(p, 'Dulian Zeqiraj', bold=True, size=sz)
    p = para(doc, 0); run(p, 'NATO-country Project Director, CERTAIN-COAST', size=sz)
    p = para(doc, 0); ph(p, '[Position, institution]', size=sz)
    p = para(doc, 8); run(p, 'dulian.zeqiraj@fgjm.edu.al  |  ', size=sz); ph(p, '[telephone]', size=sz)
    p = para(doc, 0); ph(p, '[Optional co-signature: signature, title and name, Partner-country Project Director, institution]', size=9.5)
    p = para(doc, 0, before=6); run(p, 'Enclosed: ', bold=True, size=8.8); run(p, ENCLOSED, size=8.8, color='3A4450')

    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    p = para(doc, 4); ph(p, '[Notes for use: delete this page before sending.]', size=10, bold=True)
    heading(doc, 'How to use this letter', 13, before=2)
    for i, t in enumerate(NOTES, start=1):
        q = para(doc, 3, 0.5); run(q, f'{i}. ', bold=True); ph(q, t, size=10.5)
    heading(doc, 'What makes this letter strong', 13, before=8)
    for t in STRONG:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); ph(q, t, size=10.5)
    doc.save(out)
    print('saved', out, '| budget', total)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Cover_Letter.docx')
