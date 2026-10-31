"""CERTAIN-COAST (NATO SPS proposal): Innovation. The core innovation and the three gaps it closes; the eleven technical
innovations at a glance (M1-M5, P1-P5 from the State of the Art, and V1 from the Impact section) with the novelty check
and the test of each; a figure placing them in the warning chain; one card per innovation (today, what is new, why it
matters, how it is tested, novelty check); preliminary evidence; readiness now and at the end (TRL, EU definitions);
innovations in delivery; protection and sharing; what is new in this section.
Innovation names and their references are read from sota_data.py through build_references.innovations().

usage: python build_innovation.py OUT.docx"""
import os
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade, link  # noqa: E402
import sota_data as D  # noqa: E402
import build_references as R  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')
NO_PRECEDENT = {'M2', 'P3'}
KIND = {'M': ('Mathematical', '#1F3864', '#E8EDF5'), 'P': ('Physical', '#1E8A5A', '#E6F4EC'),
        'V': ('Decision value', '#B45309', '#FCEBD9'), 'O': ('Delivery', '#52606D', '#EEF0F2')}

# code -> today (with refs), what is new, why it matters, tested by, glance: beyond today, glance: tested in
CARD = {
    'M1': ('Distribution-free risk control and online conformal prediction give long-run guarantees [19–21], and conformal '
           'coverage exists for neural-operator surrogates [23]. They assume exchangeable data or hold only on time average, have '
           'not been applied to slow physics models of coastal hazards, and say nothing beyond the observed record.',
           'Alarm k fires when the forecast indicator exceeds its threshold minus a margin mₖ made of three certified parts: the '
           'surrogate error bound, an online conformal risk-control term and an extreme-aware term. The long-run miss rate stays '
           'at or below αₖ under nonstationarity and when the forcing leaves the record.',
           'The operator knows in advance how often a warning will be missed, and therefore what silence means.',
           'D3.1 proofs (M24); D3.2 back-testing (M27, gate MS5): miss rate ≤ αₖ with Clopper–Pearson interval; D5.1 live verification (M34).',
           'Guarantees exist for exchangeable data [19–23]; here they hold for a slow physics model, under nonstationarity and beyond the record',
           'D3.1–D3.2 (M24–27); D5.1'),
    'M2': ('Guarantees are stated one alarm at a time. Our search found no warning that controls the probability that any of '
           'several dependent alarms is missed.',
           'The three alarms share inputs, so their misses are dependent. The probability that any of them is missed is held at '
           'or below Σαₖ (union bound), tightened by joint calibration, and the αₖ are allocated by consequence.',
           'A water-supply alarm and a landfall alarm are not held to the same tolerance by accident; civil protection receives one '
           'joint guarantee for the three assets.',
           'T3.3 and D3.1 proofs (M24); joint miss rate in back-testing, D3.2 (M27).',
           'Single-alarm guarantees only',
           'D3.1 (M24); D3.2 (M27)'),
    'M3': ('Optimal experimental design has a unified theory [30], and groundwater can be sensed with seismic noise and fibre '
           '[16, 17]. No study chooses InSAR, wells, chemical samples, seismic noise and fibre together to minimise the miss '
           'probability of warnings.',
           'Seven sensor types are chosen jointly to maximise the expected reduction in the miss probability of the three alarms '
           'per unit cost.',
           'Every euro of equipment is spent where it most reduces missed warnings at the three assets.',
           'D4.1 design report (M6, MS1); after installation, forecast spread reduced as the design predicted (T4.3).',
           'Design for one quantity and one sensor type [30]; here three warnings and seven sensor types',
           'D4.1 (M6)'),
    'M4': ('Neural operators approximate PDE solution maps [25], kernel methods solve nonlinear PDEs [26], and conformal '
           'prediction calibrates operator learning [23]. These give approximation guarantees, not a certificate on each forecast '
           'of a coupled variable-density and poroelastic model.',
           'A neural or kernel operator of the coupled model whose error Δₖ is bounded a posteriori for every forecast and '
           'calibrated conformally, fast enough for thousands of storylines per day.',
           'Real-time warnings and storylines become affordable without losing the guarantee.',
           'D2.2 (M22, gate MS4): coverage of Δₖ ≥ 1 − δ on held-out full-model runs; fallback: full model on a reduced domain.',
           'Approximation guarantees [25, 26]; here a certificate on each forecast of a coupled model',
           'D2.2 (M22)'),
    'M5': ('Land motion at tide gauges changes its rate with fluid withdrawal and earthquakes [4]; sea-level change and its '
           'causes are reconstructed probabilistically [5]; extremes have become more frequent with relative sea-level rise [6]. '
           'Local extreme-value models still use a straight-line land-motion trend.',
           'Extreme water levels are modelled with InSAR/GNSS land motion as a time-varying covariate of the location parameter.',
           'Design and warning levels at the port and landfall follow the land as it actually moves, including after earthquakes.',
           'D1.1 (M12) and D3.2 (M27): the extreme-water-level alarm meets its miss target; the covariate model is compared with a '
           'linear trend on held-out years.',
           'Linear land-motion trend [4–6]; here measured, time-varying land motion',
           'D1.1 (M12); D3.2 (M27)'),
    'P1': ('Subsidence is measured with InSAR and GNSS and used to set sustainable pumping [1–3], but it is not used as a '
           'time-varying boundary condition of intrusion, flooding and liquefaction models at the same site.',
           'One relative sea level, R(t) = η(t) − u_z(t), with u_z from InSAR/GNSS, drives all three hazard models.',
           'The three warnings agree on where the sea is relative to the land, so they cannot contradict each other.',
           'D1.1 (M12, MS2): vertical-velocity uncertainty of the order of 1–2 mm/yr at asset points; InSAR and GNSS agree.',
           'Subsidence measured but not used as a shared boundary [1–3]',
           'D1.1 (M12)'),
    'P2': ('Liquefaction is mapped from data with spatial correlation [10], and the water table is known to control near-surface '
           'stiffness [12]; liquefaction after the 2019 Durrës earthquake was mapped [L1]. Susceptibility stays static, with a '
           'fixed design water table.',
           'A time-varying factor of safety, FS(x, t) = CRR / CSR(z_w(x, t)), with the water table from the calibrated groundwater '
           'model and checked by seismic velocity.',
           'The port and landfall learn how their liquefaction risk changes with pumping, drought and sea level, not only in a '
           'design scenario.',
           'D2.1 (M18, gate MS3): held-out heads and seismic velocity inside the 90% predictive intervals at close to the nominal rate.',
           'Static susceptibility, fixed water table [10, 12]',
           'D2.1 (M18)'),
    'P3': ('Pumping limits are set for one effect at a time, for example subsidence from InSAR [1]. Our search found no pumping '
           'rule evaluated against intrusion, subsidence and liquefaction together.',
           'Every pumping rule is evaluated against the three hazards at once: more pumping raises intrusion and subsidence but '
           'lowers the water table and liquefaction susceptibility.',
           'The warning recommends actions that do not shift the risk from one asset to another.',
           'T5.3 (M24–33): the rule keeps the three alarms within their targets together and is agreed with the water utility.',
           'One effect at a time [1]',
           'T5.3 (M33)'),
    'P4': ('An abrupt switch to irreversible compaction has been detected with InSAR, GNSS and wells and proposed as an early '
           'warning [2]; elastic and inelastic compaction can be separated geodetically [3]. The switch is not yet a formal alarm.',
           'The switch from elastic to inelastic compaction becomes an alarm in itself, and its permanent rise of R(t) is '
           'propagated into the intrusion and flooding warnings; a pumping limit h_pc avoids it.',
           'An irreversible loss of land elevation is caught before it permanently worsens all three hazards.',
           'T1.3 and D1.1 (M12): elastic and inelastic parameters and the head limit h_pc with uncertainty.',
           'Detected after the fact [2, 3]; here a formal alarm and a pumping limit',
           'D1.1 (M12)'),
    'P5': ('Ambient seismic noise tracks groundwater levels [16], telecom fibre senses soil moisture [17], and twenty years of '
           'seismic velocity separate droughts, floods and earthquakes [18], in dry or inland basins. Never used to warn of '
           'liquefaction.',
           'Seismic velocity change is calibrated as a pore-pressure and water-table sensor in saturated coastal sediments at the '
           'port and the landfall.',
           'A dense, low-cost view of the water table where liquefaction matters, from stations and fibre that already exist.',
           'T4.4 (M12–24): relation significant against piezometers at the port or landfall; fallback: piezometers alone.',
           'Dry or inland basins [16–18]; here saturated coastal sediments',
           'T4.4 (M24)'),
    'V1': ('The cost-loss model gives the relative economic value of a forecast for a user [a, b], computed after the fact from '
           'verification data.',
           'Because the miss rate is certified, each user\'s relative value has a guaranteed lower bound, V ≥ 1 − (rF(1−s) + αₖs(1−r)) '
           '/ (min(r, s) − rs), known before live operation for every cost-loss ratio r (Impact, Section 5).',
           'Each operator knows in advance what the warning is at least worth for its own decision.',
           'T3.1 (cost-loss ratios agreed, M13–15); D5.1 (M34): bound reported for each user and compared with the observed value.',
           'Value computed after the fact [a, b]; here guaranteed in advance',
           'D5.1 (M34)'),
}

DELIVERY = [
    ('O1', 'Verification registered in advance', 'The back-testing protocol of D3.2 is published with a DOI before the results are computed, and D5.1 publishes the observed miss and false-alarm rates of the live warnings.'),
    ('O2', 'Evidence gates with fixed thresholds', 'Live operation starts and the guarantees are stated only after gates MS3, MS4 and MS5 are passed, each with a go, go-with-conditions and no-go rule fixed in advance (Milestones).'),
    ('O3', 'Transfer tested by an outsider', 'The software release D6.1 is accepted only when someone outside the team installs it and reproduces the synthetic test coast.'),
]

STAGES = [('Observe', ['P1', 'M3', 'P5']), ('Model', ['M4', 'P2', 'M5']), ('Warn', ['M1', 'M2', 'P4']),
          ('Decide', ['P3', 'V1']), ('Verify', ['O1', 'O2', 'O3'])]
SHORT = {'P1': 'One relative sea\nlevel R(t)', 'M3': 'Sensors designed\nfor the warnings', 'P5': 'Seismic velocity as\nwater-table gauge',
         'M4': 'Surrogate with\ncertified error', 'P2': 'Liquefaction safety\nfactor in time', 'M5': 'Extremes with\nland motion',
         'M1': 'Certified miss rate\nfor each alarm', 'M2': 'Joint guarantee\nfor three alarms', 'P4': 'Compaction tipping\npoint as alarm',
         'P3': 'Pumping rule safe\nfor all three', 'V1': 'Certified value\nfor each user',
         'O1': 'Verification\nfixed in advance', 'O2': 'Evidence gates,\nfixed thresholds', 'O3': 'Transfer tested\nby an outsider'}

TRL = [
    ['Certified multi-hazard warning system as a whole', 'M1–M5, P1–P5', '3', '7', 'Proof of concept on a synthetic coast → twelve months of live operation with the operators'],
    ['Certified alarm margin and joint guarantee', 'M1, M2', '3', '7', 'Synthetic tests (4–9% vs 13–21%) → back-tested and verified live'],
    ['Error-certified surrogate of the coupled model', 'M4', '3', '7', 'Method components published → used for every live forecast'],
    ['Coupled model with R(t), FS(x,t) and compaction', 'P1, P2, P4', '3', '7', 'Preparatory coupled model → calibrated at the site and run live'],
    ['Multi-hazard sensor design', 'M3', '3', '7', 'Theory and preparatory design → network installed and streaming'],
    ['Seismic velocity as coastal water-table sensor', 'P5', '3', '6', 'Shown in inland basins → calibrated at the port or landfall'],
    ['Pumping rule safe for three hazards', 'P3', '2', '6', 'Concept formulated → evaluated and agreed with the utility'],
    ['Certified value for each user', 'V1', '2', '6', 'Bound derived → reported for each user from live data'],
]

PROTECT = [
    '**Open by default.** Software under Apache-2.0 and data under CC BY 4.0 (D6.1), papers in open access, so Allies and partners can use the innovations without licence barriers.',
    '**Restricted where security requires it.** Asset-level results, sensor locations and vulnerability details stay with the operators and authorities (two-tier dissemination, Data Management Plan).',
    '**Credit and reuse.** Each release carries a DOI and a citation file; the synthetic test coast lets others reproduce the results without access to restricted site data.',
]

NEW = [
    '**Every novelty claim is checked against the literature.** Each innovation names the 2020–2026 work it goes beyond, or states that no direct precedent was found in our search (M2, P3).',
    '**Every innovation is a testable claim.** Each has a pass test, a deliverable and a month, and the tests of M1, M4 and P2 decide the evidence gates.',
    '**Innovations placed in the warning chain.** Figure 1 shows where each acts, from observing the coast to verifying the warnings, so gaps and overlaps are visible.',
    '**Readiness measured, not asserted.** The TRL of each component is given now and at the end, on the EU definitions, with the evidence behind each step.',
]

EXTRA = [
    ('a', 'Richardson, D. S. (2000). Skill and relative economic value of the ECMWF Ensemble Prediction System. Quarterly Journal of the Royal Meteorological Society 126, 649–667.', 'https://www.ecmwf.int/en/elibrary/76153-skill-and-relative-economic-value-ecmwf-ensemble-prediction-system'),
    ('b', 'Murphy, A. H. (1977). The value of climatological, categorical and probabilistic forecasts in the cost-loss ratio situation. Monthly Weather Review 105(7), 803.', None),
    ('c', 'European Commission. Horizon 2020 Work Programme 2014–2015, General Annexes, Annex G: Technology readiness levels (TRL).', 'https://ec.europa.eu/research/participants/data/ref/h2020/wp/2014_2015/annexes/h2020-wp1415-annex-g-trl_en.pdf'),
]


def names():
    out = {c: n for c, n, _ in R.innovations()}
    out['V1'] = 'A certified value of the warning for each user'
    return out


def cited_keys():
    import re
    nums = set()
    for v in CARD.values():
        for grp in re.findall(r'\[([0-9,\s–-]+)\]', ' '.join(v)):
            for part in grp.split(','):
                part = part.strip()
                if re.match(r'^\d+[–-]\d+$', part):
                    a, b = re.split(r'[–-]', part); nums.update(range(int(a), int(b) + 1))
                elif part.isdigit():
                    nums.add(int(part))
    return sorted(nums)


# ------------------------------------------------------------------------------------------------- figures
def fig_chain(path):
    Wd, Ht = 10.0, 5.12
    fig = plt.figure(figsize=(6.7, 6.7 * Ht / Wd), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, Wd); ax.set_ylim(0, Ht); ax.axis('off')
    cw, gap = 1.84, 0.12
    x0s = [0.08 + i * (cw + gap) for i in range(len(STAGES))]
    for (stage, codes), x0 in zip(STAGES, x0s):
        y0, h = Ht - 0.75, 0.6
        tip = 0.16
        pts = [(x0, y0), (x0 + cw - tip, y0), (x0 + cw + 0.02, y0 + h / 2), (x0 + cw - tip, y0 + h), (x0, y0 + h), (x0 + tip, y0 + h / 2)]
        ax.add_patch(Polygon(pts, closed=True, fc='#1F3864', ec='none'))
        ax.text(x0 + cw / 2 + 0.05, y0 + h / 2, stage, ha='center', va='center', fontsize=10, color='white', fontweight='bold')
        for j, c in enumerate(codes):
            k = KIND[c[0]]
            yb = Ht - 1.0 - (j + 1) * 1.18
            ax.add_patch(FancyBboxPatch((x0 + 0.04, yb), cw - 0.08, 1.04, boxstyle='round,pad=0,rounding_size=0.07',
                                        fc=k[2], ec=k[1], lw=1.6 if c in NO_PRECEDENT else 1.0))
            ax.text(x0 + cw / 2, yb + 0.80, c + (' ★' if c in NO_PRECEDENT else ''), ha='center', va='center', fontsize=9.5,
                    fontweight='bold', color=k[1])
            ax.text(x0 + cw / 2, yb + 0.36, SHORT[c], ha='center', va='center', fontsize=8.6, color='#1f2328', linespacing=1.1)
    lx = 0.12
    for key in ('M', 'P', 'V', 'O'):
        lab, ec, fc = KIND[key]
        ax.add_patch(FancyBboxPatch((lx, 0.08), 0.28, 0.2, boxstyle='round,pad=0,rounding_size=0.04', fc=fc, ec=ec, lw=1.0))
        ax.text(lx + 0.36, 0.18, lab, va='center', fontsize=8.8, color='#1f2328')
        lx += 0.42 + 0.115 * len(lab)
    ax.text(lx + 0.05, 0.18, '★ no direct precedent found', va='center', fontsize=8.8, color='#1f2328')
    fig.savefig(path, dpi=300); plt.close(fig)


def fig_evidence(path):
    fig, ax = plt.subplots(figsize=(6.7, 1.95), dpi=300)
    rows = [('Static alarm', 13, 21, '#9aa3ad'), ('Extreme-aware margin (M1)', 4, 9, '#1F3864')]
    for i, (lab, a, b, col) in enumerate(rows):
        ax.plot([a, b], [i, i], color=col, lw=9, solid_capstyle='butt')
        ax.text(b + 0.6, i, f'{a}–{b}%', va='center', fontsize=9.5, color=col, fontweight='bold')
    ax.set_yticks([0, 1]); ax.set_yticklabels([r[0] for r in rows], fontsize=9.5)
    ax.set_ylim(-0.6, 1.6); ax.set_xlim(0, 25)
    ax.set_xticks([0, 5, 10, 15, 20, 25]); ax.set_xticklabels(['0%', '5%', '10%', '15%', '20%', '25%'], fontsize=9)
    ax.set_xlabel('Share of seasons above the chloride limit (range over storylines)', fontsize=9.5)
    ax.grid(axis='x', color='#e3e6ea', lw=0.8); ax.set_axisbelow(True)
    for s in ('top', 'right', 'left'):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis='y', length=0)
    fig.tight_layout(); fig.savefig(path, dpi=300); plt.close(fig)


# ------------------------------------------------------------------------------------------------- Word helpers
def head_repeat(t):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t


def bullets(doc, items, size=10.5):
    for x in items:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); rich(q, x, size=size)


def main(out):
    nm = names()
    order = ['M1', 'M2', 'M3', 'M4', 'M5', 'P1', 'P2', 'P3', 'P4', 'P5', 'V1']
    work = tempfile.mkdtemp()
    f1 = os.path.join(work, 'fig_chain.png'); fig_chain(f1)
    f2 = os.path.join(work, 'fig_evidence.png'); fig_evidence(f2)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.7)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)

    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Innovation', bold=True, size=18, color='1F3864')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')

    heading(doc, '1. The core innovation', 13.5, before=2)
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'CERTAIN-COAST builds **the first early warning for critical coastal infrastructure whose miss rate is certified in '
            'advance and verified in operation**: one system for seawater intrusion at the wellfields, liquefaction at the port and '
            'the energy landfall, and extreme sea levels at all three, driven by one physical model of a subsiding coast. It closes '
            'three gaps that the State of the Art identifies in the 2020–2026 literature:', size=10.5)
    bullets(doc, ['**No early warning treats the three hazards as one system,** although they share relative sea level and pore pressure and are moved by the same pumping lever (closed by P1, P2, P3, M2).',
                  '**No warning for any of them carries a guarantee:** certified methods assume exchangeable data or hold on time average, and none covers a slow physics model or states beyond the record (closed by M1, M4, M5).',
                  '**The new sensors are measured but not designed:** no study chooses InSAR, seismic noise and fibre jointly with wells and samples to reduce missed warnings (closed by M3, P5).'])
    p = para(doc, 4, before=2); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'Numbers in brackets refer to the 30 key papers of the State of the Art; [L1] is the local Durrës study and letters to '
            'the additional sources listed at the end.', size=9.5)

    heading(doc, '2. Innovations at a glance', 13.5, before=6)
    rows = []
    for c in order:
        nov = 'No direct precedent found' if c in NO_PRECEDENT else 'Goes beyond cited work'
        rows.append([c, nm[c], CARD[c][4], nov, CARD[c][5]])
    t = head_repeat(table(doc, ['Code', 'Innovation', 'Beyond today', 'Novelty check', 'Tested in'], rows,
                          [1.2, 4.0, 6.2, 2.6, 3.0], size=8.8))
    for i, c in enumerate(order, start=1):
        shade(t.rows[i].cells[0], KIND[c[0]][2])
        if c in NO_PRECEDENT:
            shade(t.rows[i].cells[3], 'FDE7C2')
    p = para(doc, 2, before=8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(f1, width=Cm(16.5))
    p = para(doc, 4)
    run(p, 'Figure 1. Where each innovation acts in the warning chain, from observing the coast to verifying the warnings. '
           'Colour: type of innovation; ★: no direct precedent found in our search.', italic=True, size=9.5, color='52606D')

    h = heading(doc, '3. The innovations one by one', 13.5, before=10); h.paragraph_format.keep_with_next = True
    for c in order:
        today, new, why, test = CARD[c][:4]
        lab, ec, fc = KIND[c[0]]
        p = para(doc, 2, before=6); p.paragraph_format.keep_with_next = True
        run(p, f'{c}  {nm[c]}', bold=True, size=11.5, color=ec.lstrip('#'))
        run(p, f'   {lab}' + (' · no direct precedent found' if c in NO_PRECEDENT else ''), size=9.5, color='52606D')
        kv = [['Today', today], ['What is new', new], ['Why it matters', why], ['Tested by', test]]
        tt = table(doc, ['', ''], kv, [3.0, 14.0], size=9.5)
        tt._tbl.remove(tt.rows[0]._tr)
        n = len(tt.rows)
        for ri, row in enumerate(tt.rows):
            shade(row.cells[0], fc.lstrip('#'))
            for pp in row.cells[0].paragraphs:
                for rn in pp.runs:
                    rn.bold = True
            if ri < n - 1:
                for pp in row.cells[0].paragraphs + row.cells[1].paragraphs:
                    pp.paragraph_format.keep_with_next = True

    heading(doc, '4. Preliminary evidence', 13.5, before=10)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.keep_with_next = True
    rich(p, 'In preparatory synthetic tests on a coupled variable-density model, an extreme-aware margin (the third part of M1) '
            'kept the share of seasons above the chloride limit at **4–9%** across drought, human-feedback and compound storylines, '
            'where a static alarm allowed **13–21%** (Figure 2). The value bound of V1 was checked against a direct calculation '
            'of the cost-loss model (Impact, Section 5).', size=10.5)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(f2, width=Cm(15.5))
    p = para(doc, 4)
    run(p, 'Figure 2. Preparatory synthetic result: range over storylines of the share of seasons above the chloride limit. '
           'These are tests on a synthetic coast, not results at Durrës.', italic=True, size=9.5, color='52606D')

    heading(doc, '5. Readiness now and at the end of the project', 13.5, before=10)
    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.keep_with_next = True
    rich(p, 'Technology readiness levels (TRL) on the EU definitions [c]: TRL 2 technology concept formulated, TRL 3 experimental '
            'proof of concept, TRL 6 technology demonstrated in a relevant environment, TRL 7 system prototype demonstration in an '
            'operational environment. Levels are the consortium\'s estimates.', size=10.5)
    t = head_repeat(table(doc, ['Component', 'Innovations', 'TRL now', 'TRL M36', 'Evidence behind the step'], TRL,
                          [4.6, 2.4, 1.5, 1.6, 6.9], size=9))
    for r_ in t.rows[1:]:
        shade(r_.cells[2], 'EEF0F2'); shade(r_.cells[3], 'E6F4EC')

    heading(doc, '6. Innovation in delivery', 13.5, before=10)
    for code, title, text in DELIVERY:
        q = para(doc, 3, 0.5); run(q, f'{code}. ', bold=True, color='52606D'); rich(q, f'**{title}.** {text}', size=10.5)

    heading(doc, '7. Protecting and sharing the innovation', 13.5, before=8)
    bullets(doc, PROTECT)

    heading(doc, '8. What is new in this section', 13.5, before=8)
    bullets(doc, NEW)

    heading(doc, 'Sources cited', 11.5, before=8)
    groups = {it['n']: it for _, items in D.GROUPS for it in items}
    for n in cited_keys():
        it = groups[n]
        q = para(doc, 1, 0.5); run(q, f'[{n}] {R.compact(it["cite"])} ', size=9)
        link(q, it['doi'], it['doi'].replace('https://', ''), size=8.5)
    for code, cite, url in D.LOCAL:
        q = para(doc, 1, 0.5); run(q, f'[{code}] {cite} ', size=9); link(q, url, url.replace('https://', ''), size=8.5)
    for code, cite, url in EXTRA:
        q = para(doc, 1, 0.5); run(q, f'[{code}] {cite}' + (' ' if url else ''), size=9)
        if url:
            link(q, url, url.replace('https://', ''), size=8.5)
    doc.save(out)
    print('saved', out, '| innovations', len(order), '+', len(DELIVERY), '| key papers cited', len(cited_keys()))


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Innovation.docx')
