"""CERTAIN-COAST (NATO SPS proposal): Milestones. Seven milestones at a glance, an 'evidence ladder' figure (each
milestone raises the central claim one rung, with the three evidence gates, live operation and the reports to SPS),
one card per milestone (what is achieved, pass test, evidence, who decides, risks watched, if not met, what it
unlocks), three-way rules for the gates, monitoring and reporting to the SPS Programme, and what is new.
Months, names and verification are read from build_workplan.py; risk triggers from build_risk.py; deliverable names
and access from build_deliverables.py.

usage: python build_milestones.py OUT.docx"""
import os
import re
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade, link  # noqa: E402
from build_letters import ph  # noqa: E402
import build_workplan as W  # noqa: E402
import build_risk as BR  # noqa: E402
import build_deliverables as BD  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')
GATES = {'MS3', 'MS4', 'MS5'}
LIVE = (23, 34)

# rung reached when the milestone is passed (the 'evidence ladder')
RUNG = {'MS1': 'Designed', 'MS2': 'Measured', 'MS3': 'Calibrated', 'MS4': 'Certified offline',
        'MS5': 'Confirmed on the record', 'MS6': 'Verified live', 'MS7': 'Owned by the operators'}

# code -> lead WP, what is achieved, pass test, who decides, risk ids, if not met, unlocks
CARD = {
    'MS1': ('WP4',
            'The monitoring network is designed for the three warnings, not for general coverage: every candidate sensor is ranked by its value for the warnings per unit cost.',
            'D4.1 approved by the steering committee and by the operators; the chosen network fits the equipment budget; installation and access plan agreed for each site.',
            'Steering committee (Month 6 meeting); operators co-sign the installation plan',
            ['R1', 'R2', 'R19'],
            'Sensors certain to be needed are already ordered in Month 4; the design is re-run with the remaining budget; where a site is not accessible, the nearest accessible ground is used.',
            'Procurement of the remaining equipment and installation (T4.2), so that the network streams by Month 12.'),
    'MS2': ('WP4, WP1',
            'The new network streams data, and the motion of the coast is known at each asset: land motion, compaction and relative sea level R(t).',
            'D4.2: at least 90% data availability over the first month. D1.1: vertical-velocity uncertainty of the order of 1–2 mm/yr at asset points, with InSAR and GNSS agreeing within their uncertainty.',
            'Steering committee (Month 12 meeting); operators confirm the data streams',
            ['R4', 'R9', 'R10', 'R11'],
            'Sensors installed in order of value and partner equipment borrowed for the first months; GNSS and levelling where InSAR loses coherence; calibration starts on the sensors already streaming.',
            'Calibration of the coupled model (T2.3); post-event protocol on standby; first progress report to SPS (D7.2).'),
    'MS3': ('WP2',
            'One coupled model reproduces the water, salt, ground and sea-level data it has not seen, with honest uncertainty.',
            'D2.1: held-out heads, chloride, InSAR and seismic velocity fall inside the 90% predictive intervals at close to the nominal rate (gate rule in Section 3).',
            'Steering committee at the gate meeting (Month 18, with the mid-term workshop)',
            ['R3', 'R5'],
            'Two-dimensional sections through each asset; a larger error bound Δₖ; the limitation is stated in every later result.',
            'Training of the certified surrogate on the calibrated parameters (T2.4).'),
    'MS4': ('WP2, WP3',
            'Go or no-go for live operation: the fast surrogate meets its error bound, and the whole chain from data to bulletin works.',
            'D2.2: coverage of the error bound Δₖ at least 1 − δ on held-out full-model runs; thousands of storylines per day; the alarm prototype passes a dry run of [four weeks] on live data without issuing bulletins.',
            'Steering committee at the gate meeting (Month 22); the operators agree to receive bulletins',
            ['R6', 'R17'],
            'A slip of up to two months is absorbed: live operation runs from Month 25 to Month 36, still twelve months. A longer slip triggers the fallbacks of R5 and R6, decided by the steering committee.',
            'Live bulletins to the operators and civil protection from Month 23 (T5.2).'),
    'MS5': ('WP3',
            'The central promise holds on the past: over the record since 2014 and on storylines beyond it, each alarm misses no more often than its certified rate.',
            'D3.2: back-tested miss rate of each alarm at or below its target αₖ, with its Clopper–Pearson interval; false-alarm rate reported against the level agreed with the operators (T3.1). Protocol published with a DOI before the results.',
            'Steering committee at the gate meeting (Month 27); operators review the false-alarm rates',
            ['R7', 'R8'],
            'The margin is widened until the miss target is met: the target is kept and the extra false alarms are reported; if false alarms are too many, staged alerts (watch, then warning) and longer lead-time windows.',
            'The certified rates are stated in the live bulletins; the joint pumping rule is agreed with the utility (by Month 33).'),
    'MS6': ('WP5',
            'Twelve months of live warnings, every season included, with every bulletin logged and checked against what happened.',
            'D5.1: bulletin log complete; observed miss rates within their targets, with confidence intervals; post-event protocol executed for every event above its trigger.',
            'Steering committee with the operators and civil protection, who co-sign the bulletin log',
            ['R14', 'R15', 'R17', 'R21'],
            'If no strong event occurs, event-based verification rests on back-testing and storylines (T3.4) and the bulletin log continues after the hand-over; any missed exceedance is reviewed within one week.',
            'The operators issue every bulletin themselves in Months 34–36; verification statistics are published (D5.1).'),
    'MS7': ('WP6',
            'The operators own and run the system after the project, with procedures, trained staff and a maintenance agreement.',
            'D6.2: at least one operating procedure adopted per asset; the operators issue every bulletin in Months 34–36 without the project team; at least two trained people per operator; maintenance agreement signed (by Month 33); policy brief delivered to the national authorities.',
            'The operators (acceptance of the hand-over) and the steering committee (Month 36)',
            ['R16', 'R18'],
            'Extra training session and tasks reassigned within the consortium; the procedures are restated jointly with the operator; the university hosts the system until the agreement is in place.',
            'Operation after the grant at a running cost of about €{cost:,.0f} per year, shared by the operators; transfer to other coasts.'),
}

# gate -> go, go with conditions, no-go and fallback
GATE_RULES = [
    ['MS3, Month 18\nCoupled model calibrated',
     'Coverage of the 90% predictive intervals for held-out heads, chloride, InSAR and seismic velocity at least [85%]',
     'Coverage between 80% and [85%]: continue with a larger error bound Δₖ; the limitation is reported',
     'Coverage below 80% (risk trigger R5): two-dimensional sections through each asset; the surrogate is trained on them'],
    ['MS4, Month 22\nGo/no-go for live operation',
     'Error bound met (coverage at least 1 − δ); thousands of storylines per day; dry run passed',
     'Bound fails in part of the domain (Δₖ above half the margin for more than 5% of forecasts, R6): go, with the full model on a reduced domain for flagged forecasts',
     'Bound not met or dry run failed: fix and re-test within two months (live operation Months 25–36); longer: R5/R6 fallbacks'],
    ['MS5, Month 27\nMiss-rate targets confirmed',
     'Every back-tested miss rate at or below αₖ (Clopper–Pearson interval reported); false-alarm rate within the agreed level',
     'Miss targets met, false alarms above the agreed level (R7): staged alerts and longer lead-time windows; trade-off reported to the operators',
     'A miss rate above αₖ: margin widened until the target is met; the target is kept and the extra false alarms are reported'],
]

MONITOR = [
    '**Owner.** Each milestone has an owner, the lead of its work package, who files the evidence (the deliverables and the test results) in the project repository, with restricted access where the data plan requires it.',
    '**Thresholds fixed in advance.** The pass tests of the three gates are written into the first data management plan (D7.1, Month 3) and, for MS5, into the back-testing protocol published with a DOI before the results are computed.',
    '**Decision.** The steering committee decides each milestone at its regular meetings (Months 6, 12, 18, 24, 30, 36) or at a dedicated meeting for MS4, MS5 and MS6, held within [two weeks] of the due month; online meetings are allowed. Decisions and the evidence used are minuted.',
    '**Early warning on milestones.** The monthly risk review watches the risk triggers listed on each milestone card, so a milestone at risk is seen months before its date; a crossed trigger is escalated within two weeks.',
    '**Change.** A missed milestone is not moved silently: the steering committee records the cause, applies the fallback on the card and informs the SPS Programme in the next report.',
]

REPORTS = [
    ['D7.2, first progress report', 'Month 12', 'MS1, MS2', 'Status of each milestone and deliverable; revised Gantt chart; milestones of the next period'],
    ['D7.2, second progress report', 'Month 24', 'MS3, MS4', 'As above; gate decisions and any fallback applied'],
    ['D7.3, final report', 'Month 36', 'MS5, MS6, MS7 (and all earlier)', 'Status of every milestone; for any not met, why and the effect on the results'],
]

NEW = [
    '**An evidence ladder, not a list of dates.** Each milestone raises the central claim one rung, from designed to owned by the operators, so a reviewer can read at any month how much of the promise has been proven (Figure 1).',
    '**Three-way gates with thresholds fixed in advance.** Each gate has a go, a go-with-conditions and a no-go with a named fallback, so the project keeps delivering usable results even when a gate is not passed.',
    '**The promise itself is a milestone.** MS5 and MS6 test the certified miss rate of each alarm, on the record and live, with confidence intervals, rather than only the delivery of a report.',
    '**The users decide when it is ready.** The operators co-sign the milestones that change what they receive or run: the network plan (MS1), live bulletins (MS4), the bulletin log (MS6) and the hand-over (MS7).',
    '**Milestones watched like hazards.** Each milestone carries the risk triggers that bear on it, monitored monthly, so a slip is seen before the due date, not on it.',
]

SOURCES = [
    ('SPS Programme, Multi-Year Project Progress Report template (2023)',
     'https://nato.int/content/dam/nato/webready/documents/sps/grants/2023/SPS_MYP_Report_Technical_Progress.docx'),
    ('SPS Programme, Multi-Year Project Final Report template (2023)',
     'https://nato.int/content/dam/nato/webready/documents/sps/grants/2023/SPS_MYP_Report_Technical_Final.docx'),
    ('SPS Programme, Multi-Year Project Progress Report template (2016)',
     'https://nato.int/content/dam/nato/webready/documents/sps/grants/2016/SPS_MYP_Progress_Report.docx'),
]


def evidence(ev):
    names = {d['code']: d['name'] for d in BD.items()}
    return [f'{c} {names[c]}' for c in dict.fromkeys(re.findall(r'D\d\.\d', ev))]


def fig_ladder(path):
    ms = W.MILESTONES
    fig, ax = plt.subplots(figsize=(6.7, 3.9), dpi=300)
    ax.axvspan(LIVE[0], LIVE[1], ymin=0, ymax=1, color='#E3F4F5', zorder=0)
    ax.text((LIVE[0] + LIVE[1]) / 2, 0.55, 'Live operation\nMonths 23–34', ha='center', va='center', fontsize=8.8,
            color='#0E6E78', fontweight='bold', zorder=2)
    xs = [0] + [m for _, m, _, _ in ms] + [36]
    ys = list(range(len(ms) + 1)) + [len(ms)]
    ax.step(xs, ys, where='post', color='#1F3864', lw=2.0, zorder=3)
    for i, (code, m, name, _) in enumerate(ms, start=1):
        g = code in GATES
        ax.plot(m, i, marker='D' if g else 'o', ms=9 if g else 8, color='#D97706' if g else '#1F3864', mec='white',
                mew=1.2, zorder=5)
        ax.text(m - 0.6, i + 0.02, code + (' gate' if g else ''), ha='right', va='center', fontsize=8.8,
                fontweight='bold', color='#92400E' if g else '#1F3864', zorder=6,
                bbox=dict(boxstyle='square,pad=0.1', fc='white', ec='none'))
        ax.text(m + 0.5, i - 0.30, f'M{m}', ha='left', va='center', fontsize=8.6, color='#52606D', zorder=6)
    for m, lab in ((12, 'D7.2'), (24, 'D7.2'), (36, 'D7.3')):
        ax.plot(m, -0.42, marker='v', ms=7, color='#52606D', zorder=5)
        ax.text(m - 0.6, -0.42, lab, ha='right', va='center', fontsize=8.6, color='#52606D', zorder=6)
    ax.set_yticks(range(len(ms) + 1))
    ax.set_yticklabels(['Proposed'] + [RUNG[c] for c, *_ in ms], fontsize=9)
    for t, c in zip(ax.get_yticklabels()[1:], [c for c, *_ in ms]):
        if c in GATES:
            t.set_color('#92400E')
    ax.set_ylabel('Claim proven', fontsize=9.5)
    ax.set_xlim(0, 37.5); ax.set_ylim(-0.75, len(ms) + 0.55)
    ax.set_xticks([0, 6, 12, 18, 24, 30, 36]); ax.set_xticklabels([f'M{m}' for m in (0, 6, 12, 18, 24, 30, 36)], fontsize=9.3)
    ax.grid(axis='y', color='#e3e6ea', lw=0.8); ax.set_axisbelow(True)
    for s in ('top', 'right'):
        ax.spines[s].set_visible(False)
    ax.spines['left'].set_color('#9aa3ad'); ax.spines['bottom'].set_color('#9aa3ad')
    h = [plt.Line2D([], [], marker='o', color='#1F3864', ls='', ms=8, label='Milestone'),
         plt.Line2D([], [], marker='D', color='#D97706', ls='', ms=8, label='Evidence gate'),
         plt.Line2D([], [], marker='v', color='#52606D', ls='', ms=7, label='Report to SPS')]
    ax.legend(handles=h, loc='upper left', ncol=1, frameon=False, fontsize=8.8, bbox_to_anchor=(0.0, 1.0))
    fig.tight_layout()
    fig.savefig(path, dpi=300)
    plt.close(fig)


def low1(t):
    """Lower-case the first letter only for an ordinary word (keeps symbols such as Δₖ and acronyms)."""
    return t[0].lower() + t[1:] if t[0].isascii() and t[0].isalpha() and t[1:2].islower() else t


def hl_cells(t):
    """Re-write cell paragraphs that hold [placeholders] so that only the placeholders are highlighted."""
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                if '[' not in p.text or not p.runs:
                    continue
                size = p.runs[0].font.size.pt if p.runs[0].font.size else 9.5
                bold = bool(p.runs[0].bold); txt = p.text
                for rn in list(p.runs):
                    rn._element.getparent().remove(rn._element)
                ph(p, txt, size=size, bold=bold)
    return t


def head_repeat(t):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t


def main(out):
    import build_sustainability as BS
    cost = sum(float(c[1].replace(' ', '')) for c in BS.COSTS)
    risks = {r[0]: r for r in BR.all_risks()}
    work = tempfile.mkdtemp(); fig = os.path.join(work, 'fig_milestones.png'); fig_ladder(fig)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.7)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Milestones', bold=True, size=18, color='1F3864')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, f'CERTAIN-COAST has **{len(W.MILESTONES)} milestones** over 36 months. Each is a test with a stated pass level, '
            'evidence in one or more deliverables, a named decision-maker and a fallback. Three of them are **evidence gates** '
            '(MS3, MS4 and MS5 at Months 18, 22 and 27): live operation starts, and the guarantees are stated, only when they are '
            'passed. Months count from the start of the grant.', size=10.5)

    heading(doc, '1. Milestones at a glance', 13.5, before=4)
    rows = []
    for code, m, name, ev in W.MILESTONES:
        rows.append([code + ('\n(gate)' if code in GATES else ''), name, f'M{m}', CARD[code][0], ev.replace('α_k', 'αₖ'), RUNG[code]])
    t = head_repeat(table(doc, ['No.', 'Milestone', 'Month', 'Lead', 'Means of verification', 'Claim proven'], rows,
                          [1.5, 3.6, 1.35, 1.6, 6.15, 2.8], size=9))
    for i, (code, *_r) in enumerate(W.MILESTONES, start=1):
        shade(t.rows[i].cells[0], 'FDE7C2' if code in GATES else 'E8EDF5')
    para(doc, 3)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(fig, width=Cm(15.5))
    p = para(doc, 4)
    run(p, 'Figure 1. The evidence ladder. Each milestone raises the claim the project can make by one rung; orange diamonds '
           'are the evidence gates; the shaded band is the twelve months of live operation; triangles mark the reports to the '
           'SPS Programme.', italic=True, size=9.5, color='52606D')

    h = heading(doc, '2. Milestone descriptions', 13.5, before=8); h.paragraph_format.keep_with_next = True
    for code, m, name, ev in W.MILESTONES:
        lead, ach, test, who, rids, fall, unl = CARD[code]
        g = code in GATES
        p = para(doc, 2, before=6); p.paragraph_format.keep_with_next = True
        run(p, f'{code}  {name}', bold=True, size=11.5, color='92400E' if g else '1F3864')
        run(p, f'   Month {m} · {lead}' + (' · evidence gate' if g else ''), size=9.5, color='52606D')
        rk = '; '.join(f'{r} ({low1(risks[r][4])})' for r in rids)
        kv = [['Achieved', ach], ['Pass test', test], ['Evidence', '; '.join(evidence(ev))], ['Decided by', who],
              ['Risks watched', rk], ['If not met', fall], ['Unlocks', unl.format(cost=cost)]]
        tt = hl_cells(table(doc, ['', ''], kv, [3.0, 14.0], size=9.5))
        tt._tbl.remove(tt.rows[0]._tr)
        n = len(tt.rows)
        for ri, row in enumerate(tt.rows):
            shade(row.cells[0], 'FDE7C2' if g else 'E8EDF5')
            for pp in row.cells[0].paragraphs:
                for rn in pp.runs:
                    rn.bold = True
            if ri < n - 1:
                for pp in row.cells[0].paragraphs + row.cells[1].paragraphs:
                    pp.paragraph_format.keep_with_next = True

    h = heading(doc, '3. Rules for the evidence gates', 13.5, before=10); h.paragraph_format.keep_with_next = True
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.keep_with_next = True
    rich(p, 'Each gate has three outcomes, fixed before the gate. A **go** lets the next phase start as planned; a **go with '
            'conditions** lets it start with a stated change; a **no-go** applies a named fallback, so the project still '
            'delivers usable results. Values in brackets are proposals to be confirmed in the first data management plan.',
         size=10.5)
    t = hl_cells(head_repeat(table(doc, ['Gate', 'Go', 'Go with conditions', 'No-go and fallback'], GATE_RULES,
                                   [3.2, 4.4, 4.7, 4.7], size=9)))
    for i, r in enumerate(t.rows):
        if i:
            shade(r.cells[0], 'FDE7C2'); shade(r.cells[1], 'E6F4EC'); shade(r.cells[2], 'FFF6DB'); shade(r.cells[3], 'FBE9E7')
        for c in r.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True

    heading(doc, '4. Monitoring and reporting', 13.5, before=10)
    for x in MONITOR:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); ph(q, x, size=10.5)
    p = para(doc, 4, before=4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.keep_with_next = True
    rich(p, 'The SPS report templates ask, for each milestone and deliverable, its status, plans to complete any that are '
            'late, the milestones of the next period and a revised Gantt chart; the final report asks why any milestone was '
            'not met and how this affects the results [1–3]. The milestones are reported as follows; the reporting schedule '
            'itself is fixed in the Project Plan agreed with the SPS Programme.', size=10.5)
    head_repeat(table(doc, ['Report', 'Month', 'Milestones reported', 'Content'], REPORTS, [3.9, 1.9, 3.6, 7.6], size=9.5))

    h = heading(doc, '5. What is new in these milestones', 13.5, before=0); h.paragraph_format.page_break_before = True
    for x in NEW:
        q = para(doc, 3, 0.5); run(q, '• ', bold=True); rich(q, x, size=10.5)

    heading(doc, 'Sources', 11.5, before=8)
    for i, (txt, url) in enumerate(SOURCES, start=1):
        q = para(doc, 2, 0.5); run(q, f'[{i}] {txt}. ', size=9.5); link(q, url, url, size=8.5)
    doc.save(out)
    print('saved', out, '| milestones', len(W.MILESTONES), '| gates', len(GATES), '| running cost', cost)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Milestones.docx')
