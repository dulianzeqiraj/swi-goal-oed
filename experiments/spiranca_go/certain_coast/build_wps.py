"""CERTAIN-COAST (NATO SPS proposal): Work packages. Overview table, a work-package interaction diagram, an effort table
(person-months in kind, placeholders), and one description per work package: lead, participants, months, objective,
description, tasks, deliverables (type, access level), milestones, inputs and outputs, success criteria (from the
Objectives), risks (from the Risk Assessment) and what is new. Tasks, months, deliverables and milestones are read from
build_workplan.py; success criteria from build_objectives.py; risks from build_risk.py; leads from the Partnership.

usage: python build_wps.py OUT.docx"""
import os
import re
import sys
import tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_sota import run, para, heading, rich, table, shade  # noqa: E402
import build_workplan as W  # noqa: E402
import build_objectives as BO  # noqa: E402
import build_risk as BR  # noqa: E402

FULL = ('Certified early warning for seawater intrusion, liquefaction and extreme sea levels at the wellfields, port and '
        'energy landfall of a subsiding Adriatic coast')

LEAD = {'WP1': '[NPD team] or [additional NATO co-director]', 'WP2': '[PPD team]', 'WP3': '[PPD team]', 'WP4': '[NPD team]',
        'WP5': '[NPD team]', 'WP6': '[PPD team]', 'WP7': '[NPD team] (coordinator)'}
PARTS = {'WP1': 'PPD team; Albanian research partner', 'WP2': 'NPD team', 'WP3': 'NPD team; end-users (thresholds, T3.1)',
         'WP4': 'Albanian research partner; additional NATO co-director (T4.4)', 'WP5': 'End-users; civil protection; PPD team',
         'WP6': 'All partners; end-users (T6.2, T6.4)', 'WP7': 'Steering committee; all partners'}
OBJMAP = {'WP1': 'O1', 'WP2': 'O2', 'WP3': 'O3', 'WP4': 'O4', 'WP5': 'O5', 'WP6': 'O6', 'WP7': '—'}
IO = {
    'WP1': ('Sentinel-1 archive; GNSS and corner reflectors (WP4); tide-gauge record',
            'R(t) and land motion at each asset, compaction parameters and the head limit h_pc → WP2, WP3 (extremes), WP5 (pumping rule)'),
    'WP2': ('R(t) and compaction (WP1); heads, chloride and seismic velocity (WP4); operator records (access agreements, WP7)',
            'Calibrated coupled model and certified surrogate → WP3, WP5; sensitivities for the network design → WP4'),
    'WP3': ('Surrogate and error bound (WP2); thresholds agreed with the operators (T3.1); nonstationary extremes (WP1)',
            'Certified alarms and proofs → WP5; value of each sensor → WP4; verification statistics → WP6'),
    'WP4': ('Preliminary sensitivities (WP2); value of each sensor (WP3); site access (WP7)',
            'Data streams → WP1, WP2, WP5; seismic velocity as a water-table gauge → WP2'),
    'WP5': ('Certified alarms (WP3); data streams (WP4); surrogate (WP2)',
            'Bulletin log and outcomes → WP3 (margin update), WP6 (procedures, verification); joint pumping rule'),
    'WP6': ('Results of all work packages', 'Trained people; open software and data; operating procedures; policy brief'),
    'WP7': ('All work packages', 'Access agreements and data management plan; risk register; progress and final reports'),
}
NEWWP = {
    'WP1': 'Relative sea level referred to each asset, and the switch from elastic to inelastic compaction turned into a warning (P1, P4, M5).',
    'WP2': 'One model in which the same water table drives intrusion, compaction and the liquefaction factor of safety, with a certified fast copy (P2, P3, M4).',
    'WP3': 'Alarms whose miss rates are certified, each and jointly, also beyond the observed record (M1, M2).',
    'WP4': 'A network designed for the warnings across seven sensor types, and seismic noise used as a water-table gauge (M3, P5).',
    'WP5': 'A pumping rule checked against all three hazards at once, and certified miss rates verified against a live bulletin log (P3).',
    'WP6': 'A coast kit and a transfer test on a second coast within the project; verification statistics published openly.',
    'WP7': 'Risk indicators with numeric triggers watched like the hazards, and a data plan with a security gate.',
}
WP7_CRITERIA = ['Access agreements and data management plan delivered by Month 3', 'Progress reports delivered in Months 12 and 24, final report in Month 36',
                'Every risk trigger acted on within two weeks of being crossed']
DTYPE = {  # deliverable: (type, access level)
    'D1.1': ('Data, report', 'Public (regional); restricted (asset points)'),
    'D2.1': ('Model', 'Restricted (site); public (synthetic coast)'),
    'D2.2': ('Software', 'Restricted (site); public (synthetic coast)'),
    'D3.1': ('Report', 'Public'),
    'D3.2': ('Report', 'Public (aggregated); restricted annex'),
    'D4.1': ('Report', 'Restricted'),
    'D4.2': ('Demonstrator', 'Restricted'),
    'D5.1': ('Report', 'Public (verification statistics); restricted annex'),
    'D6.1': ('Software, data', 'Public'),
    'D6.2': ('Report', 'Restricted (procedures); public (policy brief)'),
    'D7.1': ('Report', 'Public (data plan); restricted (agreements)'),
    'D7.2': ('Report', 'Restricted (to NATO SPS)'),
    'D7.3': ('Report', 'Public summary; restricted full report'),
}


def deliv_wp():
    m = {}
    for wp, *_rest in W.WPS:
        for code, name, month in _rest[-1]:
            m[code] = wp
    return m


def milestones_for(wp):
    dw = deliv_wp(); out = []
    for code, month, name, ev in W.MILESTONES:
        if any(dw.get(d) == wp for d in re.findall(r'D\d\.\d', ev)):
            out.append(f'{code} (Month {month}): {name}')
    return out


def risks_for(wp):
    return [f'{r[0]} {r[1][0].lower() + r[1][1:]}' for r in BR.all_risks() if r[7] == wp]


def fig_links(path):
    Wd, Ht = 6.7, 3.9
    fig = plt.figure(figsize=(Wd, Ht), dpi=300)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, Wd); ax.set_ylim(0, Ht); ax.axis('off')

    def box(x, y, w, h, wp, sub):
        col = W.COL[wp]
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0,rounding_size=0.07', fc=col, ec='none'))
        ax.text(x + w / 2, y + h - 0.2, wp, fontsize=11, fontweight='bold', color='white', ha='center', va='center')
        ax.text(x + w / 2, y + (h - 0.3) / 2, sub, fontsize=8.8, color='white', ha='center', va='center', linespacing=1.15)

    def arr(x0, y0, x1, y1, lab=None, lx=None, ly=None, rad=0.0, ha='left'):
        ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle='-|>', mutation_scale=12, color='#1f2328', lw=1.3,
                                     connectionstyle=f'arc3,rad={rad}'))
        if lab:
            ax.text(lx, ly, lab, fontsize=8.6, color='#1f2328', ha=ha, va='center', linespacing=1.15)

    bw, bh, yb = 1.02, 0.95, 1.40
    xs = [0.12, 1.47, 2.82, 4.17, 5.52]
    names = [('WP1', 'Coastal\nmotion'), ('WP2', 'Coupled model\n+ surrogate'), ('WP3', 'Certified\nalarms'),
             ('WP5', 'Live\ndemonstration'), ('WP6', 'Training,\ntransfer')]
    for x, (wp, sub) in zip(xs, names):
        box(x, yb, bw, bh, wp, sub)
    for i in range(4):
        arr(xs[i] + bw + 0.02, yb + bh / 2, xs[i + 1] - 0.02, yb + bh / 2)
    box(0.12, 3.05, 3.72, 0.72, 'WP4', 'Monitoring network designed for the warnings')
    arr(0.63, 3.04, 0.63, yb + bh + 0.02, 'GNSS,\nreflectors', 0.70, 2.70)
    arr(1.98, 3.04, 1.98, yb + bh + 0.02, 'heads, chloride,\nseismic velocity', 2.05, 2.70)
    arr(3.33, yb + bh + 0.02, 3.33, 3.04, 'value of\neach sensor', 3.40, 2.70)
    arr(xs[3] + 0.30, yb - 0.02, xs[2] + bw - 0.30, yb - 0.02, rad=-0.5)
    ax.text((xs[2] + xs[3] + bw) / 2, 0.97, 'outcomes → margin update', fontsize=8.6, ha='center', va='center', color='#1f2328')
    ax.add_patch(FancyBboxPatch((0.12, 0.10), 6.42, 0.42, boxstyle='round,pad=0,rounding_size=0.06', fc=W.COL['WP7'], ec='none'))
    ax.text(3.33, 0.31, 'WP7  Management, data access and risk, across all work packages', fontsize=9.6, fontweight='bold',
            color='white', ha='center', va='center')
    fig.savefig(path, dpi=300)
    plt.close(fig)


def hl_table(doc, header, rows, widths, size=9.5):
    t = table(doc, header, rows, widths, size=size)
    for r in t.rows[1:]:
        for c in r.cells:
            for p in c.paragraphs:
                for rn in p.runs:
                    if '[' in rn.text:
                        rn.font.highlight_color = WD_COLOR_INDEX.YELLOW
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t


def kv_table(doc, rows, size=10):
    t = table(doc, ['', ''], rows, [4.0, 13.0], size=size)
    t._tbl.remove(t.rows[0]._tr)
    for r in t.rows:
        shade(r.cells[0], 'E8EDF5')
        for p in r.cells[0].paragraphs:
            for rn in p.runs:
                rn.bold = True
        for p in r.cells[1].paragraphs:
            for rn in p.runs:
                if '[' in rn.text:
                    rn.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return t


def sub(doc, text):
    p = para(doc, 2, before=6); run(p, text, bold=True, size=11, color='2E6DB4'); p.paragraph_format.keep_with_next = True


def bullets(doc, items, size=10):
    for t in items:
        p = para(doc, 1, 0.5); run(p, '• ', bold=True, size=size); rich(p, t, size=size)


def main(out):
    work = tempfile.mkdtemp(); fig = os.path.join(work, 'fig_wp_links.png'); fig_links(fig)
    objs = {o[0]: o for o in BO.OBJ}
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.7)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, 'CERTAIN-COAST — Work Packages', bold=True, size=18, color='1F3864')
    p = para(doc, 8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(p, FULL, bold=True, size=12, color='1F3864')
    p = para(doc, 6); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rich(p, 'CERTAIN-COAST is organised in seven work packages over 36 months: four build the science and the system (WP1 coastal '
            'motion, WP2 coupled model and surrogate, WP3 certified alarms, WP4 monitoring network), one runs it live with the '
            'operators (WP5), one makes it last (WP6), and one manages data, access and risk (WP7). Each work package below gives '
            'its tasks, deliverables, milestones, the inputs it needs and the outputs it gives, its success criteria and risks, and '
            'what is new in it. Months count from the start of the grant.', size=10.5)

    heading(doc, 'Overview', 13.5, before=4)
    rows = []
    for wp, title, obj, span, desc, tasks, dels in W.WPS:
        dl = sorted({f'{c} (M{m})' for c, n, m in dels}, key=lambda x: (x.split()[0], int(x.split('M')[-1][:-1])))
        rows.append([wp, title, LEAD[wp], f'M{span[0]}–M{span[1]}', OBJMAP[wp], ', '.join(dl), '[PM]'])
    rows.append(['', 'Total', '', 'M1–M36', '', '14 deliverables, 7 milestones', '[PM]'])
    hl_table(doc, ['', 'Work package', 'Lead', 'Months', 'Objective', 'Deliverables', 'Effort'], rows,
             [1.3, 3.6, 3.1, 1.8, 1.9, 3.8, 1.5], size=9)
    p = para(doc, 4, before=2)
    run(p, 'Effort in person-months is contributed in kind, because salaries are not eligible for SPS funding; the teams fill '
           'it in, also in the box of each work package.', italic=True, size=9, color='52606D')
    p = para(doc, 2); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(fig, width=Cm(16.0))
    p = para(doc, 4)
    run(p, 'Figure 1. How the work packages feed each other. The live operation (WP5) feeds outcomes back to the alarms (WP3), '
           'and the alarms tell the network (WP4) which sensors are worth most.', italic=True, size=9.5, color='52606D')

    for wp, title, obj, span, desc, tasks, dels in W.WPS:
        p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
        p = para(doc, 4); run(p, f'{wp}  {title}', bold=True, size=14, color=W.COL[wp].lstrip('#').upper())
        kv_table(doc, [['Lead', LEAD[wp]], ['Participants', PARTS[wp]], ['Months', f'Month {span[0]} to Month {span[1]}'],
                       ['Objective', f'{OBJMAP[wp]}: {objs[OBJMAP[wp]][1]}' if OBJMAP[wp] in objs else 'Management of the project'],
                       ['Effort', '[person-months, in kind]']], size=9.5)
        sub(doc, 'Description')
        p = para(doc, 3); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; rich(p, desc, size=10)
        sub(doc, 'Tasks')
        hl_table(doc, ['Task', 'Description', 'Months'], [[c, n, f'M{a}–M{b}'] for c, n, (a, b) in tasks], [1.4, 13.2, 2.4], size=9.5)
        sub(doc, 'Deliverables')
        seen, drows = set(), []
        for c, n, m in dels:
            if (c, m) in seen:
                continue
            seen.add((c, m)); ty, lvl = DTYPE.get(c, ('Report', 'Public'))
            drows.append([c, n, ty, lvl, f'M{m}'])
        hl_table(doc, ['', 'Deliverable', 'Type', 'Access', 'Month'], drows, [1.3, 6.0, 2.6, 5.4, 1.7], size=9.5)
        ms = milestones_for(wp)
        if ms:
            sub(doc, 'Milestones'); bullets(doc, ms)
        sub(doc, 'Inputs and outputs')
        bullets(doc, [f'**Needs:** {IO[wp][0]}', f'**Gives:** {IO[wp][1]}'])
        sub(doc, 'Success criteria')
        bullets(doc, objs[OBJMAP[wp]][3] if OBJMAP[wp] in objs else WP7_CRITERIA)
        rk = risks_for(wp)
        if rk:
            sub(doc, 'Main risks (Risk Assessment)'); bullets(doc, rk)
        sub(doc, 'What is new'); bullets(doc, [NEWWP[wp]])
    doc.save(out)
    print('saved', out)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'CERTAIN-COAST_Work_Packages.docx')
