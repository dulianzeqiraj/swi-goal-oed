"""Review 62: earthquakes, groundwater and surface deformation in coastal aquifers. 30 verified papers (2020-2026)
from Oxford, Caltech, MIT and Harvard, each with download links, key finding and what is new for the coastal system;
ends with the innovation synthesis. Data in r62_data.py. Writes a Word file (Arial, large type, clickable links)."""
import os
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import r62_data as D  # noqa: E402

FONT = 'Arial'


def run(p, text, bold=False, italic=False, size=11, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = FONT; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def link(p, url, text, size=11):
    part = p.part
    rid = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), FONT); fonts.set(qn('w:hAnsi'), FONT); rpr.append(fonts)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1F5FBF'); rpr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)


def para(doc, space_after=4, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def heading(doc, text, size=14, color='1F3864', before=10):
    p = para(doc, 6); p.paragraph_format.space_before = Pt(before)
    run(p, text, bold=True, size=size, color=color)
    return p


def main(out):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11)

    p = para(doc, 4); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, D.TITLE, bold=True, size=17, color='1F3864')
    p = para(doc, 10); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, D.SUBTITLE, italic=True, size=11.5, color='52606D')
    heading(doc, 'Why this matters for the coastal system', 13.5, before=4)
    for t in D.INTRO:
        p = para(doc, 5); run(p, t, size=11)

    for blk_title, items in D.GROUPS:
        heading(doc, blk_title, 13.5)
        for it in items:
            n = it['n']
            p = para(doc, 2); p.paragraph_format.keep_with_next = True
            run(p, f'{n}. ', bold=True, size=11.5, color='1F3864'); run(p, it['cite'], bold=True, size=11)
            p = para(doc, 2, 0.5)
            run(p, 'Institution: ', bold=True, size=10.5, color='52606D'); run(p, it['inst'], size=10.5, color='52606D')
            p = para(doc, 2, 0.5)
            run(p, 'Download: ', bold=True, size=10.5)
            link(p, it['doi'], it['doi'].replace('https://', ''), 10.5)
            if it.get('free'):
                run(p, '   |   free copy: ', size=10.5)
                link(p, it['free'], it['free_label'], 10.5)
            else:
                run(p, '   |   ' + it.get('free_note', 'no free copy found'), italic=True, size=10.5, color='52606D')
            p = para(doc, 2, 0.5)
            run(p, 'Key finding: ', bold=True, size=11); run(p, it['finding'], size=11)
            p = para(doc, 8, 0.5)
            run(p, 'What is new for our coastal system: ', bold=True, size=11, color='1E8A5A'); run(p, it['new'], size=11)

    heading(doc, 'Local context: the 2019 Durrës earthquake (studies not from the four universities)', 13.5)
    for t in D.LOCAL:
        url = t.rsplit(' ', 1)[-1]
        p = para(doc, 3, 0.5); run(p, '• ', bold=True, size=11); run(p, t[:-len(url)], size=10.5); link(p, url, url.replace('https://', ''), 10.5)
    heading(doc, D.SYN_TITLE, 14.5, color='B03A3A', before=14)
    for kind, t in D.SYNTHESIS:
        if kind == 'h':
            heading(doc, t, 12.5, before=6)
        elif kind == 'b':
            p = para(doc, 3, 0.5); run(p, '• ', bold=True, size=11); run(p, t, size=11)
        else:
            p = para(doc, 5); run(p, t, size=11)
    if getattr(D, 'NOTE', None):
        p = para(doc, 4); p.paragraph_format.space_before = Pt(10)
        run(p, D.NOTE, italic=True, size=10, color='52606D')
    doc.save(out)
    print('saved', out, sum(len(i) for _, i in D.GROUPS), 'papers')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'Review62_Earthquakes_Groundwater_Deformation.docx')
