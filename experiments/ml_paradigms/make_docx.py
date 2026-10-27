"""Put the rendered diagram on one A3 landscape page of a Word document."""
import sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

png, out = sys.argv[1], sys.argv[2]
d = Document()
s = d.sections[0]
s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = Mm(420), Mm(297)
s.left_margin = s.right_margin = Mm(15)
s.top_margin = s.bottom_margin = Mm(8)
s.header_distance = s.footer_distance = Mm(0)
p = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pf = p.paragraph_format
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing = 1.0
p.add_run().add_picture(png, width=Mm(390))
d.core_properties.title = sys.argv[3] if len(sys.argv) > 3 else 'Diagram'
d.core_properties.subject = 'Machine-learning paradigm diagram'
d.save(out)
print('wrote', out)
