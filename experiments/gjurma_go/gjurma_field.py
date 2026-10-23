"""GJURMA-GO, field part: the missing term of the transport equation, read from the multi-species chemistry of the
Fushe-Kuqe aquifer (sheet 5 of data/Fushe_Kuqe_All_Data.xlsx: Cenameri & Beqiraj 2016, four wells, major ions).

For each saline well the seawater fraction f is taken from Cl (conservative) between the fresh end-member (well 193)
and seawater; the departure of every other ion from conservative mixing (meq/L) is the integrated closure term of the
transport equation. It is split on five reactions (stoichiometric matrix N, meq-based):
  R1 Na-Ca exchange (Na -1, Ca +1), R2 Na-Mg exchange (Na -1, Mg +1), R3 K-Ca exchange (K -1, Ca +1),
  R4 sulfate reduction (SO4 -1, HCO3 +1), R5 calcite dissolution (Ca +1, HCO3 +1).
N has rank 5 on the six non-conservative ions; its left null vector is the charge balance, so the residual of the
split is the charge imbalance of the departures. Uncertainty: analytical error, fresh end-member, and seawater salinity
between standard seawater and the Adriatic. Writes gjurma_field.json.
"""
import json, os
import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, '..', '..', 'data', 'Fushe_Kuqe_All_Data.xlsx')
IONS = ['Cl', 'Na', 'K', 'Ca', 'Mg', 'SO4', 'HCO3']
MW = np.array([35.453, 22.990, 39.098, 40.078, 24.305, 96.06, 61.017])
ZA = np.array([1, 1, 1, 2, 2, 2, 1], float)                    # |charge|
SEA_MGL = np.array([19353., 10781., 399., 412., 1284., 2712., 145.])   # standard seawater
REACTIONS = ['Na-Ca exchange', 'Na-Mg exchange', 'K-Ca exchange', 'sulfate reduction', 'calcite dissolution']
# rows: Na, K, Ca, Mg, SO4, HCO3 (meq/L); columns: R1..R5
N = np.array([[-1, -1, 0, 0, 0],
              [0, 0, -1, 0, 0],
              [1, 0, 1, 0, 1],
              [0, 1, 0, 0, 0],
              [0, 0, 0, -1, 0],
              [0, 0, 0, 1, 1]], float)


def read_sheet5():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb['5_Cenameri2016_Chemistry']
    rows = list(ws.iter_rows(values_only=True))
    ih = next(i for i, r in enumerate(rows) if r and r[0] == 'Well_no')
    col = {h: i for i, h in enumerate(rows[ih])}
    out = {}
    for r in rows[ih + 1:]:
        if r[0] is None or not str(r[0]).strip().isdigit():
            continue
        out[str(r[0]).strip()] = (np.array([float(r[col[i + '_mg_L']]) for i in IONS]), r[col['Transect_position']])
    return out


def read_sheet6():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    rows = list(wb['6_Cenameri2016_Cl_Trend'].iter_rows(values_only=True))
    ih = next(i for i, r in enumerate(rows) if r and r[0] == 'Well_no')
    out = {}
    for r in rows[ih + 1:]:
        if r[0] is None or not str(r[0]).strip().isdigit():
            continue
        out[str(r[0]).strip()] = {1984: float(r[1]), 1999: float(r[2]), 2001: float(r[3])}
    return out


def split(mgl, fresh_mgl, sea_mgl):
    """Seawater fraction, departures from mixing (meq/L, six ions), reaction extents (meq/L) and the residual."""
    m = mgl / MW * ZA; fr = fresh_mgl / MW * ZA; se = sea_mgl / MW * ZA
    f = (m[0] - fr[0]) / (se[0] - fr[0])
    d = (m - (f * se + (1 - f) * fr))[1:]
    xi, *_ = np.linalg.lstsq(N, d, rcond=None)
    return f, d, xi, d - N @ xi


if __name__ == '__main__':
    S5 = read_sheet5(); S6 = read_sheet6()
    fresh = S5['193'][0]
    out = dict(source='Cenameri & Beqiraj (2016), Bull. Geol. Soc. Greece 50(2), 665-670 (sheet 5 and 6)',
               reactions=REACTIONS, N=N.tolist(), rank=int(np.linalg.matrix_rank(N)), wells={}, cl_trend=S6)
    # left null space (reaction invariant among the six ions): charge balance
    u, s, vt = np.linalg.svd(N.T)
    out['left_null'] = (vt[-1] / np.abs(vt[-1]).max()).round(6).tolist()
    rng = np.random.default_rng(2016)
    for w in ('503', '341', '340'):
        mgl, pos = S5[w]
        f, d, xi, res = split(mgl, fresh, SEA_MGL)
        draws = []
        for _ in range(4000):
            sal = rng.uniform(1.0, 21500.0 / 19353.0)                      # standard -> Adriatic salinity
            fr = fresh * np.exp(0.10 * rng.standard_normal(7))
            mg = mgl * np.exp(0.03 * rng.standard_normal(7))
            fd, dd, xd, rd = split(mg, fr, SEA_MGL * sal)
            draws.append(np.concatenate([[fd], dd, xd, [np.linalg.norm(rd)]]))
        D = np.array(draws)
        q = lambda k: [float(np.quantile(D[:, k], 0.05)), float(np.quantile(D[:, k], 0.5)), float(np.quantile(D[:, k], 0.95))]
        out['wells'][w] = dict(position=pos, f=float(f), delta_meq=dict(zip(IONS[1:], d.round(4).tolist())),
                               extents_meq=dict(zip(REACTIONS, xi.round(4).tolist())),
                               residual_meq=float(np.linalg.norm(res)), delta_norm=float(np.linalg.norm(d)),
                               explained=float(1 - np.linalg.norm(res) ** 2 / np.linalg.norm(d) ** 2),
                               charge_imbalance_of_deltas=float(d[:4].sum() - d[4:].sum()),
                               f_q=q(0), extents_q={r: q(7 + j) for j, r in enumerate(REACTIONS)})
        print(w, pos, 'f %.2f%%' % (100 * f), 'extents', np.round(xi, 2), 'residual %.3f of %.2f' % (np.linalg.norm(res), np.linalg.norm(d)))
    # analyses: charge-balance errors
    cbe = {}
    for w, (mgl, _) in S5.items():
        m = mgl / MW * ZA
        cat = m[1:5].sum(); an = m[0] + m[5] + m[6]
        cbe[w] = float(100 * (cat - an) / (cat + an))
    out['cbe_pct'] = cbe
    json.dump(out, open(os.path.join(HERE, 'gjurma_field.json'), 'w'), indent=1)
    print('charge-balance errors (%):', {k: round(v, 2) for k, v in cbe.items()})
    print('left null vector', out['left_null'])
